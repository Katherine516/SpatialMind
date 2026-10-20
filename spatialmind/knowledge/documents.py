"""Bounded, local text extraction. Documents are evidence, never executable tools."""

import csv
import hashlib
import io
import multiprocessing
import re
import zipfile
from pathlib import Path

SUPPORTED = {".txt", ".md", ".csv", ".tsv", ".docx", ".xlsx", ".pptx", ".pdf"}
MAX_FILE_BYTES = 25 * 1024 ** 2
MAX_TEXT_BYTES = 2 * 1024 ** 2
MAX_ARCHIVE_BYTES = 50 * 1024 ** 2
MAX_UNITS = 50000
PARSE_TIMEOUT = 30


class ExtractedText:
    def __init__(self):
        self.parts = []
        self.size = 0

    def add(self, locator, text):
        text = str(text).replace("\x00", "").strip()
        if not text:
            return
        # Repeat the locator in each retrieval-sized paragraph, even for long pages.
        for paragraph in re.split(r"\n\s*\n", text):
            for offset in range(0, len(paragraph), 1800):
                part = "[%s]\n%s" % (locator, paragraph[offset:offset + 1800])
                self.size += len(part.encode("utf-8")) + 2
                if self.size > MAX_TEXT_BYTES:
                    raise ValueError("Extracted text exceeds 2 MiB. Split the document before importing.")
                self.parts.append(part)


def _check_archive(payload, suffix):
    required = {".docx": "word/document.xml", ".xlsx": "xl/workbook.xml", ".pptx": "ppt/presentation.xml"}
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(e.file_size for e in entries) > MAX_ARCHIVE_BYTES:
            raise ValueError("Office archive exceeds the 50 MiB expanded / 5,000 entry limit.")
        names = {e.filename for e in entries}
        if required[suffix] not in names or "[Content_Types].xml" not in names:
            raise ValueError("The contents do not match the selected Office format.")
        if any(e.flag_bits & 1 for e in entries):
            raise ValueError("Encrypted Office documents are not supported.")
        if any("vbaproject" in name.lower() for name in names):
            raise ValueError("Macro-bearing documents are not supported. Export a macro-free copy.")


def extract_document(filename, payload):
    """Extract a single document; production callers use the isolated wrapper below."""
    name = filename.replace("\\", "/").rsplit("/", 1)[-1]
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError("Unsupported format. Use DOCX, XLSX, PPTX, PDF, CSV, TSV, TXT or MD; convert legacy DOC/XLS/PPT first.")
    if not payload or len(payload) > MAX_FILE_BYTES:
        raise ValueError("Each document must be nonempty and at most 25 MiB.")
    result = ExtractedText()
    warnings = []
    parser = "stdlib"
    if suffix in {".docx", ".xlsx", ".pptx"}:
        _check_archive(payload, suffix)
        warnings.append("Text extraction only: figures, charts, embedded objects and visual layout are not interpreted.")
    if suffix in {".txt", ".md", ".csv", ".tsv"}:
        encoding = "utf-16" if payload.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        try:
            text = payload.decode(encoding)
        except UnicodeError:
            raise ValueError("Text files must be UTF-8 or BOM-marked UTF-16. Re-export the file with that encoding.")
        if "\x00" in text:
            raise ValueError("Binary content detected in a text file.")
        if suffix in {".csv", ".tsv"}:
            cells = 0
            for row_number, row in enumerate(csv.reader(io.StringIO(text), delimiter="\t" if suffix == ".tsv" else ",", strict=True), 1):
                cells += max(len(row), 1)
                if cells > MAX_UNITS:
                    raise ValueError("Table exceeds 50,000 fields/rows. Import a smaller knowledge table, not an expression matrix.")
                result.add("Row %d" % row_number, " | ".join("%d: %s" % (i, value) for i, value in enumerate(row, 1)))
        else:
            for number, line in enumerate(text.splitlines(), 1):
                result.add("Line %d" % number, line)
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        parser = "python-docx"
        doc = Document(io.BytesIO(payload))
        for index, block in enumerate(doc.iter_inner_content(), 1):
            if isinstance(block, Paragraph):
                result.add("Paragraph %d" % index, block.text)
            elif isinstance(block, Table):
                for row_number, row in enumerate(block.rows, 1):
                    result.add("Table %d row %d" % (index, row_number), " | ".join(cell.text for cell in row.cells))
        warnings.append("Word headers, footers, comments, tracked deletions and nested tables are not imported; inspect the preview.")
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
        parser = "openpyxl"
        book = load_workbook(io.BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        cells = 0
        try:
            for sheet in book.worksheets:
                # Do not trust producer-supplied dimensions to define which rows exist.
                sheet.reset_dimensions()
                for number, row in enumerate(sheet.iter_rows(), 1):
                    cells += max(len(row), 1)
                    if cells > MAX_UNITS:
                        raise ValueError("Workbook exceeds 50,000 cells/rows. Import selected sheets or a smaller table.")
                    values = ["%s: %s" % (get_column_letter(i), cell.value) for i, cell in enumerate(row, 1) if cell.value is not None]
                    result.add("Sheet %s row %d" % (sheet.title, number), " | ".join(values))
        finally:
            book.close()
        warnings.append("All worksheets, including hidden ones, are read. Formulas are literal text, never calculated; cached results and cell comments are not imported.")
    elif suffix == ".pptx":
        from pptx import Presentation
        parser = "python-pptx"
        deck = Presentation(io.BytesIO(payload))
        if len(deck.slides) > 200:
            raise ValueError("Presentations are limited to 200 slides.")

        def shapes_text(shapes, locator):
            for shape in shapes:
                if shape.has_text_frame:
                    result.add(locator, shape.text)
                if shape.has_table:
                    for number, row in enumerate(shape.table.rows, 1):
                        result.add(locator + " table row %d" % number, " | ".join(cell.text for cell in row.cells))
                if hasattr(shape, "shapes"):
                    shapes_text(shape.shapes, locator)

        for number, slide in enumerate(deck.slides, 1):
            shapes_text(slide.shapes, "Slide %d" % number)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                result.add("Slide %d notes" % number, slide.notes_slide.notes_text_frame.text)
        warnings.append("Speaker notes and hidden slides are included. Review them before approving the source.")
    else:
        from pypdf import PdfReader
        parser = "pypdf"
        if not payload.startswith(b"%PDF-"):
            raise ValueError("The file does not have a PDF signature.")
        reader = PdfReader(io.BytesIO(payload))
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported. Export an unlocked copy you are authorized to use.")
        if len(reader.pages) > 200:
            raise ValueError("PDFs are limited to 200 pages.")
        empty = []
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if not text.strip():
                empty.append(number)
            result.add("Page %d" % number, text)
        warnings.append("PDF text order and tables may be imperfect. Figures are not interpreted and OCR is not performed.")
        if empty:
            warnings.append("No extractable text on pages: %s. OCR those pages separately if needed." % ", ".join(map(str, empty)))
    if not result.parts:
        raise ValueError("No readable text found. Image-only/scanned files need OCR before import.")
    from importlib.metadata import version
    return {"body": "\n\n".join(result.parts), "metadata": {
        "filename": name, "format": suffix.lstrip("."), "file_sha256": hashlib.sha256(payload).hexdigest(),
        "file_bytes": len(payload), "parser": parser,
        "parser_version": version(parser) if parser != "stdlib" else "builtin",
        "warnings": warnings, "extracted_bytes": result.size, "original_retained": False,
    }}


def _worker(connection, filename, payload):
    try:
        # POSIX CPU limit complements the parent's wall-clock timeout. This is
        # process isolation, not an OS security sandbox or a portable RAM quota.
        try:
            import resource
            resource.setrlimit(resource.RLIMIT_CPU, (20, 20))
        except (ImportError, ValueError, OSError):
            pass
        message = {"result": extract_document(filename, payload)}
    except ImportError as exc:
        message = {"error": "Document parser unavailable. Install requirements-app.txt (%s)." % exc.name}
    except ValueError as exc:
        message = {"error": str(exc)}
    except Exception:
        message = {"error": "Could not read this document. Check its format, remove encryption or export a fresh copy."}
    try:
        connection.send(message)
    except (BrokenPipeError, EOFError, OSError):
        pass  # The parent may have cancelled or timed out.
    finally:
        connection.close()


def extract_isolated(filename, payload):
    context = multiprocessing.get_context("spawn")
    parent, child = context.Pipe(duplex=False)
    process = context.Process(target=_worker, args=(child, filename, payload), daemon=True)
    try:
        process.start()
        child.close()
        if not parent.poll(PARSE_TIMEOUT):
            raise ValueError("Document extraction timed out after 30 seconds. Split or simplify the document.")
        try:
            message = parent.recv()
        except EOFError:
            raise ValueError("Document extraction stopped unexpectedly. Split or re-export the document.")
        if "error" in message:
            raise ValueError(message["error"])
        return message["result"]
    finally:
        parent.close()
        child.close()
        if process.pid is not None:
            process.join(timeout=1)
            if process.is_alive():
                process.terminate()
                process.join(timeout=2)
            if process.is_alive():
                process.kill()
                process.join()
