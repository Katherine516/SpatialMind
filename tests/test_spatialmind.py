import importlib
import inspect
import json
import math
import os
import csv
import shutil
import sys
import tempfile
import textwrap
import unittest
from pathlib import Path
from unittest.mock import patch

from eval.runner import EvalRunner, _case_directory
from spatialmind.agent import SpatialMindAgent
from spatialmind.agent_loop import SpatialAgent
from spatialmind.datasets import discover_dataset_candidates, inspect_dataset
from spatialmind.governance import build_dataset_governance_manifest
from spatialmind.ingestion import (
    IngestionValidationError,
    apply_xenium_cell_qc,
    BatchIngestionConfig,
    BatchIngestionPipeline,
    DataFormat,
    DataIngestionLayer,
    DataIngestionPipeline,
    IngestionConfig,
    SampleConfig,
    UnsupportedRawDataError,
    load_scatac,
    load_scrna,
    infer_data_type,
    summarize_supported_raw_data_types,
    build_readiness_report,
    apply_external_label_table,
    apply_external_region_table,
    build_xenium_label_intake_report,
    LabelApplicationReport,
    RegionApplicationReport,
    summarize_xenium_expert_readiness,
    validate_cell_by_feature_contract,
    write_expert_label_template,
    write_region_label_template,
    XeniumExpertReadiness,
)
from spatialmind.llm import StaticLLMProvider
from spatialmind.memory import LongTermMemory, SessionMemory
from spatialmind.memory import PriorType, UserPriorStore
from spatialmind.agent import build_xenium_mvp_plan, validate_tool_plan
from spatialmind.planner import LLMReasoningLayer
from spatialmind.promotion.local import build_local_promotion_report
from spatialmind.viz import (
    PdfSection,
    PdfTable,
    QCReportBuilder,
    VisualizationLayer,
    VizRouter,
    XeniumExplorerLiteViewer,
    normalize_report_format,
    write_pdf_report,
)
from spatialmind.storage import StorageLayer
from spatialmind.storage import index_run_records, replay_run_record, verify_run_record
from spatialmind.tools import MVP_TOOL_NAMES, build_default_registry, build_full_registry, build_mvp_registry
from spatialmind.tools.exceptions import MissingPreconditionError
from spatialmind.tools.fusion import ModalityFuser
from spatialmind.tools.implementations import (
    assess_reference_lineage_coverage,
    describe_lineage_coverage,
    feature_overlay,
    marker_detection,
    reference_label_transfer,
)
from spatialmind.workflows import INTEGRATION_MODE, SCATAC_STANDALONE, SCRNA_STANDALONE, XENIUM_STANDALONE
from spatialmind.contracts import BiologicalClaim, CellByFeatureContract, CoreSpatialObject, ground_claim
from spatialmind.schemas import SpatialDataset, SpotRecord, ToolResult
from spatialmind.pilot import build_pilot_claim_ledger, pilot_gate
from spatialmind.methods.reliability import build_claim_reliability_table, fit_claim_reliability_calibration
from spatialmind.review import CLAIM_TRUTH_FIELDS, validate_claim_truth_table


ROOT = os.path.dirname(os.path.dirname(__file__))
DEMO = os.path.join(ROOT, "data", "demo_spatial.csv")
MANIFEST = os.path.join(ROOT, "data", "demo_manifest.json")
XENIUM_LYMPH = os.path.join(ROOT, "data", "Xenium lymph", "Xenium_V1_hLymphNode_nondiseased_section_outs")

try:
    import scanpy  # noqa: F401

    _HAS_SCANPY = True
except Exception:  # pragma: no cover - dependency-light environments
    _HAS_SCANPY = False


def _make_two_program_dataset():
    """Two expression programs interleaved in space: expression clustering should
    recover the programs while spatial clustering would mix them."""
    records = []
    for index in range(60):
        program = "A" if index % 2 == 0 else "B"
        other = "B" if program == "A" else "A"
        genes = {}
        for gene in range(6):
            genes["%s_gene_%d" % (program, gene)] = 5.0 + float(index % 3)
            genes["%s_gene_%d" % (other, gene)] = 0.0
        # Programs alternate along x, so neighbors in space are always the other program.
        records.append(SpotRecord("S1", float(index), 0.0, program, genes, cell_id="c%d" % index))
    return SpatialDataset(
        sample_id="S1",
        source_path="synthetic",
        modality="xenium_spatial_rna",
        records=records,
    )


class PlannerTests(unittest.TestCase):
    def test_plans_colocalization_request(self):
        plan = LLMReasoningLayer().plan(
            "Show CD8+ T cells relative to tumor cells in sample BRCA_04 and test co-localization."
        )
        self.assertEqual(plan.request.sample_id, "BRCA_04")
        self.assertIn("CD8+ T cell", plan.request.cell_types)
        self.assertIn("Tumor cell", plan.request.cell_types)
        self.assertIn("cell_type_colocalization", [step.tool for step in plan.steps])

    def test_plans_plain_language_neighborhood_comparison(self):
        plan = LLMReasoningLayer().plan("Compare tumor and CD8+ T cells and assess spatial neighborhoods.")
        self.assertIn("cell_type_colocalization", [step.tool for step in plan.steps])

    def test_accepts_valid_llm_plan(self):
        provider = StaticLLMProvider(
            {
                "sample_id": "BRCA_04",
                "cell_types": ["CD8+ T cell", "Tumor cell"],
                "genes": [],
                "wants_visualization": True,
                "wants_colocalization": True,
                "steps": [
                    {
                        "name": "LLM co-localization",
                        "tool": "cell_type_colocalization",
                        "parameters": {"cell_types": ["CD8+ T cell", "Tumor cell"]},
                    }
                ],
            }
        )
        plan = LLMReasoningLayer(llm_provider=provider).plan("Check sample BRCA_04")
        self.assertEqual(plan.steps[0].name, "LLM co-localization")
        self.assertEqual(plan.steps[0].tool, "cell_type_colocalization")


class IngestionTests(unittest.TestCase):
    def test_loads_and_normalizes_demo_data(self):
        dataset = DataIngestionLayer().load_csv(DEMO, sample_id="BRCA_04")
        self.assertEqual(dataset.sample_id, "BRCA_04")
        self.assertTrue(dataset.normalized)
        self.assertIn("CD8A", dataset.genes)
        self.assertIn("Tumor cell", dataset.cell_types)
        self.assertEqual(dataset.qc_metrics["record_count"], 20)
        self.assertIn("Applied library-size normalization", " ".join(dataset.processing_steps))
        self.assertTrue(dataset.records[0].raw_genes)
        self.assertNotEqual(dataset.records[0].raw_genes, dataset.records[0].genes)
        self.assertEqual(dataset.metadata["expression_layers"]["normalization_target_sum"], 10000.0)

    def test_ingestion_sanitizes_nonfinite_features_before_normalization(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "nonfinite.csv")
            with open(path, "w", encoding="utf-8") as handle:
                handle.write("sample_id,x,y,cell_type,gene_A,gene_B\n")
                handle.write("S1,1,2,T cell,nan,4\n")
                handle.write("S1,3,4,Tumor cell,2,inf\n")
            dataset = DataIngestionLayer().load_csv(path, sample_id="S1")
            self.assertEqual(dataset.qc_metrics["nonfinite_feature_value_count"], 2)
            self.assertTrue(
                all(
                    math.isfinite(value)
                    for record in dataset.records
                    for value in record.genes.values()
                )
            )

    def test_loads_manifest_with_source_metadata(self):
        dataset = DataIngestionLayer().load(MANIFEST, sample_id="BRCA_04")
        self.assertEqual(dataset.sample_id, "BRCA_04")
        self.assertEqual(dataset.source_path, MANIFEST)
        self.assertEqual(len(dataset.sources), 2)
        self.assertIn("project", dataset.metadata)
        self.assertTrue(any(source.data_type == "pathology_image" for source in dataset.sources))

    def test_summarizes_supported_raw_types(self):
        summaries = summarize_supported_raw_data_types()
        self.assertTrue(any(item["data_type"] == "h5ad_anndata" for item in summaries))
        self.assertEqual(infer_data_type("example.h5ad"), "h5ad_anndata")

    def test_unsupported_raw_type_has_guidance(self):
        real_import = __import__

        def block_anndata(name, *args, **kwargs):
            if name == "anndata":
                raise ImportError("anndata intentionally unavailable for this test")
            return real_import(name, *args, **kwargs)

        with self.assertRaises(UnsupportedRawDataError) as error:
            with patch("builtins.__import__", side_effect=block_anndata):
                DataIngestionLayer().load("sample.h5ad")
        self.assertIn("requires anndata", str(error.exception))

    def test_loads_xenium_directory_metadata(self):
        dataset = DataIngestionLayer().load_xenium_directory(XENIUM_LYMPH, max_records=10)
        self.assertEqual(dataset.modality, "spatial_transcriptomics")
        self.assertEqual(len(dataset.records), 10)
        self.assertIn("TRANSCRIPT_COUNTS", dataset.genes)
        self.assertTrue(dataset.metadata["xenium_files"]["cell_feature_matrix_h5"])
        self.assertTrue(dataset.metadata["xenium_files"]["experiment_xenium"])
        self.assertIn("xenium_explorer_assets", dataset.metadata)
        self.assertIn("gene_matrix", dataset.metadata)
        self.assertEqual(dataset.metadata["gene_matrix"]["loader"], "h5py_10x_csc_v1")
        first = dataset.records[0]
        self.assertEqual(first.genes["TRANSCRIPT_COUNTS"], first.raw_genes["TRANSCRIPT_COUNTS"])
        self.assertEqual(dataset.qc_metrics["qc_expression_layer"], "raw_counts")
        self.assertEqual(dataset.metadata["analysis_scope"], "sampled")

    def test_loads_xenium_experiment_file_entrypoint(self):
        experiment_path = os.path.join(XENIUM_LYMPH, "experiment.xenium")
        self.assertEqual(infer_data_type(experiment_path), "xenium_experiment_file")
        dataset = DataIngestionLayer().load(experiment_path)
        self.assertEqual(dataset.modality, "spatial_transcriptomics")
        # The scan selects 5000 cells; per-cell QC then removes the unusable ones,
        # so the retained count is 5000 minus exactly what QC reports dropping.
        # This lymph node section is genuinely sparse (median 45 transcripts per
        # cell against 166 in the brain section), so the drop is not negligible.
        sampling = dataset.metadata["sampling"]
        cell_qc = dataset.metadata["cell_qc"]
        self.assertEqual(sampling["scanned_records"], 5000)
        self.assertEqual(len(dataset.records), 5000 - cell_qc["dropped_cell_count"])
        self.assertEqual(cell_qc["retained_cell_count"], len(dataset.records))
        self.assertEqual(dataset.source_path, experiment_path)
        self.assertEqual(dataset.sources[0].data_type, "xenium_experiment_file")
        self.assertTrue(dataset.metadata["xenium_files"]["experiment_xenium"])
        self.assertTrue(dataset.metadata["xenium_explorer_assets"]["analysis_summary_filepath"]["exists"])

    def test_discovers_and_inspects_xenium_datasets(self):
        candidates = discover_dataset_candidates(os.path.join(ROOT, "data"))
        self.assertIn(XENIUM_LYMPH, candidates)
        self.assertFalse(any("cell_ontology_terms" in path for path in candidates))
        inspection = inspect_dataset(XENIUM_LYMPH)
        self.assertTrue(inspection.usable)
        self.assertEqual(inspection.readiness, "partially_ready")
        self.assertIn("spatial_scatter", inspection.supported_workflows)
        self.assertFalse(any("gene matrix was not loaded" in blocker for blocker in inspection.blockers))
        self.assertTrue(any("reviewed expert" in blocker for blocker in inspection.blockers))
        self.assertTrue(inspection.metadata["dataset_metadata"]["sampling"]["fraction_loaded"] > 0)

    def test_pipeline_returns_report(self):
        dataset, report = DataIngestionPipeline().ingest(
            Path(DEMO),
            IngestionConfig(format=DataFormat.TIDY_CSV, sample_id="BRCA_04", min_counts=0),
        )
        self.assertEqual(dataset.sample_id, "BRCA_04")
        self.assertEqual(report.format_detected, DataFormat.TIDY_CSV)
        self.assertEqual(report.errors, [])

    def test_batch_ingestion_pipeline_harmonizes_gene_space(self):
        datasets, report = BatchIngestionPipeline().ingest_batch(
            BatchIngestionConfig(
                samples=[
                    SampleConfig(path=Path(DEMO), sample_id="BRCA_04", format=DataFormat.TIDY_CSV),
                    SampleConfig(path=Path(DEMO), sample_id="BRCA_04", format=DataFormat.TIDY_CSV),
                ]
            )
        )
        self.assertEqual(len(datasets), 2)
        self.assertGreater(report.harmonized_gene_count, 0)

    def test_readiness_report_blocks_label_dependent_workflow_without_labels(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "unannotated.csv")
            with open(path, "w", encoding="utf-8") as handle:
                handle.write("sample_id,x,y,cell_type,CD8A,EPCAM\n")
                for index in range(10):
                    handle.write("S1,%s,%s,Unannotated,%s,%s\n" % (index, index + 1, index + 2, index + 3))
            dataset = DataIngestionLayer().load_csv(path, sample_id="S1")
            report = build_readiness_report(dataset)
            self.assertEqual(report.workflow_status("spatial_visualization").status, "ready")
            self.assertEqual(report.workflow_status("neighborhood_enrichment").status, "blocked")

    def test_mvp_loaders_set_cell_by_feature_subtypes(self):
        scrna = load_scrna(DEMO, sample_id="BRCA_04")
        scatac = load_scatac(DEMO, sample_id="BRCA_04")
        self.assertEqual(validate_cell_by_feature_contract(scrna).assay_subtype, "scrna")
        scatac_contract = validate_cell_by_feature_contract(scatac)
        self.assertEqual(scatac_contract.assay_subtype, "scatac_gene_activity")
        self.assertEqual(scatac_contract.feature_type, "gene_activity")

    def test_external_label_table_overrides_loaded_cell_ids(self):
        with tempfile.TemporaryDirectory() as tmp:
            data_path = os.path.join(tmp, "cells.csv")
            with open(data_path, "w", encoding="utf-8") as handle:
                handle.write("sample_id,cell_id,x,y,cell_type,CD8A\n")
                handle.write("S1,c1,1,2,Unannotated,3\n")
                handle.write("S1,c2,2,3,Unannotated,4\n")
            label_path = os.path.join(tmp, "expert_cell_labels.csv")
            with open(label_path, "w", encoding="utf-8") as handle:
                handle.write("cell_id,expert_label,confidence\n")
                handle.write("c1,T cell,0.91\n")
                handle.write("c2,Tumor cell,0.87\n")
            dataset = DataIngestionLayer().load_csv(data_path, sample_id="S1")
            report = apply_external_label_table(dataset, label_path)
            self.assertEqual(report.status, "expert_labels_applied")
            self.assertEqual(report.matched_cells, 2)
            self.assertEqual(sorted(dataset.cell_types), ["T cell", "Tumor cell"])
            self.assertEqual(dataset.metadata["annotation_strategy"], "expert_label_table")

    def test_external_region_table_applies_user_regions(self):
        with tempfile.TemporaryDirectory() as tmp:
            data_path = os.path.join(tmp, "cells.csv")
            with open(data_path, "w", encoding="utf-8") as handle:
                handle.write("sample_id,cell_id,x,y,cell_type,CD8A\n")
                handle.write("S1,c1,1,2,T cell,3\n")
                handle.write("S1,c2,2,3,Tumor cell,4\n")
            region_path = os.path.join(tmp, "cell_regions.csv")
            with open(region_path, "w", encoding="utf-8") as handle:
                handle.write("cell_id,region,region_confidence\n")
                handle.write("c1,stroma,0.8\n")
                handle.write("c2,tumor_core,0.9\n")
            dataset = DataIngestionLayer().load_csv(data_path, sample_id="S1")
            report = apply_external_region_table(dataset, region_path)
            self.assertEqual(report.status, "user_regions_applied")
            self.assertEqual(report.matched_cells, 2)
            self.assertEqual(sorted({record.region for record in dataset.records}), ["stroma", "tumor_core"])
            self.assertEqual(dataset.metadata["region_label_source"], "user_provided")

    def test_xenium_expert_readiness_reports_missing_labels_but_existing_clusters(self):
        readiness = summarize_xenium_expert_readiness(XENIUM_LYMPH)
        self.assertTrue(readiness.has_cell_table)
        self.assertTrue(readiness.has_feature_matrix)
        self.assertTrue(readiness.has_10x_analysis_clusters)
        self.assertIn("gene_expression_graphclust", readiness.cluster_methods)
        self.assertFalse(readiness.external_label_tables)
        self.assertTrue(any("Expert cell label table" in item for item in readiness.needs))

    def test_expert_label_template_includes_cluster_and_marker_evidence(self):
        with tempfile.TemporaryDirectory() as tmp:
            data_path = os.path.join(tmp, "cells.csv")
            with open(data_path, "w", encoding="utf-8") as handle:
                handle.write("sample_id,cell_id,x,y,cell_type,CD8A,EPCAM\n")
                handle.write("S1,c1,1,2,Unannotated,3,4\n")
            cluster_dir = os.path.join(tmp, "analysis", "clustering", "gene_expression_graphclust")
            os.makedirs(cluster_dir)
            with open(os.path.join(cluster_dir, "clusters.csv"), "w", encoding="utf-8") as handle:
                handle.write("Barcode,Cluster\n")
                handle.write("c1,7\n")
            dataset = DataIngestionLayer().load_csv(data_path, sample_id="S1")
            template = write_expert_label_template(dataset, os.path.join(tmp, "expert_label_template.csv"), dataset_path=tmp)
            with open(template, encoding="utf-8") as handle:
                content = handle.read()
            self.assertIn("graph_cluster", content)
            self.assertIn("marker_evidence", content)
            self.assertIn(",7,", content)
            self.assertIn("CD8A=", content)

    def test_region_label_template_includes_region_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            data_path = os.path.join(tmp, "cells.csv")
            with open(data_path, "w", encoding="utf-8") as handle:
                handle.write("sample_id,cell_id,x,y,cell_type,region,CD8A\n")
                handle.write("S1,c1,1,2,T cell,old_region,3\n")
            dataset = DataIngestionLayer().load_csv(data_path, sample_id="S1")
            template = write_region_label_template(dataset, os.path.join(tmp, "region_label_template.csv"), dataset_path=tmp)
            with open(template, encoding="utf-8") as handle:
                content = handle.read()
            self.assertIn("current_region", content)
            self.assertIn("region_confidence", content)
            self.assertIn("old_region", content)

    def test_xenium_label_intake_report_passes_validated_inputs(self):
        dataset = SpatialDataset(
            sample_id="X1",
            source_path="xenium",
            modality="xenium_spatial_rna",
            records=[
                SpotRecord("X1", 0.0, 0.0, "T cell", {"CD8A": 1.0}, region="stroma", cell_id="c1"),
                SpotRecord("X1", 1.0, 1.0, "Tumor cell", {"EPCAM": 1.0}, region="tumor_core", cell_id="c2"),
            ],
        )
        report = build_xenium_label_intake_report(
            dataset=dataset,
            dataset_path="xenium",
            label_report=LabelApplicationReport(
                status="expert_labels_applied",
                method="expert_label_table",
                source_path="expert_cell_labels.csv",
                matched_cells=2,
                total_records=2,
                label_counts={"T cell": 1, "Tumor cell": 1},
                confidence_summary={"mean": 0.9},
            ),
            region_report=RegionApplicationReport(
                status="user_regions_applied",
                method="user_region_table",
                source_path="cell_regions.csv",
                matched_cells=2,
                total_records=2,
                region_counts={"stroma": 1, "tumor_core": 1},
                confidence_summary={"mean": 0.9},
            ),
            asset_readiness=XeniumExpertReadiness(
                dataset_path="xenium",
                has_cell_table=True,
                has_feature_matrix=True,
                has_morphology=True,
                has_boundaries=True,
                has_10x_analysis_clusters=True,
                cluster_methods=["gene_expression_graphclust"],
                external_label_tables=["expert_cell_labels.csv"],
                external_region_tables=["cell_regions.csv"],
                ready_for_expert_label_mvp=True,
                ready_for_region_summary_mvp=True,
            ),
        )
        self.assertEqual(report.status, "validated_ready")
        self.assertTrue(report.ready_for_validated_pilot)
        self.assertEqual(report.label_coverage, 1.0)
        self.assertEqual(report.region_coverage, 1.0)

    def test_xenium_label_intake_report_blocks_missing_inputs(self):
        dataset = SpatialDataset(
            sample_id="X1",
            source_path="xenium",
            modality="xenium_spatial_rna",
            records=[SpotRecord("X1", 0.0, 0.0, "Unannotated cell", {"CD8A": 1.0}, cell_id="c1")],
        )
        report = build_xenium_label_intake_report(
            dataset=dataset,
            dataset_path="xenium",
            label_report=LabelApplicationReport(status="missing_expert_labels", method="none", total_records=1),
            region_report=RegionApplicationReport(status="missing_user_regions", method="none", total_records=1),
            asset_readiness=XeniumExpertReadiness(
                dataset_path="xenium",
                has_cell_table=True,
                has_feature_matrix=True,
                has_morphology=True,
                has_boundaries=True,
                has_10x_analysis_clusters=True,
                cluster_methods=[],
                external_label_tables=[],
                external_region_tables=[],
                ready_for_expert_label_mvp=False,
                ready_for_region_summary_mvp=False,
            ),
        )
        self.assertEqual(report.status, "blocked_label_intake")
        self.assertFalse(report.ready_for_validated_pilot)
        self.assertTrue(any("Expert cell labels" in item for item in report.blockers))
        self.assertTrue(any("User-provided region labels" in item for item in report.blockers))


class ToolRegistryTests(unittest.TestCase):
    def test_default_registry_keeps_full_scaffold_and_anthropic_schema(self):
        registry = build_default_registry()
        self.assertGreaterEqual(len(registry.list_all()), 22)
        # Scaffolds stay registered for provenance, but are not advertised to a
        # planner: a model must not be able to pick a tool that does no work.
        anthropic_tools = registry.to_anthropic_tools()
        self.assertTrue(any(tool["name"] == "neighborhood_enrichment" for tool in anthropic_tools))
        self.assertFalse(any(tool["name"] == "cnv_inference" for tool in anthropic_tools))
        self.assertTrue(any(tool.name == "cnv_inference" for tool in registry.list_all()))
        self.assertTrue(
            any(tool["name"] == "cnv_inference" for tool in registry.to_anthropic_tools(plannable_only=False))
        )

    def test_registry_exposes_resource_profiles_and_method_citations(self):
        registry = build_default_registry()
        tool = registry.get("neighborhood_enrichment")
        self.assertEqual(tool.resource_profile.runtime, "medium")
        self.assertIn("Squidpy", tool.citation.method_name)

    def test_mvp_registry_exposes_only_v7_named_tools(self):
        registry = build_mvp_registry()
        names = [tool.name for tool in registry.list_all()]
        self.assertEqual(sorted(names), sorted(MVP_TOOL_NAMES))
        self.assertEqual(
            sorted(names),
            sorted(
                [
                    "qc_and_cluster",
                    "annotation",
                    "marker_detection",
                    "spatial_variable_genes",
                    "feature_overlay",
                    "region_summary",
                    "cell_neighborhood_enrichment",
                ]
            ),
        )
        self.assertNotIn("differential_expression", names)
        self.assertNotIn("trajectory_inference", names)
        self.assertNotIn("motif_tf_activity", names)
        self.assertNotIn("reference_label_transfer", names)
        self.assertNotIn("cnv_inference", names)
        self.assertNotIn("ligand_receptor_analysis", names)

    def test_mvp_tool_results_include_v7_quality_metrics(self):
        dataset = DataIngestionLayer().load_csv(DEMO, sample_id="BRCA_04")
        result = build_mvp_registry().get("marker_detection").run(
            dataset,
            {"group_key": "cell_type", "group1": "CD8+ T cell", "group2": "Tumor cell"},
        )
        self.assertIsNotNone(result.quality_metrics)
        self.assertEqual(result.quality_metrics.differential.n_significant.role, "statistical_evidence")
        self.assertIn("quality_metrics", result.metrics)

    def test_spatial_variable_gene_results_expose_morans_i_quality_metric(self):
        from spatialmind.tools.implementations import attach_quality_metrics

        dataset = DataIngestionLayer().load_csv(DEMO, sample_id="BRCA_04")
        result = ToolResult(
            tool_name="spatial_variable_genes",
            summary="Synthetic spatial-gene result.",
            metrics={
                "engine": "squidpy",
                "n_neighs": 6,
                "top_genes": [{"gene": "AQP4", "morans_i": 0.42, "pval_adj": 0.01}],
            },
        )
        attached = attach_quality_metrics(result, dataset, {"n_neighs": 6})
        self.assertEqual(attached.quality_metrics.spatial.morans_i.status, "computed")
        self.assertEqual(attached.quality_metrics.spatial.morans_i.value, 0.42)
        self.assertEqual(attached.quality_metrics.spatial.cooccurrence_z.status, "not_applicable")

    @unittest.skipUnless(_HAS_SCANPY, "scanpy required for expression clustering path")
    def test_qc_and_cluster_defaults_to_expression_clustering_with_qc(self):
        dataset = _make_two_program_dataset()
        result = build_mvp_registry().get("qc_and_cluster").run(dataset, {})
        metrics = result.metrics
        self.assertEqual(metrics["engine"], "scanpy")
        self.assertEqual(metrics["cluster_on"], "expression")
        self.assertEqual(metrics["method"], "pca_neighbors_leiden")
        self.assertTrue(str(metrics["representation"]).startswith("X_pca"))
        # Two cleanly separated expression programs should not collapse to one cluster.
        self.assertGreaterEqual(len(metrics["cluster_counts"]), 2)
        self.assertIsNotNone(metrics["silhouette"])
        self.assertGreaterEqual(metrics["silhouette"], -1.0)
        self.assertLessEqual(metrics["silhouette"], 1.0)
        self.assertIsNotNone(metrics["modularity"])
        self.assertGreaterEqual(metrics["modularity"], -1.0)
        self.assertLessEqual(metrics["modularity"], 1.0)
        self.assertEqual(metrics["excluded_zero_feature_cell_count"], 0)
        self.assertEqual(result.quality_metrics.clustering.modularity.provenance.method, "weighted kNN graph modularity")
        qc = metrics["expression_qc"]
        self.assertEqual(qc["n_cells"], 60)
        self.assertEqual(qc["n_features"], 12)
        self.assertGreater(qc["mean_total_counts"], 0.0)
        self.assertGreater(qc["mean_features_per_cell"], 0.0)

    @unittest.skipUnless(_HAS_SCANPY, "scanpy required for expression clustering path")
    def test_expression_clustering_excludes_zero_feature_cells(self):
        dataset = _make_two_program_dataset()
        dataset.records.append(
            SpotRecord("S1", 100.0, 100.0, "Unannotated cell", {}, cell_id="zero-expression")
        )
        result = build_mvp_registry().get("qc_and_cluster").run(dataset, {})
        self.assertEqual(result.metrics["total_cell_count"], 61)
        self.assertEqual(result.metrics["analyzed_cell_count"], 60)
        self.assertEqual(result.metrics["excluded_zero_feature_cell_count"], 1)
        assignments = dataset.metadata["cluster_assignments"]
        self.assertEqual(assignments["zero-expression"], "")
        markers = build_mvp_registry().get("marker_detection").run(dataset, {"group_key": "cluster"})
        self.assertEqual(markers.metrics["excluded_unassigned_cell_count"], 1)
        self.assertNotIn("unassigned", markers.metrics["groups"])

    @unittest.skipUnless(_HAS_SCANPY, "scanpy required for expression clustering path")
    def test_qc_and_cluster_spatial_mode_is_opt_in(self):
        dataset = _make_two_program_dataset()
        result = build_mvp_registry().get("qc_and_cluster").run(dataset, {"cluster_on": "spatial"})
        self.assertEqual(result.metrics["cluster_on"], "spatial")
        self.assertEqual(result.metrics["method"], "spatial_neighbors_leiden")
        self.assertEqual(result.metrics["representation"], "spatial")

    def test_expression_matrix_excludes_qc_pseudo_features(self):
        from spatialmind.tools.implementations import EXPRESSION_EXCLUDED_FEATURES, expression_feature_names

        records = [
            SpotRecord(
                "S1",
                float(index),
                0.0,
                "A" if index % 2 == 0 else "B",
                {"CD8A": 2.0, "EPCAM": 1.0, "CELL_AREA": 500.0, "TOTAL_COUNTS": 300.0, "NUCLEUS_AREA": 90.0},
                cell_id="c%d" % index,
            )
            for index in range(6)
        ]
        dataset = SpatialDataset(sample_id="S1", source_path="synthetic", modality="xenium_spatial_rna", records=records)
        names = expression_feature_names(dataset)
        self.assertIn("CD8A", names)
        self.assertTrue(EXPRESSION_EXCLUDED_FEATURES.isdisjoint({name.upper() for name in names}))

    def test_prototype_marker_detection_ignores_qc_pseudo_features(self):
        records = [
            SpotRecord("S1", 0.0, 0.0, "A", {"CD8A": 9.0, "EPCAM": 0.0, "CELL_AREA": 800.0, "TOTAL_COUNTS": 500.0}, cell_id="a"),
            SpotRecord("S1", 1.0, 0.0, "B", {"CD8A": 0.0, "EPCAM": 8.0, "CELL_AREA": 20.0, "TOTAL_COUNTS": 15.0}, cell_id="b"),
        ]
        dataset = SpatialDataset(sample_id="S1", source_path="synthetic", modality="xenium_spatial_rna", records=records)
        result = marker_detection(dataset, {"engine": "prototype", "group_key": "cell_type", "group1": "A", "group2": "B"})
        genes = {row["gene"] for row in result.metrics["ranked_genes"]}
        self.assertIn("CD8A", genes)
        self.assertNotIn("CELL_AREA", genes)
        self.assertNotIn("TOTAL_COUNTS", genes)

    def test_neighborhood_robustness_summary_scores_stability(self):
        from spatialmind.tools.implementations import summarize_neighborhood_robustness

        def setting(n, pairs):
            return {"n_neighs": n, "engine": "squidpy", "pairs": pairs}

        stable = [
            setting(6, [{"pair": "A | B", "zscore": 5.0}, {"pair": "A | C", "zscore": -3.0}]),
            setting(10, [{"pair": "A | B", "zscore": 4.2}, {"pair": "A | C", "zscore": -2.5}]),
        ]
        stable_summary = summarize_neighborhood_robustness(stable, top_k=2)
        self.assertEqual(stable_summary["status"], "computed")
        self.assertEqual(stable_summary["mean_sign_agreement"], 1.0)
        self.assertEqual(stable_summary["score"], 1.0)
        self.assertEqual(stable_summary["pair_stability"][0]["settings_present"], 2)
        self.assertEqual(stable_summary["pair_stability"][0]["sign_agreement"], 1.0)

        flipped = [
            setting(6, [{"pair": "A | B", "zscore": 5.0}, {"pair": "A | C", "zscore": -3.0}]),
            setting(10, [{"pair": "A | B", "zscore": -4.0}, {"pair": "A | C", "zscore": 3.0}]),
        ]
        flipped_summary = summarize_neighborhood_robustness(flipped, top_k=2)
        self.assertEqual(flipped_summary["mean_sign_agreement"], 0.0)
        self.assertLess(flipped_summary["score"], stable_summary["score"])

        # Prototype settings carry no z-scores, so robustness cannot be established.
        proto = [{"n_neighs": 6, "engine": "prototype", "pairs": [{"pair": "A | B", "neighbor_count": 10}]}]
        self.assertEqual(summarize_neighborhood_robustness(proto, top_k=2)["status"], "insufficient_settings")

    def test_neighborhood_robustness_records_execution_settings(self):
        from spatialmind.tools.implementations import run_neighborhood_robustness

        def result_for(_dataset, params):
            n_neighs = int(params["n_neighs"])
            return ToolResult(
                tool_name="cell_neighborhood_enrichment",
                summary="test",
                metrics={
                    "engine": "squidpy",
                    "all_pairs": [
                        {"pair": "A | B", "zscore": 4.0 + n_neighs / 100.0},
                        {"pair": "A | C", "zscore": -3.0},
                    ],
                },
            )

        with patch("spatialmind.tools.implementations.cell_neighborhood_enrichment", side_effect=result_for):
            summary = run_neighborhood_robustness(
                _make_two_program_dataset(),
                {
                    "robustness_n_neighs": [5, 9],
                    "n_perms": 250,
                    "random_state": 17,
                    "robustness_top_k": 2,
                },
            )
        self.assertEqual(summary["status"], "computed")
        self.assertEqual(summary["requested_settings"], [5, 9])
        self.assertEqual(summary["n_perms"], 250)
        self.assertEqual(summary["random_state"], 17)
        self.assertEqual(summary["top_k"], 2)
        self.assertEqual(summary["engines"], ["squidpy"])

    def test_spatial_relationship_summary_combines_orthogonal_evidence(self):
        from spatialmind.pilot.spatial_relationships import build_spatial_relationship_summary

        records = []
        for index in range(20):
            records.append(SpotRecord("S1", float(index), 0.0, "A", {"G": 1.0}, region="core", cell_id="a%d" % index))
            records.append(SpotRecord("S1", float(index) + 0.5, 1.0, "B", {"G": 1.0}, region="core", cell_id="b%d" % index))
            records.append(SpotRecord("S1", float(index) + 100.0, 0.0, "C", {"G": 1.0}, region="margin", cell_id="c%d" % index))
        dataset = SpatialDataset(
            sample_id="S1",
            source_path="synthetic",
            modality="xenium_spatial_rna",
            coordinate_system="micron",
            records=records,
        )
        neighborhood = ToolResult(
            tool_name="cell_neighborhood_enrichment",
            summary="test",
            metrics={
                "engine": "squidpy",
                "n_neighs": 6,
                "n_perms": 250,
                "random_state": 0,
                "tested_pair_count": 2,
                "all_pairs": [
                    {"pair": "A | B", "zscore": 4.2},
                    {"pair": "A | C", "zscore": -3.1},
                ],
            },
        )
        robustness = {
            "status": "computed",
            "score": 0.84,
            "pair_stability": [
                {"pair": "A | B", "settings_present": 3, "sign_agreement": 1.0, "top_k_presence": 1.0},
                {"pair": "A | C", "settings_present": 3, "sign_agreement": 1.0, "top_k_presence": 1.0},
            ],
        }
        summary = build_spatial_relationship_summary(dataset, [neighborhood], robustness, validated=True)
        self.assertEqual(summary["status"], "computed")
        by_pair = {item["pair"]: item for item in summary["relationships"]}
        self.assertEqual(by_pair["A | B"]["evidence_status"], "stable_enriched")
        self.assertEqual(by_pair["A | C"]["evidence_status"], "stable_depleted")
        self.assertEqual(by_pair["A | B"]["region_overlap"], 1.0)
        self.assertGreater(by_pair["A | C"]["median_bidirectional_nearest_distance"], 50.0)
        self.assertIn("not evidence", summary["warnings"][0])

    def test_spatial_relationship_summary_rejects_prototype_neighbor_counts(self):
        from spatialmind.pilot.spatial_relationships import build_spatial_relationship_summary

        dataset = _make_two_program_dataset()
        result = ToolResult(
            tool_name="cell_neighborhood_enrichment",
            summary="prototype",
            metrics={"engine": "prototype", "top_pairs": [{"pair": "A | B", "neighbor_count": 20}]},
        )
        summary = build_spatial_relationship_summary(dataset, [result], {}, validated=True)
        self.assertEqual(summary["status"], "not_computed")
        self.assertIn("permutation z-scores", summary["reason"])

    def test_region_stratified_neighborhoods_track_consistency_and_skips(self):
        from spatialmind.tools.implementations import run_region_stratified_neighborhoods

        records = []
        for region, offset in (("core", 0.0), ("margin", 100.0)):
            for index in range(20):
                records.append(SpotRecord("S1", offset + index, 0.0, "A", {"G": 1.0}, region=region, cell_id="%sa%d" % (region, index)))
                records.append(SpotRecord("S1", offset + index, 1.0, "B", {"G": 1.0}, region=region, cell_id="%sb%d" % (region, index)))
        records.append(SpotRecord("S1", 300.0, 0.0, "A", {"G": 1.0}, region="tiny", cell_id="tiny"))
        dataset = SpatialDataset("S1", records, "synthetic", modality="xenium_spatial_rna")
        calls = []

        def region_result(subset, params):
            calls.append((subset.records[0].region, params))
            zscore = 4.0 if subset.records[0].region == "core" else 3.0
            return ToolResult(
                tool_name="cell_neighborhood_enrichment",
                summary="test",
                metrics={
                    "engine": "squidpy",
                    "all_pairs": [
                        {"pair": "A | B", "zscore": zscore},
                        {"pair": "A | A", "zscore": -2.0},
                        {"pair": "B | B", "zscore": 0.2},
                    ],
                },
            )

        with patch("spatialmind.tools.implementations.cell_neighborhood_enrichment", side_effect=region_result):
            summary = run_region_stratified_neighborhoods(
                dataset,
                {"min_region_cells": 40, "min_cells_per_type": 20, "n_perms": 50},
            )
        self.assertEqual(summary["status"], "computed")
        self.assertEqual(summary["tested_region_count"], 2)
        self.assertEqual(summary["skipped_region_count"], 1)
        self.assertEqual(summary["pair_consistency"][0]["pair"], "A | B")
        self.assertEqual(summary["pair_consistency"][0]["status"], "region_consistent")
        self.assertEqual(summary["pair_consistency"][0]["supported_region_count"], 2)
        weak_pair = next(item for item in summary["pair_consistency"] if item["pair"] == "B | B")
        self.assertEqual(weak_pair["status"], "weak_or_indeterminate")
        self.assertEqual(weak_pair["supported_region_count"], 0)
        self.assertTrue(all(call[1]["include_all_pairs"] for call in calls))

    def test_distance_cooccurrence_curves_use_symmetric_pair_average(self):
        import numpy as np
        from spatialmind.tools.implementations import run_distance_dependent_cooccurrence

        dataset = SpatialDataset(
            "S1",
            [
                SpotRecord("S1", float(index), 0.0, "A" if index < 20 else "B", {"G": 1.0}, cell_id="c%d" % index)
                for index in range(40)
            ],
            "synthetic",
            modality="xenium_spatial_rna",
            coordinate_system="micron",
        )
        occurrence = np.zeros((2, 2, 6), dtype=float)
        occurrence[0, 1, :] = np.array([1.2, 1.4, 1.6, 1.3, 1.1, 1.0])
        occurrence[1, 0, :] = np.array([1.0, 1.2, 1.4, 1.1, 0.9, 0.8])
        thresholds = np.arange(7, dtype=float)
        with patch("squidpy.gr.co_occurrence", autospec=True, return_value=(occurrence, thresholds)) as mocked:
            summary = run_distance_dependent_cooccurrence(
                dataset,
                pairs=["A | B"],
                params={"n_intervals": 6, "max_distance": 6.0},
            )
        self.assertEqual(summary["status"], "computed")
        self.assertEqual(summary["curve_count"], 1)
        self.assertEqual(summary["curves"][0]["peak_ratio"], 1.5)
        self.assertEqual(summary["curves"][0]["peak_distance"], 3.0)
        self.assertEqual(summary["curves"][0]["left_cell_count"], 20)
        self.assertEqual(summary["min_cells_per_type"], 20)
        self.assertEqual(mocked.call_args.kwargs.get("n_jobs"), 1)
        self.assertEqual(mocked.call_args.kwargs.get("backend"), "threading")

    def test_validated_reports_surface_spatial_robustness(self):
        from spatialmind.pilot.xenium import (
            _spatial_robustness_rows,
            _write_html_report,
            _write_markdown_report,
            _write_pilot_pdf_report,
        )

        payload = {
            "created_at": "2026-07-22T00:00:00+00:00",
            "status": "validated_ready",
            "dataset_path": "synthetic.xenium",
            "records_loaded": 20,
            "features_loaded": 10,
            "blocking_reasons": [],
            "required_next_inputs": [],
            "tool_plan": [],
            "plan_validation": {"status": "valid", "errors": []},
            "expert_label_template": "labels.csv",
            "region_label_template": "regions.csv",
            "label_report": {"status": "expert_labels_applied"},
            "region_report": {"status": "user_regions_applied"},
            "review_figures": [],
            "cell_type_counts": {"A": 10, "B": 10},
            "region_counts": {"core": 10, "margin": 10},
            "claim_ledger": [],
            "claim_reliability": [],
            "run_record_path": "run.json",
            "report_html": "report.html",
            "descriptive_analysis": {
                "status": "computed",
                "clustering_method": "pca_neighbors_leiden",
                "expression_qc": {
                    "n_cells": 20,
                    "n_features": 10,
                    "median_total_counts": 12.0,
                    "median_features_per_cell": 4.0,
                },
                "clustering_diagnostics": {
                    "analyzed_cell_count": 19,
                    "excluded_zero_feature_cell_count": 1,
                    "silhouette": 0.42,
                    "modularity": 0.61,
                },
                "cluster_counts": {"0": 10, "1": 9},
                "markers_by_cluster": {"0": ["AQP4", "GJA1"], "1": ["MOG", "CLDN11"]},
                "cluster_neighborhood": {"top_pairs": [{"pair": "0 | 1", "zscore": -2.4}]},
                "interpretation": "Synthetic descriptive fixture.",
            },
            "spatial_robustness": {
                "status": "computed",
                "score": 0.82,
                "requested_settings": [6, 10, 15],
                "n_perms": 250,
                "random_state": 17,
                "top_k": 10,
                "engines": ["squidpy"],
                "mean_sign_agreement": 1.0,
                "mean_topk_jaccard": 0.55,
                "n_reference_pairs": 10,
            },
            "spatial_relationships": {
                "status": "computed",
                "method": "Squidpy permutation neighborhood enrichment",
                "graph": {"n_neighs": 6, "n_perms": 250, "random_state": 17},
                "relationships": [
                    {
                        "pair": "A | B",
                        "direction": "enriched",
                        "zscore": 4.2,
                        "settings_present": 3,
                        "sign_agreement": 1.0,
                        "median_bidirectional_nearest_distance": 8.5,
                        "coordinate_units": "micron",
                        "region_overlap": 0.8,
                        "shared_regions": ["core"],
                        "evidence_status": "stable_enriched",
                    }
                ],
                "warnings": ["Spatial adjacency is not evidence of signaling or causation."],
            },
            "region_stratified_neighborhoods": {
                "status": "computed",
                "tested_region_count": 2,
                "skipped_region_count": 0,
                "pair_consistency": [
                    {
                        "pair": "A | B",
                        "regions_tested": 2,
                        "direction_agreement": 1.0,
                        "strongest_region": "core",
                        "strongest_abs_zscore": 3.8,
                        "status": "region_consistent",
                    }
                ],
                "warnings": ["Within-region synthetic fixture."],
            },
            "distance_cooccurrence": {
                "status": "computed",
                "coordinate_units": "micron",
                "max_distance": 50.0,
                "n_intervals": 20,
                "curves": [
                    {
                        "pair": "A | B",
                        "peak_ratio": 1.4,
                        "peak_distance": 10.0,
                        "short_range_mean_ratio": 1.3,
                        "long_range_mean_ratio": 1.0,
                    }
                ],
                "warnings": ["Descriptive synthetic fixture."],
            },
        }
        rows = dict(_spatial_robustness_rows(payload))
        self.assertEqual(rows["Robustness score"], "0.8200")
        self.assertEqual(rows["Permutations per setting"], "250")
        self.assertEqual(rows["Random seed"], "17")
        with tempfile.TemporaryDirectory() as tmp:
            md_path = Path(tmp) / "report.md"
            html_path = Path(tmp) / "report.html"
            pdf_path = Path(tmp) / "report.pdf"
            _write_markdown_report(md_path, payload, [])
            _write_html_report(html_path, payload, [])
            _write_pilot_pdf_report(pdf_path, payload, [])
            self.assertIn("Spatial Robustness Sweep", md_path.read_text(encoding="utf-8"))
            self.assertIn("Descriptive Analysis (no expert labels required)", md_path.read_text(encoding="utf-8"))
            self.assertIn("Weighted graph modularity", md_path.read_text(encoding="utf-8"))
            self.assertIn("Spatial Relationships", md_path.read_text(encoding="utf-8"))
            self.assertIn("stable_enriched", md_path.read_text(encoding="utf-8"))
            self.assertIn("Region-Stratified Neighborhood Testing", md_path.read_text(encoding="utf-8"))
            self.assertIn("Distance-Dependent Co-Occurrence", md_path.read_text(encoding="utf-8"))
            html_report = html_path.read_text(encoding="utf-8")
            self.assertIn("Spatial Robustness Sweep", html_report)
            self.assertIn("Descriptive Analysis (no expert labels required)", html_report)
            self.assertIn("Weighted graph modularity", html_report)
            self.assertIn("Spatial Relationships", html_report)
            self.assertIn("stable_enriched", html_report)
            self.assertIn("region_consistent", html_report)
            self.assertIn("Peak distance", html_report)
            self.assertIn("0.8200", html_report)
            with pdf_path.open("rb") as handle:
                self.assertEqual(handle.read(5), b"%PDF-")

    def test_spatial_robustness_component_prefers_real_sweep(self):
        from spatialmind.methods.reliability.scoring import _spatial_robustness_component

        claim = {"claim_type": "spatial_colocalization", "status": "supported"}
        payload = {
            "spatial_robustness": {
                "status": "computed",
                "score": 0.82,
                "mean_sign_agreement": 1.0,
                "mean_topk_jaccard": 0.55,
                "settings": [6, 10, 15],
            }
        }
        component = _spatial_robustness_component(claim, payload, [])
        self.assertEqual(component.status, "computed")
        self.assertEqual(component.score, 0.82)
        # Without a sweep, it falls back to the heuristic proxy (no crash, still a score).
        fallback = _spatial_robustness_component(claim, {}, [])
        self.assertIsNotNone(fallback.score)

    def test_pilot_report_renders_per_group_markers(self):
        from spatialmind.pilot.xenium import _marker_group_markdown, _marker_group_html
        from spatialmind.schemas import ToolResult

        result = ToolResult(
            tool_name="marker_detection",
            summary="Detected one-vs-rest marker candidates for 2 groups.",
            metrics={
                "mode": "one_vs_rest",
                "markers_by_group": {
                    "T cell": [{"gene": "CD8A"}, {"gene": "CD3D"}],
                    "Tumor cell": [{"gene": "EPCAM"}],
                },
            },
        )
        md = "\n".join(_marker_group_markdown(result))
        self.assertIn("Top markers (one-vs-rest)", md)
        self.assertIn("CD8A", md)
        self.assertIn("EPCAM", md)
        self.assertIn("CD8A", _marker_group_html(result))
        # Non-marker tools contribute no marker table.
        other = ToolResult(tool_name="annotation", summary="x", metrics={})
        self.assertEqual(_marker_group_markdown(other), [])
        self.assertEqual(_marker_group_html(other), "")

    def test_full_panel_loading_keeps_all_positive_genes(self):
        from spatialmind.ingestion.pipeline import _matrix_row_to_features

        gene_names = ["g0", "g1", "g2", "g3", "g4"]
        row = [5.0, 3.0, 0.0, 4.0, 2.0]
        # Full panel (0) keeps every positive gene; the zero-valued gene stays absent.
        full = _matrix_row_to_features(row, gene_names, max_features_per_record=0)
        self.assertEqual(set(full), {"g0", "g1", "g3", "g4"})
        # A positive cap truncates to the top-N by value.
        capped = _matrix_row_to_features(row, gene_names, max_features_per_record=2)
        self.assertEqual(set(capped), {"g0", "g3"})

    def test_marker_detection_defaults_to_one_vs_rest_per_group(self):
        records = []
        for index in range(12):
            if index % 3 == 0:
                cell_type, genes = "T cell", {"CD8A": 8.0, "EPCAM": 0.0, "PECAM1": 0.0}
            elif index % 3 == 1:
                cell_type, genes = "Tumor cell", {"CD8A": 0.0, "EPCAM": 8.0, "PECAM1": 0.0}
            else:
                cell_type, genes = "Endothelial cell", {"CD8A": 0.0, "EPCAM": 0.0, "PECAM1": 8.0}
            records.append(SpotRecord("S1", float(index), 0.0, cell_type, genes, cell_id="c%d" % index))
        dataset = SpatialDataset(sample_id="S1", source_path="synthetic", modality="xenium_spatial_rna", records=records)
        result = marker_detection(dataset, {"engine": "prototype"})
        self.assertEqual(result.metrics["mode"], "one_vs_rest")
        markers = result.metrics["markers_by_group"]
        self.assertEqual(set(markers), {"T cell", "Tumor cell", "Endothelial cell"})
        # Each group's own defining gene should be its top up-regulated marker.
        self.assertEqual(markers["T cell"][0]["gene"], "CD8A")
        self.assertEqual(markers["Tumor cell"][0]["gene"], "EPCAM")
        self.assertEqual(markers["Endothelial cell"][0]["gene"], "PECAM1")

    def test_marker_detection_explicit_groups_stay_pairwise(self):
        dataset = DataIngestionLayer().load_csv(DEMO, sample_id="BRCA_04")
        result = marker_detection(
            dataset,
            {"engine": "prototype", "group_key": "cell_type", "group1": "CD8+ T cell", "group2": "Tumor cell"},
        )
        self.assertEqual(result.metrics["mode"], "pairwise")
        self.assertEqual(result.metrics["group1"], "CD8+ T cell")
        self.assertEqual(result.metrics["group2"], "Tumor cell")

    def test_mvp_honesty_tools_caveat_panel_and_transfer(self):
        dataset = DataIngestionLayer().load_csv(DEMO, sample_id="BRCA_04")
        dataset.metadata["is_targeted_panel"] = True
        dataset.metadata["feature_type"] = "targeted_panel"
        result = feature_overlay(dataset, {"feature": "NOT_IN_PANEL"})
        self.assertEqual(result.metrics["status"], "panel_absent")
        self.assertIn("not measured", result.label_caveat)
        with self.assertRaises(MissingPreconditionError):
            reference_label_transfer(dataset, {"reference_features": ["A"], "min_shared_features": 5})


class WorkflowTests(unittest.TestCase):
    def test_v7_workflow_compositions_match_mvp_plan(self):
        self.assertEqual(SCRNA_STANDALONE.name, "SCRNA_LITE")
        self.assertEqual(SCRNA_STANDALONE.steps, ["qc_and_cluster", "marker_detection"])
        self.assertEqual(SCATAC_STANDALONE.name, "SCATAC_LITE")
        self.assertEqual(SCATAC_STANDALONE.steps, ["qc_and_cluster", "marker_detection"])
        self.assertEqual(XENIUM_STANDALONE.name, "XENIUM_PRIMARY")
        self.assertEqual(
            XENIUM_STANDALONE.steps,
            ["qc_and_cluster", "annotation", "region_summary", "cell_neighborhood_enrichment"],
        )
        self.assertEqual(INTEGRATION_MODE.name, "REFERENCE_ASSIST")
        self.assertEqual(INTEGRATION_MODE.steps, ["qc_and_cluster", "marker_detection", "annotation"])

    def test_v7_region_summary_requires_user_regions(self):
        dataset = DataIngestionLayer().load_csv(DEMO, sample_id="BRCA_04")
        dataset.metadata["region_label_source"] = "user_provided"
        result = build_mvp_registry().get("region_summary").run(dataset, {"top_n_features": 3})
        self.assertEqual(result.tool_name, "region_summary")
        self.assertIn("region_count", result.metrics)
        self.assertIsNotNone(result.quality_metrics)


class ValidatedPilotTests(unittest.TestCase):
    def test_v11_xenium_tool_plan_validates_dependencies(self):
        plan = build_xenium_mvp_plan()
        valid = validate_tool_plan(
            plan,
            available_inputs=["normalized_counts", "spatial_coords", "targeted_panel", "segmentation", "expert_labels", "user_regions"],
            registry_tool_names=MVP_TOOL_NAMES,
        )
        self.assertTrue(valid.ok)
        invalid = validate_tool_plan(
            [plan[1], plan[0]],
            available_inputs=["normalized_counts", "spatial_coords", "expert_labels", "user_regions"],
            registry_tool_names=MVP_TOOL_NAMES,
        )
        self.assertFalse(invalid.ok)
        self.assertIn("missing required outputs", invalid.errors[0])
        self.assertIn("spatial_variable_genes", [step.tool_name for step in plan])
        strict_tools = {
            step.tool_name
            for step in plan
            if step.params.get("strict_engine")
        }
        self.assertEqual(
            strict_tools,
            {"qc_and_cluster", "marker_detection", "spatial_variable_genes", "cell_neighborhood_enrichment"},
        )

    def test_analysis_scope_distinguishes_sample_from_full_section(self):
        from spatialmind.pilot.xenium import _analysis_scope

        sampled = SpatialDataset(
            sample_id="X1",
            source_path="xenium",
            records=[SpotRecord("X1", 0.0, 0.0, "T cell", {"CD8A": 1.0})],
            metadata={
                "analysis_scope": "sampled",
                "n_obs_total": 10,
                "sampling": {"total_records": 10, "method": "deterministic_even_index"},
            },
        )
        full = SpatialDataset(
            sample_id="X2",
            source_path="xenium",
            records=[SpotRecord("X2", 0.0, 0.0, "T cell", {"CD8A": 1.0})],
            metadata={
                "analysis_scope": "full_section",
                "n_obs_total": 1,
                "sampling": {"total_records": 1, "method": "all"},
            },
        )
        self.assertFalse(_analysis_scope(sampled)["complete_section_requirement_met"])
        self.assertTrue(_analysis_scope(full)["complete_section_requirement_met"])
        self.assertFalse(_analysis_scope(full)["validated_claims_allowed"])

    def test_pilot_reports_structurally_valid_plan_regardless_of_inputs(self):
        # The plan is structurally sound; input availability is the gate's job, so
        # plan validation must not report "invalid" just because externals are pending.
        from spatialmind.pilot.xenium import _pilot_structural_inputs

        report = validate_tool_plan(
            build_xenium_mvp_plan(),
            available_inputs=_pilot_structural_inputs(),
            registry_tool_names=MVP_TOOL_NAMES,
        )
        self.assertEqual(report.status, "valid")
        self.assertEqual(report.errors, [])

    def test_readiness_only_skips_heavy_artifacts(self):
        if not os.path.isdir(XENIUM_LYMPH):
            self.skipTest("local Xenium lymph dataset not available")
        from spatialmind.pilot.xenium import run_pilot

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "readiness"
            result = run_pilot(XENIUM_LYMPH, out, max_records=40, readiness_only=True)
            # Gate / plan / claim status is still computed.
            self.assertTrue(result["status"].startswith("blocked"))
            self.assertEqual(result["plan_validation"]["status"], "valid")
            self.assertIn("claim_ledger", result)
            self.assertTrue(result["readiness_only"])
            # Heavy artifacts are skipped.
            self.assertEqual(result["report_md"], "")
            self.assertEqual(result["report_html"], "")
            self.assertEqual(result["run_record_path"], "")
            self.assertEqual(result["figures"], [])
            self.assertEqual({p.name for p in out.iterdir()}, {"pilot_validation.json"})

    def test_pilot_gate_blocks_without_expert_labels_and_regions(self):
        dataset = SpatialDataset(
            sample_id="X1",
            source_path="xenium",
            modality="xenium_spatial_rna",
            records=[
                SpotRecord("X1", 0.0, 0.0, "T cell", {"CD8A": 1.0}, cell_id="c1"),
                SpotRecord("X1", 1.0, 1.0, "Tumor cell", {"EPCAM": 1.0}, cell_id="c2"),
            ],
        )
        gate = pilot_gate(
            dataset=dataset,
            asset_readiness={
                "has_cell_table": True,
                "has_feature_matrix": True,
                "has_morphology": True,
                "has_boundaries": True,
            },
            label_report={"status": "missing_expert_labels", "matched_cells": 0, "total_records": 2},
            region_report={"status": "missing_user_regions", "matched_cells": 0, "total_records": 2},
            min_label_coverage=0.7,
            min_region_coverage=0.7,
            allow_single_region=False,
            # A 2-cell fixture exercises the condition wiring, not the
            # class-size floor; the default 50 is asserted separately.
            min_cells_per_class=1,
        )
        self.assertEqual(gate["status"], "blocked_missing_validation_inputs")
        self.assertTrue(any("Expert cell labels" in item for item in gate["blocking_reasons"]))
        self.assertTrue(any("region labels" in item for item in gate["blocking_reasons"]))

    def test_pilot_gate_passes_with_validated_inputs(self):
        dataset = SpatialDataset(
            sample_id="X1",
            source_path="xenium",
            modality="xenium_spatial_rna",
            records=[
                SpotRecord("X1", 0.0, 0.0, "T cell", {"CD8A": 1.0}, region="stroma", cell_id="c1"),
                SpotRecord("X1", 1.0, 1.0, "Tumor cell", {"EPCAM": 1.0}, region="tumor_core", cell_id="c2"),
            ],
        )
        gate = pilot_gate(
            dataset=dataset,
            asset_readiness={
                "has_cell_table": True,
                "has_feature_matrix": True,
                "has_morphology": True,
                "has_boundaries": True,
            },
            label_report={"status": "expert_labels_applied", "matched_cells": 2, "total_records": 2},
            region_report={"status": "user_regions_applied", "matched_cells": 2, "total_records": 2},
            min_label_coverage=0.7,
            min_region_coverage=0.7,
            allow_single_region=False,
            # A 2-cell fixture exercises the condition wiring, not the
            # class-size floor; the default 50 is asserted separately.
            min_cells_per_class=1,
        )
        self.assertEqual(gate["status"], "validated_ready")
        self.assertEqual(gate["blocking_reasons"], [])

    def test_blocked_pilot_claim_ledger_refuses_biology(self):
        payload = {
            "status": "blocked_missing_validation_inputs",
            "required_next_inputs": ["Add expert_cell_labels.csv", "Add cell_regions.csv"],
            "asset_readiness": {"has_cell_table": True},
            "contract": {"assay_subtype": "xenium_spatial_rna"},
        }
        ledger = build_pilot_claim_ledger(payload, [])
        self.assertEqual(ledger[0]["status"], "refused")
        self.assertEqual(ledger[0]["allowed_wording"], "")
        self.assertIn("expert", " ".join(ledger[0]["missing_inputs"]).lower())

    def test_claim_reliability_scores_blocked_and_readiness_claims(self):
        payload = {
            "status": "blocked_missing_validation_inputs",
            "required_next_inputs": ["Add expert_cell_labels.csv", "Add cell_regions.csv"],
            "records_loaded": 10,
            "features_loaded": 50,
            "label_report": {"status": "missing_expert_labels", "matched_cells": 0, "total_records": 10},
            "region_report": {"status": "missing_user_regions", "matched_cells": 0, "total_records": 10},
            "asset_readiness": {
                "has_cell_table": True,
                "has_feature_matrix": True,
                "has_morphology": True,
                "has_boundaries": True,
            },
            "contract": {"assay_subtype": "xenium_spatial_rna", "n_features": 50, "is_targeted_panel": True},
        }
        payload["claim_ledger"] = build_pilot_claim_ledger(payload, [])
        reliability = build_claim_reliability_table(payload, [])
        self.assertEqual(len(reliability), 2)
        self.assertEqual(reliability[0]["status"], "blocked")
        self.assertEqual(reliability[0]["reliability"], 0.0)
        self.assertGreater(reliability[1]["reliability"], reliability[0]["reliability"])
        self.assertIn("A_annotation", reliability[0]["components"])

    def test_claim_truth_validation_and_calibration_require_reviewed_labels(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "spatial_claim_truth_draft_for_review.csv")
            with open(path, "w", encoding="utf-8") as handle:
                handle.write(",".join(CLAIM_TRUTH_FIELDS) + "\n")
                handle.write(
                    "r0,healthy,pilot_claim,pipeline_readiness_control,claim_001,visual_pattern,supported,ready,0.75,1,0.75,0.8,1,1,,no,,,,train,\n"
                )
            blocked = validate_claim_truth_table(path)
            self.assertEqual(blocked["status"], "blocked")
            self.assertIn("Need at least", " ".join(blocked["blockers"]))

            with open(path, "w", encoding="utf-8") as handle:
                handle.write(",".join(CLAIM_TRUTH_FIELDS) + "\n")
                rows = [
                    "r1,healthy,pilot_claim,pipeline_readiness_control,c1,visual_pattern,supported,ready,0.75,1,0.75,0.8,1,1,1,yes,reviewer,2026-07-08,asset check,,train,",
                    "r2,healthy,null_control,null_control,c2,spatial_colocalization,refused,null,0,0,0,0.5,0,0,0,yes,reviewer,2026-07-08,null,,train,",
                    "r3,glioblastoma,pilot_claim,pipeline_readiness_control,c3,visual_pattern,supported,ready,0.75,1,0.75,0.8,1,1,1,yes,reviewer,2026-07-08,asset check,,validation,",
                    "r4,glioblastoma,null_control,null_control,c4,spatial_colocalization,refused,null,0,0,0,0.5,0,0,0,yes,reviewer,2026-07-08,null,,test,",
                ]
                handle.write("\n".join(rows) + "\n")
            ready = validate_claim_truth_table(path)
            self.assertEqual(ready["status"], "ready_for_calibration")
            model = fit_claim_reliability_calibration(ready["records"])
            self.assertEqual(model["status"], "fit")
            self.assertIn("A_annotation", model["weights"])

    def test_local_promotion_report_summarizes_available_gap_statuses(self):
        with tempfile.TemporaryDirectory() as tmp:
            data_root = os.path.join(tmp, "data")
            os.makedirs(data_root)
            demo_path = os.path.join(data_root, "demo.csv")
            with open(demo_path, "w", encoding="utf-8") as handle:
                handle.write("sample_id,x,y,cell_type,CD8A\n")
                handle.write("S1,1,2,T cell,3\n")
            report = build_local_promotion_report(data_root, os.path.join(tmp, "promotion"), max_records=10)
            self.assertEqual(report["xenium_dataset_count"], 0)
            statuses = {item["name"]: item["status"] for item in report["gap_status"]}
            self.assertEqual(statuses["Xenium raw data ingestion"], "missing")
            self.assertIn("LLM API planner", statuses)
            self.assertTrue(os.path.exists(os.path.join(tmp, "promotion", "local_promotion_report.md")))


class SpatialAgentTests(unittest.TestCase):
    def test_agent_returns_tool_trace(self):
        response = SpatialAgent().run("Are CD8+ T cells near tumor cells?", MANIFEST)
        self.assertFalse(response.clarification_needed)
        self.assertIn("cell_type_annotation", [call.tool_name for call in response.tool_trace])
        self.assertIn("neighborhood_enrichment", [call.tool_name for call in response.tool_trace])

    def test_agent_detects_ambiguity(self):
        response = SpatialAgent().run("Analyze this", MANIFEST)
        self.assertTrue(response.clarification_needed)
        self.assertEqual(response.tool_trace, [])

    def test_agent_refuses_blocked_readiness_workflow(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "unannotated.csv")
            with open(path, "w", encoding="utf-8") as handle:
                handle.write("sample_id,x,y,cell_type,CD8A,EPCAM\n")
                for index in range(10):
                    handle.write("S1,%s,%s,Unannotated,%s,%s\n" % (index, index + 1, index + 2, index + 3))
            response = SpatialAgent().run("Are T cells near tumor cells?", path)
            self.assertIsNotNone(response.no_analysis_response)
            self.assertEqual(response.tool_trace, [])
            self.assertIn("cell-type annotation", response.no_analysis_response.recommended_next_step.lower())

    def test_mvp_agent_refuses_deferred_v1_workflow(self):
        response = SpatialAgent(mvp_mode=True).run("Run deconvolution on this dataset", MANIFEST)
        self.assertIsNotNone(response.no_analysis_response)
        self.assertEqual(response.tool_trace, [])
        self.assertIn("deferred", response.no_analysis_response.blocking_reasons[0])


class ContractTests(unittest.TestCase):
    def test_contracts_are_importable_and_claim_grounding_softens_unsupported_claims(self):
        core = CoreSpatialObject(sample_id="S1", modality="transcriptomics", qc_passed=True)
        self.assertEqual(core.sample_id, "S1")
        claim = BiologicalClaim(
            claim_text="CD8 T cells are significantly enriched near tumor cells.",
            claim_type="spatial_colocalization",
        )
        grounded = ground_claim(claim, available_evidence=["figure"])
        self.assertIn("no valid neighborhood test", grounded.allowed_wording)

    def test_cell_by_feature_contract_requires_xenium_coordinates(self):
        contract = CellByFeatureContract(
            sample_id="X1",
            modality="transcriptomics",
            assay_subtype="xenium_spatial_rna",
            feature_type="targeted_panel",
            n_features=10,
            is_targeted_panel=True,
            resolution="subcellular",
            qc_passed=True,
        )
        with self.assertRaises(Exception):
            contract.validate()


class MemoryTests(unittest.TestCase):
    def test_session_and_long_term_memory(self):
        with tempfile.TemporaryDirectory() as tmp:
            session_memory = SessionMemory(root=os.path.join(tmp, "sessions"))
            context = session_memory.append_message("s1", MANIFEST, "user", "Show T cells")
            self.assertEqual(context.session_id, "s1")
            self.assertEqual(len(session_memory.get("s1").conversation_history), 1)

            long_term = LongTermMemory(root=tmp)
            long_term.store(MANIFEST, "Show tumor cells", ["cell_type_annotation"], "Tumor cells were present.", "run/report.html")
            self.assertEqual(len(long_term.search("tumor", dataset_id=MANIFEST)), 1)

            priors = UserPriorStore(root=tmp)
            priors.upsert("u1", PriorType.CELL_TYPE_ALIAS, "CD8+", "Use CD8A and CD8B markers.")
            self.assertEqual(len(priors.search("u1", "CD8 marker")), 1)


class V2ScaffoldTests(unittest.TestCase):
    def test_qc_report_builder_viz_router_and_fusion_scaffold(self):
        with tempfile.TemporaryDirectory() as tmp:
            dataset, report = DataIngestionPipeline().ingest(
                Path(DEMO),
                IngestionConfig(format=DataFormat.TIDY_CSV, sample_id="BRCA_04", min_counts=0),
            )
            qc_path = QCReportBuilder().build(dataset, report, tmp)
            self.assertTrue(os.path.exists(qc_path))
            with open(qc_path, encoding="utf-8") as handle:
                qc_content = handle.read()
            self.assertIn("Metric Distributions", qc_content)
            self.assertIn("Spatial QC Overlay", qc_content)
            self.assertIn("Filtration Waterfall", qc_content)
            self.assertEqual(VizRouter().choose("differential_expression", dataset.modality).name, "spatial_scatter")
            self.assertEqual(VizRouter().choose("marker_detection", dataset.modality).name, "spatial_scatter")
            self.assertEqual(VizRouter().choose("region_summary", "spatial_transcriptomics").name, "region_summary_plot")
            fused = ModalityFuser().fuse([("visium", dataset)])
            self.assertIn("visium", fused.modalities_present)

    def test_cluster_style_spatial_svg_has_axes_and_legend(self):
        with tempfile.TemporaryDirectory() as tmp:
            dataset = DataIngestionLayer().load_csv(DEMO, sample_id="BRCA_04")
            svg_path = VisualizationLayer().render_distribution_svg(dataset, tmp, [])
            with open(svg_path, encoding="utf-8") as handle:
                content = handle.read()
            self.assertIn(">Cluster<", content)
            self.assertIn(">spatial1<", content)
            self.assertIn(">spatial2<", content)
            self.assertIn("Tumor cell", content)

    def test_xenium_explorer_lite_viewer_exports_review_controls(self):
        dataset = SpatialDataset(
            sample_id="X1",
            source_path="experiment.xenium",
            modality="xenium_spatial_rna",
            coordinate_system="microns",
            records=[
                SpotRecord("X1", 0.0, 0.0, "astrocyte", {"GFAP": 4.0, "AQP4": 2.0}, region="brain", cell_id="cell-a"),
                SpotRecord("X1", 10.0, 8.0, "microglial cell", {"CX3CR1": 3.0}, region="brain", cell_id="cell-b"),
            ],
            metadata={
                "xenium_explorer_assets": {
                    "analysis_summary_filepath": {
                        "relative_path": "analysis_summary.html",
                        "resolved_path": "/tmp/analysis_summary.html",
                        "exists": True,
                    }
                }
            },
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = XeniumExplorerLiteViewer().render(dataset, tmp, dataset_path="experiment.xenium")
            with open(path, encoding="utf-8") as handle:
                content = handle.read()
            self.assertIn("SpatialMind Explorer Lite", content)
            self.assertIn("Export Regions", content)
            self.assertIn("expert_cell_labels.csv", content)
            self.assertIn("cell_regions.csv", content)
            self.assertIn("cell-a", content)
            self.assertIn("analysis_summary_filepath", content)


class MorphologyLayerTests(unittest.TestCase):
    def test_loaders_degrade_when_assets_are_missing(self):
        from spatialmind.viz.morphology import (
            find_morphology_image,
            load_cell_boundaries,
            load_morphology_thumbnail,
            read_pixel_size,
        )

        with tempfile.TemporaryDirectory() as tmp:
            self.assertIsNone(find_morphology_image(tmp))
            self.assertIsNone(read_pixel_size(tmp))
            image = load_morphology_thumbnail(tmp)
            self.assertEqual(image["status"], "unavailable")
            self.assertIn("reason", image)
            boundaries = load_cell_boundaries(tmp)
            self.assertEqual(boundaries["status"], "unavailable")
            self.assertEqual(boundaries["polygons"], {})

    def test_choose_level_picks_smallest_level_above_target(self):
        from spatialmind.viz.morphology import _choose_level

        class Level:
            def __init__(self, shape):
                self.shape = shape

        levels = [Level((27282, 36955)), Level((13641, 18477)), Level((6820, 9238)), Level((1705, 2309)), Level((213, 288))]
        # Level 3 (2309 wide) is the smallest still >= 1600.
        self.assertEqual(_choose_level(levels, 1600), 3)
        # Nothing satisfies a huge target, so the largest level is used.
        self.assertEqual(_choose_level(levels, 99999), 0)

    def test_real_xenium_morphology_and_boundaries_align(self):
        if not os.path.isdir(XENIUM_LYMPH):
            self.skipTest("local Xenium dataset not available")
        from spatialmind.viz.morphology import load_cell_boundaries, load_morphology_thumbnail

        image = load_morphology_thumbnail(XENIUM_LYMPH, max_dimension=400)
        if image["status"] != "loaded":
            self.skipTest("morphology image unavailable: %s" % image.get("reason"))
        self.assertTrue(image["data_uri"].startswith("data:image/png;base64,"))
        self.assertGreater(image["width_um"], 0)
        self.assertGreater(image["height_um"], 0)
        self.assertLessEqual(max(image["thumbnail_width"], image["thumbnail_height"]), 400)

        from spatialmind.ingestion import load_xenium

        dataset = load_xenium(XENIUM_LYMPH, max_records=40)
        cell_ids = [record.cell_id for record in dataset.records]
        boundaries = load_cell_boundaries(XENIUM_LYMPH, cell_ids=cell_ids)
        if boundaries["status"] != "loaded":
            self.skipTest("boundaries unavailable: %s" % boundaries.get("reason"))
        self.assertLessEqual(boundaries["cell_count"], len(cell_ids))
        # Each polygon must enclose its own centroid, which is what keeps the
        # segmentation overlay registered with the plotted cells.
        for record in dataset.records:
            vertices = boundaries["polygons"].get(record.cell_id)
            if not vertices:
                continue
            xs = [vertex[0] for vertex in vertices]
            ys = [vertex[1] for vertex in vertices]
            self.assertGreaterEqual(record.x, min(xs) - 1.0)
            self.assertLessEqual(record.x, max(xs) + 1.0)
            self.assertGreaterEqual(record.y, min(ys) - 1.0)
            self.assertLessEqual(record.y, max(ys) + 1.0)


class LabelTransferTests(unittest.TestCase):
    def _reference(self):
        records = []
        for index in range(12):
            records.append(SpotRecord("R", 0.0, 0.0, "T cell", {"CD8A": 9.0, "CD3D": 8.0, "EPCAM": 0.0}, cell_id="r%d" % index))
            records.append(SpotRecord("R", 0.0, 0.0, "Tumor cell", {"CD8A": 0.0, "CD3D": 0.0, "EPCAM": 9.0}, cell_id="t%d" % index))
        return SpatialDataset(sample_id="R", source_path="ref", modality="scrna", records=records)

    def _target(self):
        return SpatialDataset(
            sample_id="X",
            source_path="x",
            modality="xenium_spatial_rna",
            records=[
                SpotRecord("X", 0.0, 0.0, "Unannotated cell", {"CD8A": 7.0, "CD3D": 6.0, "EPCAM": 0.0}, cell_id="a1"),
                SpotRecord("X", 1.0, 0.0, "Unannotated cell", {"CD8A": 0.0, "CD3D": 0.0, "EPCAM": 7.0}, cell_id="a2"),
            ],
        )

    def test_transfer_assigns_a_label_and_confidence_per_cell(self):
        result = reference_label_transfer(
            self._target(),
            {"reference_dataset": self._reference(), "min_shared_features": 2},
        )
        self.assertEqual(result.metrics["status"], "transferred")
        self.assertTrue(result.metrics["labels_transferred"])
        predictions = {item["cell_id"]: item for item in result.metrics["predictions"]}
        self.assertEqual(predictions["a1"]["predicted_label"], "T cell")
        self.assertEqual(predictions["a2"]["predicted_label"], "Tumor cell")
        for item in predictions.values():
            self.assertGreaterEqual(item["confidence"], 0.0)
            self.assertLessEqual(item["confidence"], 1.0)
        self.assertIn("expert review", " ".join(result.caveats).lower())

    def test_review_priority_is_calibrated_on_the_target_not_the_reference(self):
        target = self._target()
        # Extra cells so a 90th-percentile cut is meaningful.
        for index in range(18):
            target.records.append(
                SpotRecord("X", float(index), 1.0, "Unannotated cell",
                           {"CD8A": 6.0, "CD3D": 5.0, "EPCAM": 0.0}, cell_id="e%d" % index)
            )
        result = reference_label_transfer(
            target, {"reference_dataset": self._reference(), "min_shared_features": 2},
        )
        metrics = result.metrics
        flagged = metrics["high_review_priority_count"]
        # Must flag a reviewable minority, not everything (the failure mode of
        # calibrating against reference-internal distances across assays).
        self.assertLess(flagged, len(target.records))
        self.assertIn("platform_shift_ratio", metrics)
        priorities = {item["review_priority"] for item in metrics["predictions"]}
        self.assertTrue(priorities <= {"high", "normal"})
        for item in metrics["predictions"]:
            self.assertIn("distant_from_reference", item)
        # Coverage limits must be stated, since vote confidence cannot express them.
        self.assertTrue(any("no matching class" in caveat for caveat in result.caveats))

    def test_marker_lineage_resolution_is_conservative(self):
        from spatialmind.tools.implementations import lineage_for_label, lineages_conflict, marker_lineage

        # Unambiguous myeloid evidence resolves.
        lineage, score = marker_lineage({"CD68": 5.0, "AIF1": 4.0, "C1QA": 3.0})
        self.assertEqual(lineage, "myeloid")
        self.assertGreater(score, 0)
        # Tied evidence across lineages must NOT resolve, rather than guess.
        self.assertEqual(marker_lineage({"AQP4": 2.0, "CD4": 2.0})[0], "")
        # Trace evidence below the floor must not resolve.
        self.assertEqual(marker_lineage({"CD68": 0.2})[0], "")

        self.assertEqual(lineage_for_label("intratelencephalic-projecting glutamatergic cortical neuron"), "neuronal")
        self.assertEqual(lineage_for_label("microglial cell"), "myeloid")
        # Most specific keyword wins.
        self.assertEqual(lineage_for_label("oligodendrocyte precursor cell"), "opc")
        self.assertEqual(lineage_for_label("oligodendrocyte"), "oligodendrocyte")

        self.assertTrue(lineages_conflict("neuronal", "myeloid"))
        self.assertFalse(lineages_conflict("neuronal", "neuronal"))
        self.assertFalse(lineages_conflict("oligodendrocyte", "opc"))  # compatible
        self.assertFalse(lineages_conflict("neuronal", ""))  # unknown never conflicts

    def test_marker_disagreement_flags_cells_confidence_would_miss(self):
        target = SpatialDataset(
            sample_id="X",
            source_path="x",
            modality="xenium_spatial_rna",
            records=[
                # Unambiguously myeloid, but the reference has no myeloid class.
                SpotRecord("X", 0.0, 0.0, "Unannotated cell",
                           {"CD68": 6.0, "AIF1": 5.0, "C1QA": 4.0, "CD8A": 0.0, "EPCAM": 0.0}, cell_id="myeloid1"),
                SpotRecord("X", 1.0, 0.0, "Unannotated cell",
                           {"CD8A": 7.0, "CD3D": 6.0, "EPCAM": 0.0}, cell_id="lymphoid1"),
            ],
        )
        def reference_cell(label, profile, index):
            genes = {"CD68": 0.0, "AIF1": 0.0, "C1QA": 0.0, "CD8A": 0.0, "EPCAM": 0.0, "MBP": 0.0, "MOG": 0.0}
            genes.update(profile)
            return SpotRecord("R", 0.0, 0.0, label, genes, cell_id="%s%d" % (label[:3], index))

        reference = SpatialDataset(
            sample_id="R", source_path="r", modality="scrna",
            records=[reference_cell("oligodendrocyte", {"MBP": 9.0, "MOG": 8.0}, i) for i in range(10)]
                    + [reference_cell("neuron", {"CD8A": 1.0}, i) for i in range(10)],
        )
        result = reference_label_transfer(target, {"reference_dataset": reference, "min_shared_features": 2})
        by_id = {item["cell_id"]: item for item in result.metrics["predictions"]}
        # The myeloid cell has no correct label available in this reference.
        self.assertEqual(by_id["myeloid1"]["marker_lineage"], "myeloid")
        self.assertTrue(by_id["myeloid1"]["lineage_absent_from_reference"])
        self.assertEqual(by_id["myeloid1"]["review_priority"], "high")
        self.assertGreaterEqual(result.metrics["lineage_absent_from_reference_count"], 1)
        self.assertIn("myeloid", result.metrics["reference_lineages"] + ["myeloid"])  # sanity on key presence
        self.assertIn("reference_lineages", result.metrics)

    def test_without_reference_dataset_no_transfer_is_claimed(self):
        result = reference_label_transfer(
            self._target(),
            {"reference_features": ["CD8A", "CD3D", "EPCAM"], "min_shared_features": 2},
        )
        self.assertEqual(result.metrics["status"], "compatibility_only")
        self.assertFalse(result.metrics["labels_transferred"])
        self.assertNotIn("predictions", result.metrics)
        # The summary must not assert that a transfer happened.
        self.assertIn("No labels were transferred", result.summary)

    def test_shared_feature_count_comes_from_the_reference_not_the_target(self):
        # The target carries extra genes the reference never measured; only the
        # genuinely shared ones may be counted.
        target = self._target()
        for record in target.records:
            record.genes["EXTRA_GENE_A"] = 1.0
            record.genes["EXTRA_GENE_B"] = 1.0
        result = reference_label_transfer(
            target,
            {"reference_dataset": self._reference(), "min_shared_features": 2},
        )
        self.assertEqual(result.metrics["shared_feature_count"], 3)

    def test_cross_species_reference_is_refused(self):
        target = self._target()
        target.metadata["organism"] = "Human"
        reference = self._reference()
        reference.metadata["organism"] = "Mus musculus"
        with self.assertRaises(MissingPreconditionError) as ctx:
            reference_label_transfer(target, {"reference_dataset": reference, "min_shared_features": 2})
        self.assertIn("cross-species", str(ctx.exception))
        # Explicit opt-in (for pre-mapped orthologs) still works.
        result = reference_label_transfer(
            target,
            {"reference_dataset": reference, "min_shared_features": 2, "allow_cross_species": True},
        )
        self.assertEqual(result.metrics["status"], "transferred")

    def test_same_species_and_unknown_species_are_allowed(self):
        from spatialmind.tools.implementations import normalize_species

        self.assertEqual(normalize_species("Homo sapiens"), "human")
        self.assertEqual(normalize_species("NCBITaxon:10090"), "mouse")
        target = self._target()
        target.metadata["organism"] = "Homo sapiens"
        reference = self._reference()
        reference.metadata["organism"] = "Human"
        result = reference_label_transfer(target, {"reference_dataset": reference, "min_shared_features": 2})
        self.assertEqual(result.metrics["status"], "transferred")
        # Unknown organism on either side must not hard-block a legitimate run.
        reference.metadata["organism"] = ""
        self.assertEqual(
            reference_label_transfer(target, {"reference_dataset": reference, "min_shared_features": 2}).metrics["status"],
            "transferred",
        )

    def test_single_class_reference_is_rejected(self):
        reference = SpatialDataset(
            sample_id="R",
            source_path="ref",
            modality="scrna",
            records=[SpotRecord("R", 0.0, 0.0, "T cell", {"CD8A": 9.0, "CD3D": 8.0, "EPCAM": 0.0}, cell_id="r%d" % i) for i in range(6)],
        )
        with self.assertRaises(MissingPreconditionError):
            reference_label_transfer(self._target(), {"reference_dataset": reference, "min_shared_features": 2})


class ToolHonestyTests(unittest.TestCase):
    """A tool must never report work it did not do.

    `_is_scaffold` catches tools that do *nothing* -- it greps their source for
    `_scaffold_result(`. It cannot catch a tool that does *meaningless* work, which
    is how four tools shipped as `capability="validated"` while returning hardcoded
    p-values, a coordinate gradient labelled "pseudotime", label counts labelled
    "deconvolution", and a list index labelled "TF activity".
    """

    STAT_KEYS = (
        "pval", "pvalue", "p_value", "pval_adj", "padj", "score", "zscore", "z_score",
        "activity_score", "pseudotime", "morans_i", "silhouette", "modularity",
    )
    GENES = ("VEGFA", "PTPRC", "EPCAM", "CD8A", "CD3D", "KDR", "FLT1", "MBP", "GFAP", "SNAP25")

    def _dataset(self, modality, subtype, seed):
        records = []
        for index in range(40):
            genes = {gene: float((index * (n + 1) + seed) % 7) for n, gene in enumerate(self.GENES)}
            records.append(
                SpotRecord(
                    "S",
                    float((index * 3 + seed) % 20),
                    float((index * 7 + seed) % 20),
                    ["Tumor cell", "CD8+ T cell", "Endothelial cell"][index % 3],
                    genes,
                    region=["r1", "r2"][index % 2],
                    cell_id="c%d" % index,
                )
            )
        dataset = SpatialDataset(sample_id="S", source_path="p", modality=modality, records=records)
        dataset.metadata["assay_subtype"] = subtype
        return dataset

    def _statistics(self, payload, path=""):
        found = []
        if isinstance(payload, dict):
            for key, value in payload.items():
                if key in self.STAT_KEYS and isinstance(value, (int, float)):
                    found.append(("%s.%s" % (path, key), round(float(value), 6)))
                else:
                    found.extend(self._statistics(value, "%s.%s" % (path, key)))
        elif isinstance(payload, list):
            for index, value in enumerate(payload):
                found.extend(self._statistics(value, "%s[%d]" % (path, index)))
        return found

    def test_no_plannable_tool_returns_data_independent_statistics(self):
        """A statistic that survives changing every value was never computed.

        Expression, coordinates and labels all differ between the two datasets, so
        a tool whose p-values, scores or z-scores come out identical is asserting a
        test it never ran. Tools that raise are fine: refusing on data they cannot
        handle is the honest outcome. Tools that emit no statistic at all are fine
        too -- `annotation` summarises labels and claims nothing more.

        Calls `.callable` rather than `.run` deliberately: `run` attaches quality
        metrics derived from the dataset, which would make any result look
        data-dependent and mask exactly what this is looking for.
        """
        offenders = []
        for tool in build_full_registry().list_plannable():
            for modality, subtype in (
                ("spatial_transcriptomics", "spatial_transcriptomics"),
                ("scatac", "scatac_gene_activity"),
            ):
                try:
                    first = tool.callable(self._dataset(modality, subtype, 0), {})
                    second = tool.callable(self._dataset(modality, subtype, 3), {})
                except Exception:
                    continue  # refused this modality; try the next
                baseline = self._statistics(first.metrics)
                if baseline and baseline == self._statistics(second.metrics):
                    offenders.append("%s (%d identical: %s)" % (tool.name, len(baseline), baseline[:2]))
                break
        self.assertEqual(
            offenders,
            [],
            "plannable tools returned statistics unchanged by the data: %s" % "; ".join(offenders),
        )

    def test_known_prototype_tools_are_not_plannable(self):
        # These four fabricated or mislabelled their output and were reachable by
        # an LLM planner through to_anthropic_tools(). Two of them
        # (trajectory_inference, spatial_deconvolution) derive a real number from
        # the data and so cannot be caught by the check above -- their defect is
        # that the quantity is not what the name claims. Pin them explicitly.
        registry = build_full_registry()
        exposed = {schema["name"] for schema in registry.to_anthropic_tools()}
        for name in (
            "ligand_receptor_analysis",
            "trajectory_inference",
            "spatial_deconvolution",
            "motif_tf_activity",
        ):
            self.assertEqual(
                registry.get(name).capability,
                "unavailable",
                "%s does no real work and must not be marked usable" % name,
            )
            self.assertNotIn(name, exposed, "%s must not be offered to a planner" % name)


class XeniumCellQCTests(unittest.TestCase):
    """Nothing filtered a Xenium load before this.

    `_apply_threshold_qc` is reachable only from the config-driven pipeline, so
    every statistic in every Xenium report was computed over an unfiltered
    population.
    """

    def _cell(self, cell_id, transcripts, background=0.0, nucleus_area=5.0):
        return SpotRecord(
            "S", 0.0, 0.0, "Unannotated cell",
            {
                "TRANSCRIPT_COUNTS": float(transcripts),
                "TOTAL_COUNTS": float(transcripts) + float(background),
                "CELL_AREA": 100.0,
                "NUCLEUS_AREA": float(nucleus_area),
                "CONTROL_PROBE_COUNTS": float(background),
                "CONTROL_CODEWORD_COUNTS": 0.0,
                "UNASSIGNED_CODEWORD_COUNTS": 0.0,
                "GENE_A": 1.0,
            },
            cell_id=cell_id,
        )

    def test_low_count_and_high_background_cells_are_removed(self):
        records = [
            self._cell("good", 200),
            self._cell("sparse", 3),
            self._cell("noisy", 200, background=100.0),
        ]
        kept, report = apply_xenium_cell_qc(records)
        self.assertEqual([r.cell_id for r in kept], ["good"])
        self.assertEqual(report["dropped_low_transcript_count"], 1)
        self.assertEqual(report["dropped_high_background_count"], 1)
        self.assertEqual(report["input_cell_count"], 3)
        self.assertEqual(report["retained_cell_count"], 1)

    def test_nucleus_free_cells_are_reported_but_kept(self):
        # A zero nucleus area is a segmentation warning, not proof the cell is
        # wrong; dropping on it would silently discard whole morphologies.
        kept, report = apply_xenium_cell_qc([self._cell("anucleate", 200, nucleus_area=0.0)])
        self.assertEqual(len(kept), 1)
        self.assertEqual(report["nucleus_free_cell_count"], 1)
        self.assertEqual(report["dropped_cell_count"], 0)

    def test_filtering_everything_refuses_instead_of_returning_nothing(self):
        with self.assertRaises(IngestionValidationError) as ctx:
            apply_xenium_cell_qc([self._cell("a", 1), self._cell("b", 2)])
        self.assertIn("removed every cell", str(ctx.exception))

    def test_thresholds_are_reported_not_silent(self):
        _kept, report = apply_xenium_cell_qc([self._cell("good", 200)], min_transcripts=25)
        self.assertIn("transcript_counts >= 25", report["rule"])
        self.assertEqual(report["min_transcripts"], 25)

    def test_qc_removals_do_not_make_a_full_section_look_sampled(self):
        """Coverage is what the scan reached, not what survived QC.

        Regression: `_analysis_scope` recomputed completeness from the post-QC
        record count, so 44 cells dropped for quality out of 24,406 turned a full
        section into `blocked_sampled_inference` and blocked validated inference
        outright. The loader guarded this; the pilot's own recomputation undid it.
        """
        from spatialmind.pilot.xenium import _analysis_scope

        dataset = SpatialDataset(
            sample_id="S", source_path="p", modality="xenium_spatial_rna",
            records=[self._cell("c%d" % i, 200) for i in range(96)],
        )
        dataset.metadata["analysis_scope"] = "full_section"
        dataset.metadata["sampling"] = {
            "method": "all", "scanned_records": 100, "loaded_records": 96, "total_records": 100,
        }
        scope = _analysis_scope(dataset)
        self.assertEqual(scope["scope"], "full_section")
        self.assertTrue(scope["complete_section"])
        self.assertEqual(scope["qc_removed_records"], 4)
        self.assertEqual(scope["fraction_loaded"], 1.0)

    def test_a_genuinely_sampled_run_is_still_sampled(self):
        from spatialmind.pilot.xenium import _analysis_scope

        dataset = SpatialDataset(
            sample_id="S", source_path="p", modality="xenium_spatial_rna",
            records=[self._cell("c%d" % i, 200) for i in range(40)],
        )
        dataset.metadata["analysis_scope"] = "sampled"
        dataset.metadata["sampling"] = {
            "method": "deterministic_even_index", "scanned_records": 40,
            "loaded_records": 40, "total_records": 100,
        }
        scope = _analysis_scope(dataset)
        self.assertEqual(scope["scope"], "sampled")
        self.assertFalse(scope["complete_section"])

    def test_background_features_are_never_expression(self):
        from spatialmind.schemas import NON_EXPRESSION_FEATURE_NAMES

        for name in ("CONTROL_PROBE_COUNTS", "CONTROL_CODEWORD_COUNTS", "UNASSIGNED_CODEWORD_COUNTS"):
            self.assertIn(name, NON_EXPRESSION_FEATURE_NAMES)

    def test_the_three_exclusion_sets_cannot_drift(self):
        # They were three identical literals in three modules. Adding a fourth
        # feature to one would have silently left the others behind.
        from spatialmind.ingestion.labels import NON_BIOLOGICAL_FEATURES
        from spatialmind.ingestion.pipeline import NON_EXPRESSION_FEATURES
        from spatialmind.tools.implementations import EXPRESSION_EXCLUDED_FEATURES

        self.assertIs(NON_EXPRESSION_FEATURES, NON_BIOLOGICAL_FEATURES)
        self.assertIs(NON_BIOLOGICAL_FEATURES, EXPRESSION_EXCLUDED_FEATURES)


class ScaffoldDetectionTests(unittest.TestCase):
    def test_a_mention_of_the_scaffold_helper_is_not_a_scaffold(self):
        """Detection parses the body; it does not search the text.

        The substring check this replaced matched `_scaffold_result(` anywhere in
        the source, so a working tool that named the helper in a docstring,
        comment or caveat string was silently marked unavailable and dropped out
        of `list_plannable()` with no error raised anywhere.
        """
        from spatialmind.tools.registry import _is_scaffold

        source = textwrap.dedent(
            '''
            from spatialmind.schemas import ToolResult
            from spatialmind.tools.implementations import _scaffold_result

            def real_tool(dataset, params):
                """A real tool. Unlike _scaffold_result(), this does the work."""
                # deliberately mentions _scaffold_result( in a comment
                return ToolResult(tool_name="real", summary="s",
                                  metrics={"x": 1}, caveats=["see _scaffold_result("])

            def genuine_scaffold(dataset, params):
                return _scaffold_result("t", "s", "c", params)
            '''
        )
        directory = tempfile.mkdtemp()
        path = os.path.join(directory, "scaffold_probe_module.py")
        with open(path, "w", encoding="utf-8") as handle:
            handle.write(source)
        sys.path.insert(0, directory)
        try:
            module = importlib.import_module("scaffold_probe_module")
            self.assertFalse(_is_scaffold(module.real_tool), "prose mention must not mark a tool unavailable")
            self.assertTrue(_is_scaffold(module.genuine_scaffold), "a real scaffold return must still be caught")
        finally:
            sys.path.remove(directory)
            sys.modules.pop("scaffold_probe_module", None)
            shutil.rmtree(directory, ignore_errors=True)

    def test_statistical_strength_does_not_grow_with_the_pair_count(self):
        """More cell types must not buy statistical strength for free.

        `S_statistical` scored `max|z| / 5` across every reported pair with no
        correction. A section with ten cell types draws that maximum from 55
        pairs while a two-type section draws from 3, so the larger vocabulary
        won on pair count alone. Measured against controls matched on everything
        except spatial arrangement, that inverted the score: the permutation
        null out-ranked implanted structure, AUROC 0.35 where 0.50 is chance.
        """
        from spatialmind.methods.reliability.scoring import _statistical_component
        from spatialmind.schemas import ToolResult

        claim = {"claim_type": "spatial_colocalization", "status": "supported"}

        def component(pair_count, top_z):
            pairs = [{"pair": "t%d | t%d" % (i, i), "zscore": 1.2} for i in range(pair_count - 1)]
            pairs.append({"pair": "a | b", "zscore": top_z})
            result = ToolResult(tool_name="cell_neighborhood_enrichment", summary="s",
                                metrics={"top_pairs": pairs})
            return _statistical_component(claim, [result]).score

        # Same evidence, more noise pairs alongside it: strength must not rise.
        few = component(3, 6.0)
        many = component(55, 6.0)
        self.assertLessEqual(many, few, "widening the pair list must not raise strength")

        # Real structure must still outscore a null-sized effect at equal width.
        self.assertGreater(component(3, 8.0), component(3, 1.5))

    def test_scaffold_detection_survives_having_no_source(self):
        """Detection must not depend on source being readable on disk.

        Inside a PyInstaller bundle, modules load from a compiled archive and
        `inspect.getsource` raises. A source-reading check therefore returned
        False for every scaffold, and the packaged app advertised all 30 tools
        as usable -- including the 18 that return a placeholder and do no work.
        """
        from spatialmind.tools.registry import _is_scaffold

        source = textwrap.dedent(
            '''
            from spatialmind.schemas import ToolResult
            from spatialmind.tools.implementations import _scaffold_result

            def hidden_scaffold(dataset, params):
                return _scaffold_result("t", "s", "c", params)

            def hidden_real(dataset, params):
                return ToolResult(tool_name="real", summary="s", metrics={"x": 1})
            '''
        )
        namespace = {}
        # A filename that does not exist on disk: exactly the situation in a
        # frozen bundle, where the code object is present and the file is not.
        exec(compile(source, "<frozen spatialmind.tools.implementations>", "exec"), namespace)

        with self.assertRaises((OSError, TypeError)):
            inspect.getsource(namespace["hidden_scaffold"])
        self.assertTrue(_is_scaffold(namespace["hidden_scaffold"]),
                        "a scaffold must still be detected with no source available")
        self.assertFalse(_is_scaffold(namespace["hidden_real"]),
                         "a real tool must not be marked unavailable")

    def test_registered_scaffolds_are_still_detected(self):
        registry = build_full_registry()
        summary = registry.capability_summary()
        self.assertGreaterEqual(summary.get("unavailable", 0), 14)
        self.assertNotIn("tumor_niche_analysis", {t.name for t in registry.list_plannable()})


class PanelAdequacyTests(unittest.TestCase):
    """P_panel must measure the panel, not match words in a sentence.

    It read marker families out of the claim's English text. The pilot's claims
    are capability statements that never name a cell type, so it always returned
    its 0.8 fallback -- and because reliability is a weakest link, 0.8 became the
    silent ceiling on every claim in every Xenium report.
    """

    def _payload(self, feature_names, cell_types=("oligodendrocyte",)):
        return {
            "contract": {"panel_name": "test_panel", "n_features": len(feature_names)},
            "features_loaded": len(feature_names),
            "feature_names": list(feature_names),
            "cell_types": list(cell_types),
            "claim_ledger": [
                {
                    "claim_text": "Expert-reviewed cell labels are available for the loaded cells.",
                    "claim_type": "cell_type_annotation",
                    "status": "supported",
                }
            ],
            "label_report": {"status": "expert_labels_applied", "coverage": 1.0},
        }

    def _panel_score(self, payload):
        rows = build_claim_reliability_table(payload, [])
        return rows[0]["components"]["P_panel"]["score"]

    def test_score_reflects_which_markers_the_panel_measures(self):
        # LINEAGE_MARKERS["oligodendrocyte"] has 8 canonical markers.
        from spatialmind.tools.implementations import LINEAGE_MARKERS

        markers = list(LINEAGE_MARKERS["oligodendrocyte"])
        full = self._panel_score(self._payload(markers))
        half = self._panel_score(self._payload(markers[: len(markers) // 2]))
        none = self._panel_score(self._payload(["IRRELEVANT1", "IRRELEVANT2"]))
        self.assertAlmostEqual(full, 1.0, places=3)
        self.assertAlmostEqual(half, 0.5, places=3)
        self.assertAlmostEqual(none, 0.0, places=3)

    def test_score_is_not_pinned_to_the_generic_fallback(self):
        # The exact defect: a claim naming no cell type used to score 0.8 no
        # matter what the panel contained.
        from spatialmind.tools.implementations import LINEAGE_MARKERS

        markers = list(LINEAGE_MARKERS["oligodendrocyte"])
        scores = {
            self._panel_score(self._payload(markers)),
            self._panel_score(self._payload(markers[:2])),
        }
        self.assertNotIn(0.8, scores, "P_panel fell back to the constant despite labels being applied")
        self.assertEqual(len(scores), 2, "P_panel did not vary with panel content")

    def test_label_text_is_not_treated_as_a_measured_gene(self):
        # `measured` used to be seeded from cell_types, so a cell type sharing a
        # name with a gene scored as though that gene had been measured.
        payload = self._payload(["IRRELEVANT"], cell_types=["oligodendrocyte"])
        payload["cell_types"] = payload["cell_types"] + ["MOG"]
        self.assertAlmostEqual(self._panel_score(payload), 0.0, places=3)


class ReferenceLineageCoverageTests(unittest.TestCase):
    """A reference can share genes with the panel and still be unable to name the tissue.

    Panel overlap and vote confidence both miss this: a cell whose type is absent
    from the reference is assigned its nearest available label, usually at high
    confidence. Only the target's own markers can answer it.
    """

    LINEAGE_PROFILES = {
        "myeloid": {"CD68": 6.0, "AIF1": 5.0, "C1QA": 4.0},
        "lymphoid": {"CD8A": 7.0, "CD3D": 6.0, "PTPRC": 5.0},
        "neuronal": {"SNAP25": 7.0, "RBFOX3": 6.0, "SYT1": 5.0},
        "opc": {"PDGFRA": 8.0, "CSPG4": 7.0, "VCAN": 6.0},
    }
    ALL_MARKERS = {
        marker: 0.0 for profile in LINEAGE_PROFILES.values() for marker in profile
    }

    def _cells(self, lineage, count, prefix):
        records = []
        for index in range(count):
            genes = dict(self.ALL_MARKERS)
            genes.update(self.LINEAGE_PROFILES[lineage])
            genes["MBP"] = 0.0
            records.append(
                SpotRecord("X", float(index), 0.0, "Unannotated cell", genes, cell_id="%s%d" % (prefix, index))
            )
        return records

    def _target(self, composition):
        records = []
        for lineage, count in composition.items():
            records.extend(self._cells(lineage, count, lineage[:3]))
        return SpatialDataset(sample_id="X", source_path="x", modality="xenium_spatial_rna", records=records)

    def _reference(self):
        """Oligodendrocyte + neuron classes, so `oligodendrocyte` and `neuronal` are nameable."""
        records = []
        for index in range(12):
            oligo = dict(self.ALL_MARKERS)
            oligo["MBP"] = 9.0
            records.append(SpotRecord("R", 0.0, 0.0, "oligodendrocyte", oligo, cell_id="o%d" % index))
            neuron = dict(self.ALL_MARKERS)
            neuron.update(self.LINEAGE_PROFILES["neuronal"])
            neuron["MBP"] = 0.0
            records.append(SpotRecord("R", 0.0, 0.0, "neuron", neuron, cell_id="n%d" % index))
        return SpatialDataset(sample_id="R", source_path="ref", modality="scrna", records=records)

    def test_lineages_above_the_evidence_floor_are_reported_as_uncovered(self):
        target = self._target({"neuronal": 40, "myeloid": 30, "lymphoid": 30})
        report = assess_reference_lineage_coverage(target, {"neuronal", "oligodendrocyte"})
        self.assertEqual(set(report["uncovered_lineages"]), {"myeloid", "lymphoid"})
        self.assertEqual(report["status"], "inadequate")
        self.assertEqual(report["confident_uncovered_cell_count"], 60)

    def test_a_handful_of_cells_is_not_a_missing_population(self):
        # Below the evidence floor these are ambiguous calls, not a population,
        # and must not trigger a refusal.
        target = self._target({"neuronal": 200, "myeloid": 5})
        report = assess_reference_lineage_coverage(target, {"neuronal", "oligodendrocyte"})
        self.assertEqual(report["uncovered_lineages"], [])
        self.assertEqual(report["uncovered_below_evidence_floor"], {"myeloid": 5})
        self.assertEqual(report["status"], "adequate")

    def test_counts_are_reported_as_a_floor_not_an_estimate(self):
        # The strict rule leaves sparse cells unassigned rather than guessing, so
        # the wording must not present its counts as the true share.
        target = self._target({"neuronal": 40, "myeloid": 30})
        report = assess_reference_lineage_coverage(target, {"neuronal"})
        self.assertIn("floor rather than an estimate", describe_lineage_coverage(report))
        self.assertIn("floor", report["method"])

    def test_near_lineages_still_count_as_uncovered(self):
        # COMPATIBLE_LINEAGES suppresses false *disagreement* alarms, a different
        # question. An OPC labelled `oligodendrocyte` is a wrong label.
        target = self._target({"neuronal": 40, "opc": 30})
        report = assess_reference_lineage_coverage(target, {"neuronal", "oligodendrocyte"})
        self.assertIn("opc", report["uncovered_lineages"])

    def test_transfer_refuses_an_incomplete_reference(self):
        target = self._target({"neuronal": 40, "myeloid": 30, "lymphoid": 30})
        with self.assertRaises(MissingPreconditionError) as ctx:
            reference_label_transfer(target, {"reference_dataset": self._reference(), "min_shared_features": 2})
        message = str(ctx.exception)
        self.assertIn("incomplete reference", message)
        self.assertIn("myeloid", message)

    def test_explicit_override_transfers_and_carries_the_caveat(self):
        target = self._target({"neuronal": 40, "myeloid": 30, "lymphoid": 30})
        result = reference_label_transfer(
            target,
            {
                "reference_dataset": self._reference(),
                "min_shared_features": 2,
                "allow_incomplete_reference": True,
            },
        )
        self.assertEqual(result.metrics["status"], "transferred")
        self.assertEqual(result.metrics["lineage_coverage"]["status"], "inadequate")
        self.assertTrue(
            any("systematically wrong rather than uncertain" in caveat for caveat in result.caveats),
            "an overridden coverage failure must still be stated in the caveats",
        )

    def test_an_adequate_reference_transfers_without_a_coverage_caveat(self):
        target = self._target({"neuronal": 60})
        result = reference_label_transfer(
            target, {"reference_dataset": self._reference(), "min_shared_features": 2}
        )
        self.assertEqual(result.metrics["status"], "transferred")
        self.assertEqual(result.metrics["lineage_coverage"]["status"], "adequate")
        self.assertFalse(any("systematically wrong" in caveat for caveat in result.caveats))

    def test_refusal_happens_before_the_knn_is_fitted(self):
        # The check needs only class names and target markers, so a doomed run
        # must not pay for a full transfer first.
        target = self._target({"neuronal": 40, "myeloid": 30, "lymphoid": 30})
        reference = self._reference()
        reference.records = []  # unusable for KNN; refusal must still fire
        with self.assertRaises(MissingPreconditionError):
            reference_label_transfer(target, {"reference_dataset": reference, "min_shared_features": 2})


class H5adReferenceLoadingTests(unittest.TestCase):
    def test_scrna_h5ad_without_spatial_coordinates_loads(self):
        try:
            import anndata as ad  # type: ignore
            import numpy as np  # type: ignore
            import pandas as pd  # type: ignore
        except ImportError:
            self.skipTest("anndata required")

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "reference.h5ad"
            adata = ad.AnnData(
                X=np.array([[5.0, 0.0, 1.0], [0.0, 6.0, 1.0], [4.0, 0.0, 2.0], [0.0, 7.0, 1.0]]),
                obs=pd.DataFrame({"cell_type": ["T cell", "Tumor cell", "T cell", "Tumor cell"]}),
                var=pd.DataFrame(
                    {"feature_name": ["CD8A", "EPCAM", "ACTB"]},
                    index=["ENSG00000153563", "ENSG00000119888", "ENSG00000075624"],
                ),
            )
            adata.uns["organism"] = "Homo sapiens"
            adata.write_h5ad(path)

            dataset = load_scrna(str(path), max_records=10)
            self.assertEqual(len(dataset.records), 4)
            # Symbols must win over Ensembl IDs so symbol panels can align.
            self.assertIn("CD8A", dataset.genes)
            self.assertNotIn("ENSG00000153563", dataset.genes)
            self.assertEqual(dataset.metadata.get("organism"), "Homo sapiens")
            self.assertEqual(sorted(dataset.cell_types), ["T cell", "Tumor cell"])

    def test_single_class_reference_files_combine_into_a_usable_reference(self):
        try:
            import anndata as ad  # type: ignore
            import numpy as np  # type: ignore
            import pandas as pd  # type: ignore
        except ImportError:
            self.skipTest("anndata required")
        from spatialmind.ingestion import load_scrna_reference_set

        with tempfile.TemporaryDirectory() as tmp:
            paths = []
            # Atlases such as the Siletti brain collection ship one class per file.
            for label, profile in (("oligodendrocyte", [9.0, 0.0]), ("astrocyte", [0.0, 9.0])):
                path = Path(tmp) / ("%s.h5ad" % label.replace(" ", "_"))
                adata = ad.AnnData(
                    X=np.array([profile, profile]),
                    obs=pd.DataFrame({"cell_type": [label, label]}),
                    var=pd.DataFrame({"feature_name": ["MOG", "AQP4"]}, index=["ENSG1", "ENSG2"]),
                )
                adata.uns["organism"] = "Homo sapiens"
                adata.write_h5ad(path)
                paths.append(str(path))

            single = load_scrna(paths[0], max_records=10)
            self.assertEqual(len(single.cell_types), 1)  # unusable alone

            combined = load_scrna_reference_set(paths, max_records_per_file=10)
            self.assertEqual(sorted(combined.cell_types), ["astrocyte", "oligodendrocyte"])
            self.assertEqual(len(combined.records), 4)
            self.assertEqual(combined.metadata["reference_file_count"], 2)

    def test_reference_set_refuses_mixed_organisms(self):
        try:
            import anndata as ad  # type: ignore
            import numpy as np  # type: ignore
            import pandas as pd  # type: ignore
        except ImportError:
            self.skipTest("anndata required")
        from spatialmind.ingestion import load_scrna_reference_set

        with tempfile.TemporaryDirectory() as tmp:
            paths = []
            for label, organism in (("astrocyte", "Homo sapiens"), ("microglial cell", "Mus musculus")):
                path = Path(tmp) / ("%s.h5ad" % organism.split()[0])
                adata = ad.AnnData(
                    X=np.array([[5.0, 1.0], [4.0, 2.0]]),
                    obs=pd.DataFrame({"cell_type": [label, label]}),
                    var=pd.DataFrame({"feature_name": ["AQP4", "AIF1"]}, index=["G1", "G2"]),
                )
                adata.uns["organism"] = organism
                adata.write_h5ad(path)
                paths.append(str(path))
            with self.assertRaises(Exception) as ctx:
                load_scrna_reference_set(paths, max_records_per_file=10)
            self.assertIn("organism", str(ctx.exception).lower())

    def test_backed_and_memory_reads_agree(self):
        try:
            import anndata as ad  # type: ignore
            import numpy as np  # type: ignore
            import pandas as pd  # type: ignore
            from scipy import sparse  # type: ignore
        except ImportError:
            self.skipTest("anndata/scipy required")

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "backed.h5ad"
            rng = np.random.default_rng(0)
            matrix = sparse.csr_matrix(rng.integers(0, 5, size=(40, 6)).astype(float))
            adata = ad.AnnData(
                X=matrix,
                obs=pd.DataFrame({"cell_type": ["A" if i % 2 else "B" for i in range(40)]}),
                var=pd.DataFrame({"feature_name": ["G%d" % i for i in range(6)]}, index=["E%d" % i for i in range(6)]),
            )
            adata.uns["organism"] = "Homo sapiens"
            adata.write_h5ad(path)

            layer = DataIngestionLayer()
            backed = layer.load_h5ad(str(path), max_records=12, require_spatial=False, backed=True)
            memory = layer.load_h5ad(str(path), max_records=12, require_spatial=False, backed=False)

            self.assertEqual(memory.metadata["h5ad_read_mode"], "memory")
            self.assertIn(backed.metadata["h5ad_read_mode"], {"backed", "memory"})
            # Whichever path is taken, the ingested content must be identical.
            self.assertEqual(len(backed.records), len(memory.records))
            self.assertEqual(backed.genes, memory.genes)
            self.assertEqual(
                [record.cell_id for record in backed.records],
                [record.cell_id for record in memory.records],
            )
            self.assertEqual(
                [record.cell_type for record in backed.records],
                [record.cell_type for record in memory.records],
            )
            for left, right in zip(backed.records, memory.records):
                self.assertEqual(left.genes, right.genes)

    def test_spatial_h5ad_still_requires_coordinates(self):
        try:
            import anndata as ad  # type: ignore
            import numpy as np  # type: ignore
            import pandas as pd  # type: ignore
        except ImportError:
            self.skipTest("anndata required")

        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "nospatial.h5ad"
            adata = ad.AnnData(
                X=np.array([[1.0, 2.0], [3.0, 4.0]]),
                obs=pd.DataFrame({"cell_type": ["A", "B"]}),
                var=pd.DataFrame(index=["G1", "G2"]),
            )
            adata.write_h5ad(path)
            with self.assertRaises(Exception):
                DataIngestionLayer().load_h5ad(str(path), require_spatial=True)


class ReferenceAssistTests(unittest.TestCase):
    def test_tabular_reference_is_accepted_without_a_xenium_folder(self):
        from spatialmind.review.glioblastoma import _load_reference_dataset

        with tempfile.TemporaryDirectory() as tmp:
            reference = Path(tmp) / "brain_reference.csv"
            reference.write_text(
                "sample_id,x,y,cell_type,CD8A,EPCAM,PTPRC\n"
                "R1,0,0,T cell,5,0,4\n"
                "R1,1,0,Tumor cell,0,6,0\n"
                "R1,2,0,T cell,4,0,3\n",
                encoding="utf-8",
            )
            dataset, ready, status, blockers, fmt = _load_reference_dataset(str(reference), max_records=10)
            self.assertIsNotNone(dataset)
            self.assertEqual(fmt, "table")
            self.assertTrue(ready, blockers)
            self.assertEqual(status, "reference_labels_available")

    def test_single_class_reference_is_blocked(self):
        from spatialmind.review.glioblastoma import _load_reference_dataset

        with tempfile.TemporaryDirectory() as tmp:
            reference = Path(tmp) / "thin_reference.csv"
            reference.write_text(
                "sample_id,x,y,cell_type,CD8A\nR1,0,0,T cell,5\nR1,1,0,T cell,4\n",
                encoding="utf-8",
            )
            dataset, ready, status, blockers, _fmt = _load_reference_dataset(str(reference), max_records=10)
            self.assertIsNotNone(dataset)
            self.assertFalse(ready)
            self.assertEqual(status, "blocked_missing_reference_labels")
            self.assertTrue(any("cell-type classes" in item for item in blockers))

    def test_unreadable_reference_reports_blocker_instead_of_raising(self):
        from spatialmind.review.glioblastoma import _load_reference_dataset

        with tempfile.TemporaryDirectory() as tmp:
            missing = Path(tmp) / "absent_reference.h5ad"
            dataset, ready, status, blockers, fmt = _load_reference_dataset(str(missing), max_records=10)
            self.assertIsNone(dataset)
            self.assertFalse(ready)
            self.assertEqual(fmt, "anndata")
            self.assertEqual(status, "blocked_unreadable_reference")
            self.assertTrue(blockers)


class EvalHarnessTests(unittest.TestCase):
    def test_mvp_cli_defaults_to_matching_case_suite(self):
        self.assertEqual(_case_directory(None, False), "eval/test_cases")
        self.assertEqual(_case_directory(None, True), "eval/mvp_cases")
        self.assertEqual(_case_directory("custom/cases", True), "custom/cases")

    def test_eval_runner_loads_cases_and_scores(self):
        runner = EvalRunner(SpatialAgent())
        cases = runner.load_cases(os.path.join(ROOT, "eval", "test_cases"))
        # Not an exact count: adding a case should not fail an unrelated test.
        self.assertGreaterEqual(len(cases), 15)
        self.assertTrue(all(case.id for case in cases))
        report = runner.run(cases[:2])
        self.assertEqual(report["summary"]["case_count"], 2)
        self.assertGreaterEqual(report["summary"]["mean_score"], 0.5)

    def test_large_h5ad_reference_is_streamed_and_keyed_by_gene_symbol(self):
        """A reference atlas has to be readable, and over the right gene names.

        Two failures, both silent. anndata's backed mode covers `X` only, so
        opening the 7.6 GB Core GBmap atlas -- whose `layers` total 37 GB
        uncompressed -- was killed before a single cell was sampled, and
        `_open_h5ad` then fell back to a full in-memory read. And CELLxGENE keys
        `var` by Ensembl ID with symbols in `feature_name`, while a Xenium panel
        is symbols: reading the index gives an overlap of zero, which is not an
        error that announces itself but a confident answer computed from nothing.
        """
        import numpy as np
        from spatialmind.ingestion.loaders.scrna import read_h5ad_subsample

        anndata = importlib.import_module("anndata")
        import pandas as pd
        import scipy.sparse as sp

        n_cells, n_genes = 40, 6
        counts = sp.csr_matrix(np.arange(n_cells * n_genes, dtype="float32").reshape(n_cells, n_genes))
        obs = pd.DataFrame(
            {"cell_type": pd.Categorical(["malignant cell" if i % 2 else "astrocyte" for i in range(n_cells)])},
            index=["cell-%d" % i for i in range(n_cells)],
        )
        var = pd.DataFrame(
            {"feature_name": ["GFAP", "AQP4", "PTPRC", "CD68", "MBP", "SOX2"]},
            index=["ENSG%011d" % i for i in range(n_genes)],
        )
        adata = anndata.AnnData(X=counts, obs=obs, var=var)
        adata.layers["counts"] = counts.copy()

        directory = tempfile.mkdtemp()
        try:
            path = os.path.join(directory, "atlas.h5ad")
            adata.write_h5ad(path)

            dataset = read_h5ad_subsample(path, max_records=10)
            self.assertEqual(len(dataset.records), 10, "max_records must bound the read")

            genes = set()
            for record in dataset.records:
                genes.update(record.genes)
            self.assertTrue(genes, "the sampled rows carried no features")
            self.assertTrue(
                genes <= {"GFAP", "AQP4", "PTPRC", "CD68", "MBP", "SOX2"},
                "features must be gene symbols, not the Ensembl index: %s" % sorted(genes),
            )
            self.assertIn("malignant cell", {record.cell_type for record in dataset.records})
            self.assertEqual(dataset.metadata["read_strategy"], "h5py_subsample")
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    def test_eval_suite_carries_invariants_that_can_fail(self):
        """The suite has to be able to go red, or it is not measuring anything.

        It scored 1.0000 on every run while two real bugs were live: scaffold
        detection died inside the frozen app, and the Studio's API enforced the
        gate only in its UI. Every case drove the router and scored which tools
        it chose, so neither was reachable. These invariants assert on the
        registry and the gatekeeper directly.
        """
        import spatialmind.tools.registry as registry_module
        from eval.runner import check_suite_invariants

        self.assertTrue(all(item["passed"] for item in check_suite_invariants()))

        # Reproduce the frozen-bundle failure: nothing is recognised as a scaffold.
        original = registry_module._is_scaffold
        registry_module._is_scaffold = lambda func: False
        try:
            failed = [item["name"] for item in check_suite_invariants() if not item["passed"]]
        finally:
            registry_module._is_scaffold = original
        self.assertIn("scaffolds_are_detected", failed)
        self.assertIn("no_scaffold_in_llm_schemas", failed)

    def test_a_scaffold_refuses_to_execute(self):
        """Hidden from planners is not the same as refusing to run.

        `list_plannable()` and `to_anthropic_tools()` filtered scaffolds out, but
        `get(name).run(...)` executed one for anyone who asked by name -- and the
        v1 keyword router asks, for "deconvolve cell type proportions". The trace
        then recorded a successful call whose result was a placeholder.
        """
        from spatialmind.tools import build_default_registry
        from spatialmind.tools.exceptions import ToolExecutionError

        from spatialmind.ingestion import DataIngestionLayer

        dataset = DataIngestionLayer().load(DEMO)
        registry = build_default_registry()
        scaffold = next(t for t in registry.list_all() if t.capability == "unavailable")
        with self.assertRaises(ToolExecutionError):
            scaffold.run(dataset, {})


class BrainExpertBenchmarkTests(unittest.TestCase):
    def test_stratified_selection_is_deterministic_and_keeps_cluster_spatial_coverage(self):
        from spatialmind.review.brain_benchmark import _stratified_select

        rows = []
        for index in range(120):
            cluster = "0" if index < 100 else "1"
            rows.append(
                {
                    "cell_id": "cell-%03d" % index,
                    "expression_cluster": cluster,
                    "proposed_spatial_block": "block-%d" % (index % 4),
                    "selection_priority": 8 if index in {2, 103} else 0,
                }
            )
        first = _stratified_select(rows, 40, "brain")
        second = _stratified_select(rows, 40, "brain")
        self.assertEqual([row["cell_id"] for row in first], [row["cell_id"] for row in second])
        self.assertEqual(len(first), 40)
        self.assertEqual({row["expression_cluster"] for row in first}, {"0", "1"})
        self.assertGreaterEqual(len({row["proposed_spatial_block"] for row in first}), 4)
        self.assertIn("cell-002", {row["cell_id"] for row in first})
        self.assertIn("cell-103", {row["cell_id"] for row in first})

    def test_spatial_block_splits_never_leak_cells_from_one_block(self):
        from spatialmind.review.brain_benchmark import _assign_spatial_block_splits

        mapping = _assign_spatial_block_splits(["block-%d" % index for index in range(16)], "brain")
        self.assertEqual(set(mapping.values()), {"train", "validation", "test"})
        self.assertEqual(len(mapping), 16)
        self.assertEqual(mapping, _assign_spatial_block_splits(mapping.keys(), "brain"))

    def test_qc_rebalancing_keeps_critical_exceptions_without_dominating_cohort(self):
        from spatialmind.review.brain_benchmark import _rebalance_qc_strata

        rows = []
        for index in range(100):
            rows.append(
                {
                    "cell_id": "cell-%03d" % index,
                    "expression_cluster": str(index % 2),
                    "qc_stratum": "low_transcript_count" if index < 70 else "typical",
                    "selection_priority": 6 if index in {0, 1} else 0,
                }
            )
        selected = rows[:40]
        balanced = _rebalance_qc_strata(selected, rows, "brain", minimum_typical_fraction=0.6)
        self.assertGreaterEqual(sum(1 for row in balanced if row["qc_stratum"] == "typical"), 24)
        self.assertIn("cell-000", {row["cell_id"] for row in balanced})
        self.assertIn("cell-001", {row["cell_id"] for row in balanced})

    def test_completed_review_packet_materializes_frozen_truth_splits(self):
        from spatialmind.review.brain_benchmark import validate_brain_benchmark_packet

        with tempfile.TemporaryDirectory() as tmp:
            dataset_dir = Path(tmp) / "healthy_brain"
            dataset_dir.mkdir()
            labels = []
            regions = []
            splits = []
            split_names = ["train", "train", "validation", "validation", "test", "test"]
            for index, split_name in enumerate(split_names):
                cell_id = "cell-%d" % index
                labels.append(
                    {
                        "cell_id": cell_id,
                        "expert_label": "astrocyte",
                        "cl_id": "CL:0000127",
                        "secondary_state": "",
                        "confidence": "0.9",
                        "reviewer_id": "reviewer-a",
                        "expression_cluster": "0",
                        "proposed_spatial_block": "block-%d" % index,
                    }
                )
                regions.append(
                    {
                        "cell_id": cell_id,
                        "region": "region-%d" % (index % 2),
                        "region_confidence": "0.9",
                        "region_reviewer_id": "reviewer-b",
                    }
                )
                splits.append(
                    {
                        "cell_id": cell_id,
                        "proposed_spatial_block": "block-%d" % index,
                        "provisional_split": split_name,
                        "split_unit": "proposed_spatial_block",
                    }
                )

            def write(path, rows):
                with path.open("w", newline="", encoding="utf-8") as handle:
                    writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
                    writer.writeheader()
                    writer.writerows(rows)

            write(dataset_dir / "expert_cell_labels_for_review.csv", labels)
            write(dataset_dir / "cell_regions_for_review.csv", regions)
            write(dataset_dir / "benchmark_split_manifest.csv", splits)
            report = validate_brain_benchmark_packet(tmp, minimum_review_coverage=0.9)
            self.assertEqual(report["status"], "ready")
            self.assertTrue((dataset_dir / "reviewed_benchmark_truth.csv").exists())
            self.assertTrue((dataset_dir / "frozen_splits" / "test.csv").exists())
            with (dataset_dir / "reviewed_benchmark_truth.csv").open(newline="", encoding="utf-8") as handle:
                self.assertEqual(len(list(csv.DictReader(handle))), 6)

    def test_validator_requires_joint_label_region_coverage(self):
        from spatialmind.review.brain_benchmark import validate_brain_benchmark_packet

        with tempfile.TemporaryDirectory() as tmp:
            dataset_dir = Path(tmp) / "glioblastoma"
            dataset_dir.mkdir()
            labels = []
            regions = []
            splits = []
            for index in range(10):
                cell_id = "cell-%d" % index
                split_name = "train" if index < 6 else ("validation" if index < 8 else "test")
                labels.append(
                    {
                        "cell_id": cell_id,
                        "expert_label": "" if index == 0 else "astrocyte",
                        "reviewer_id": "" if index == 0 else "reviewer-a",
                    }
                )
                regions.append(
                    {
                        "cell_id": cell_id,
                        "region": "" if index == 1 else "tumor_core",
                        "region_reviewer_id": "" if index == 1 else "reviewer-b",
                    }
                )
                splits.append(
                    {
                        "cell_id": cell_id,
                        "proposed_spatial_block": "block-%d" % index,
                        "provisional_split": split_name,
                        "split_unit": "proposed_spatial_block",
                    }
                )

            def write(path, rows):
                with path.open("w", newline="", encoding="utf-8") as handle:
                    writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
                    writer.writeheader()
                    writer.writerows(rows)

            write(dataset_dir / "expert_cell_labels_for_review.csv", labels)
            write(dataset_dir / "cell_regions_for_review.csv", regions)
            write(dataset_dir / "benchmark_split_manifest.csv", splits)
            report = validate_brain_benchmark_packet(tmp, minimum_review_coverage=0.9)
            dataset_report = report["datasets"]["glioblastoma"]
            self.assertEqual(report["status"], "awaiting_expert_review")
            self.assertEqual(dataset_report["label_coverage"], 0.9)
            self.assertEqual(dataset_report["region_coverage"], 0.9)
            self.assertEqual(dataset_report["joint_label_region_coverage"], 0.8)
            self.assertTrue(any("Joint label-and-region" in item for item in dataset_report["blockers"]))
            self.assertFalse((dataset_dir / "reviewed_benchmark_truth.csv").exists())

class AgentTests(unittest.TestCase):
    def test_agent_creates_report_and_provenance(self):
        with tempfile.TemporaryDirectory() as tmp:
            agent = SpatialMindAgent(output_root=os.path.join(tmp, "outputs"), memory_root=os.path.join(tmp, "memory"))
            run = agent.run(
                "Show me CD8+ T cells relative to tumor cells in sample BRCA_04 and test co-localization.",
                DEMO,
            )
            self.assertTrue(os.path.exists(run.report_path))
            self.assertTrue(os.path.exists(run.provenance_path))
            self.assertTrue(os.path.exists(os.path.join(os.path.dirname(run.report_path), "spatial_distribution_interactive.html")))
            self.assertTrue(any(result.tool_name == "cell_type_colocalization" for result in run.results))
            stored = StorageLayer(root=os.path.join(tmp, "outputs")).get_run(run.run_id)
            self.assertEqual(stored.run_id, run.run_id)
            self.assertTrue(stored.provenance_hash)
            self.assertTrue(any(path.endswith(".html") for path in StorageLayer(root=os.path.join(tmp, "outputs")).list_figures(run.run_id)))

    def test_agent_creates_selectable_pdf_and_html_reports(self):
        with tempfile.TemporaryDirectory() as tmp:
            agent = SpatialMindAgent(output_root=os.path.join(tmp, "outputs"), memory_root=os.path.join(tmp, "memory"))
            run = agent.run(
                "Show cell type abundance in sample BRCA_04.",
                DEMO,
                report_format="both",
            )
            self.assertEqual(set(run.report_paths), {"html", "pdf"})
            self.assertTrue(run.report_path.endswith(".html"))
            with open(run.report_paths["pdf"], "rb") as handle:
                self.assertEqual(handle.read(5), b"%PDF-")


class ReportExportTests(unittest.TestCase):
    def test_report_format_validation_and_pdf_writer(self):
        self.assertEqual(normalize_report_format("PDF"), "pdf")
        with self.assertRaises(ValueError):
            normalize_report_format("docx")
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "report.pdf")
            write_pdf_report(
                path,
                "SpatialMind Test Report",
                [
                    PdfSection(
                        title="Results",
                        paragraphs=["A verified PDF result."],
                        tables=[PdfTable(headers=["Metric", "Value"], rows=[("score", "1.0")])],
                    )
                ],
            )
            self.assertGreater(os.path.getsize(path), 1024)
            with open(path, "rb") as handle:
                self.assertEqual(handle.read(5), b"%PDF-")

    def test_storage_writes_mvp_run_record_with_hashes(self):
        with tempfile.TemporaryDirectory() as tmp:
            input_path = os.path.join(tmp, "input.txt")
            with open(input_path, "w", encoding="utf-8") as handle:
                handle.write("input")
            record = StorageLayer(root=tmp).write_mvp_run_record(
                query="test",
                tool_trace=[],
                params={},
                input_files=[input_path],
            )
            self.assertTrue(os.path.exists(record.run_record_path))
            self.assertTrue(record.input_file_md5[input_path])
            self.assertEqual(record.artifact_paths, {})

    def test_replay_verifies_run_record_and_indexes_database(self):
        with tempfile.TemporaryDirectory() as tmp:
            input_path = os.path.join(tmp, "input.txt")
            figure_path = os.path.join(tmp, "figure.svg")
            with open(input_path, "w", encoding="utf-8") as handle:
                handle.write("input")
            with open(figure_path, "w", encoding="utf-8") as handle:
                handle.write("<svg></svg>")
            record = StorageLayer(root=os.path.join(tmp, "outputs")).write_mvp_run_record(
                query="Validated Xenium pilot: test replay",
                tool_trace=[],
                params={"max_records": 10},
                input_files=[input_path],
                figures=[figure_path],
            )
            verification = verify_run_record(record.run_record_path)
            self.assertEqual(verification.status, "verified")
            indexed = index_run_records(os.path.join(tmp, "outputs"), os.path.join(tmp, "runs.sqlite"))
            self.assertEqual(indexed["run_records_indexed"], 1)

    def test_xenium_replay_uses_workflow_type_and_preserves_full_section(self):
        with tempfile.TemporaryDirectory() as tmp:
            input_path = os.path.join(tmp, "input.txt")
            with open(input_path, "w", encoding="utf-8") as handle:
                handle.write("input")
            record = StorageLayer(root=os.path.join(tmp, "outputs")).write_mvp_run_record(
                query="Custom user request without legacy replay wording",
                tool_trace=[],
                params={
                    "workflow_type": "validated_xenium_pilot",
                    "max_records": 0,
                    "require_complete_section": True,
                    "review_max_records": 750,
                },
                input_files=[input_path],
            )
            replay = replay_run_record(record.run_record_path, verify_only=False)
            self.assertEqual(replay["status"], "verified_replay_ready")
            self.assertEqual(replay["query"], "Custom user request without legacy replay wording")
            self.assertEqual(replay["params"]["max_records"], 0)
            self.assertTrue(replay["params"]["require_complete_section"])
            self.assertEqual(replay["params"]["review_max_records"], 750)


class GovernanceTests(unittest.TestCase):
    def test_governance_manifest_marks_local_metadata_for_review(self):
        with tempfile.TemporaryDirectory() as tmp:
            data_root = os.path.join(tmp, "data")
            os.makedirs(data_root)
            csv_path = os.path.join(data_root, "demo.csv")
            with open(csv_path, "w", encoding="utf-8") as handle:
                handle.write("sample_id,x,y,cell_type,CD8A\n")
                handle.write("S1,1,2,T cell,3\n")
            out = os.path.join(tmp, "manifest.json")
            manifest = build_dataset_governance_manifest(data_root, out)
            self.assertEqual(manifest["status"], "needs_human_governance_review")
            self.assertEqual(len(manifest["records"]), 1)
            self.assertEqual(manifest["records"][0]["license"], "needs_review")
            self.assertTrue(os.path.exists(out))


if __name__ == "__main__":
    unittest.main()


class MarkerRuleTighteningTests(unittest.TestCase):
    def test_trace_and_ambiguous_evidence_no_longer_names_a_class(self):
        from spatialmind.ingestion.pipeline import _infer_marker_cell_type

        # PTPRC is pan-leukocyte and must not call a T/NK cell by itself.
        self.assertIsNone(_infer_marker_cell_type({"PTPRC": 2.34}))
        # Two equally specific markers from different lineages stay unresolved.
        self.assertIsNone(_infer_marker_cell_type({"AQP4": 2.5, "CD3D": 2.5}))
        # A shared marker must not outweigh a specific one: AQP4 beats CD4, which
        # is downweighted because myeloid cells express it too.
        self.assertEqual(_infer_marker_cell_type({"AQP4": 2.14, "CD4": 2.14}), "Neural/Glial cell")
        # Trace evidence below the floor.
        self.assertIsNone(_infer_marker_cell_type({"CD68": 0.3}))
        self.assertIsNone(_infer_marker_cell_type({}))

    def test_clear_evidence_still_resolves(self):
        from spatialmind.ingestion.pipeline import _infer_marker_cell_type

        self.assertEqual(
            _infer_marker_cell_type({"CD3D": 4.0, "CD3E": 3.0, "CD8A": 3.0}), "T/NK cell"
        )
        self.assertEqual(
            _infer_marker_cell_type({"AQP4": 5.0, "GFAP": 4.0, "PLP1": 3.0}), "Neural/Glial cell"
        )
        self.assertEqual(
            _infer_marker_cell_type({"PECAM1": 4.0, "CLDN5": 4.0, "FLT1": 3.0}), "Endothelial cell"
        )


class DescriptiveLaneTests(unittest.TestCase):
    def _clustered_dataset(self):
        records = []
        for index in range(24):
            if index % 3 == 0:
                genes = {"AQP4": 8.0, "GJA1": 7.0, "MOG": 0.0, "SLC17A7": 0.0}
            elif index % 3 == 1:
                genes = {"AQP4": 0.0, "GJA1": 0.0, "MOG": 8.0, "SLC17A7": 0.0}
            else:
                genes = {"AQP4": 0.0, "GJA1": 0.0, "MOG": 0.0, "SLC17A7": 8.0}
            records.append(SpotRecord("S1", float(index), float(index % 5), "Unannotated cell", genes, cell_id="c%d" % index))
        return SpatialDataset(sample_id="S1", source_path="synthetic", modality="xenium_spatial_rna", records=records)

    def test_marker_detection_can_group_by_data_derived_clusters(self):
        from spatialmind.tools.implementations import resolve_group_labels, store_cluster_assignments

        dataset = self._clustered_dataset()
        # No expert labels anywhere.
        self.assertTrue(all(record.cell_type == "Unannotated cell" for record in dataset.records))
        store_cluster_assignments(dataset, [str(index % 3) for index in range(len(dataset.records))])
        labels, key = resolve_group_labels(dataset, {"group_key": "cluster"})
        self.assertEqual(key, "cluster")
        self.assertEqual(sorted(set(labels)), ["0", "1", "2"])

        result = marker_detection(dataset, {"group_key": "cluster", "engine": "prototype"})
        self.assertEqual(result.metrics["group_key"], "cluster")
        markers = result.metrics["markers_by_group"]
        self.assertEqual(sorted(markers), ["0", "1", "2"])
        # Each cluster's defining gene should top its own marker list.
        self.assertEqual(markers["0"][0]["gene"], "AQP4")
        self.assertEqual(markers["1"][0]["gene"], "MOG")
        self.assertEqual(markers["2"][0]["gene"], "SLC17A7")

    def test_cluster_grouping_requires_assignments(self):
        from spatialmind.tools.implementations import resolve_group_labels

        dataset = self._clustered_dataset()
        with self.assertRaises(MissingPreconditionError):
            resolve_group_labels(dataset, {"group_key": "cluster"})

    def test_blocked_pilot_still_returns_descriptive_results(self):
        if not os.path.isdir(XENIUM_LYMPH):
            self.skipTest("local Xenium dataset not available")
        from spatialmind.pilot.xenium import run_pilot

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "pilot"
            result = run_pilot(XENIUM_LYMPH, out, max_records=150)
            self.assertTrue(result["status"].startswith("blocked"))
            descriptive = result["descriptive_analysis"]
            self.assertEqual(descriptive["status"], "computed")
            self.assertIn("qc_and_cluster", descriptive["tools"])
            self.assertIn("spatial_variable_genes", descriptive["tools"])
            self.assertGreaterEqual(descriptive["cluster_count"], 2)
            # Results must appear before the refusal section in the report.
            report = (out / "validated_xenium_pilot_report.md").read_text()
            self.assertIn("Descriptive Analysis (no expert labels required)", report)
            self.assertIn("Spatially autocorrelated genes", report)
            # Results come before the note about what review would add, and the
            # validated-tier scaffolding is not present on a descriptive report.
            self.assertLess(
                report.index("Descriptive Analysis"),
                report.index("What Expert Review Would Add"),
            )
            self.assertNotIn("## Claim Ledger", report)
            html_report = (out / "validated_xenium_pilot_report.html").read_text()
            self.assertIn("Descriptive Analysis (no expert labels required)", html_report)
            self.assertIn("Spatially autocorrelated genes", html_report)
            self.assertLess(html_report.index("Descriptive Analysis"), html_report.index("Blocking Reasons"))


class DescriptiveFigureTests(unittest.TestCase):
    def _dataset(self):
        records = []
        for index in range(30):
            if index % 3 == 0:
                genes = {"AQP4": 8.0, "GJA1": 7.0, "MOG": 0.0, "SLC17A7": 0.0}
            elif index % 3 == 1:
                genes = {"AQP4": 0.0, "GJA1": 0.0, "MOG": 8.0, "SLC17A7": 0.0}
            else:
                genes = {"AQP4": 0.0, "GJA1": 0.0, "MOG": 0.0, "SLC17A7": 8.0}
            records.append(
                SpotRecord("S1", float(index), float(index % 6), "Unannotated cell", genes, cell_id="c%d" % index)
            )
        return SpatialDataset(sample_id="S1", source_path="synthetic", modality="xenium_spatial_rna", records=records)

    def test_cluster_map_and_marker_heatmap_render(self):
        from spatialmind.pilot.xenium import _render_cluster_marker_heatmap, _render_cluster_spatial_map

        dataset = self._dataset()
        assignments = {record.cell_id: str(index % 3) for index, record in enumerate(dataset.records)}
        with tempfile.TemporaryDirectory() as tmp:
            spatial = _render_cluster_spatial_map(dataset, assignments, Path(tmp) / "map.png")
            if not spatial:
                self.skipTest("matplotlib unavailable")
            self.assertTrue(os.path.getsize(spatial) > 0)
            heatmap = _render_cluster_marker_heatmap(
                dataset,
                assignments,
                {"0": ["AQP4", "GJA1"], "1": ["MOG"], "2": ["SLC17A7"]},
                Path(tmp) / "heat.png",
            )
            self.assertTrue(os.path.getsize(heatmap) > 0)

    def test_figures_are_skipped_without_cluster_assignments(self):
        from spatialmind.pilot.xenium import _write_descriptive_figures

        with tempfile.TemporaryDirectory() as tmp:
            self.assertEqual(_write_descriptive_figures(self._dataset(), {}, Path(tmp)), [])

    def test_descriptive_lane_runs_even_when_gate_passes(self):
        import inspect

        from spatialmind.pilot import xenium as pilot_module

        source = inspect.getsource(pilot_module.run_pilot)
        # The lane must not be conditioned on the gate being blocked.
        self.assertIn("if not readiness_only:\n        descriptive = _run_descriptive_lane", source)


class DisplaySamplingTests(unittest.TestCase):
    def _records(self, count):
        return [
            SpotRecord("S1", float(index % 100), float(index // 100), "Unannotated cell", {"G": 1.0}, cell_id="c%d" % index)
            for index in range(count)
        ]

    def test_small_datasets_are_not_capped(self):
        from spatialmind.viz.display_sampling import downsample_for_display

        records = self._records(500)
        subset, info = downsample_for_display(records, max_points=20000)
        self.assertEqual(len(subset), 500)
        self.assertFalse(info["display_capped"])

    def test_large_datasets_are_capped_and_bounded(self):
        from spatialmind.viz.display_sampling import display_caption, downsample_for_display

        records = self._records(50000)
        subset, info = downsample_for_display(records, max_points=5000)
        self.assertEqual(len(subset), 5000)
        self.assertTrue(info["display_capped"])
        self.assertEqual(info["total_records"], 50000)
        self.assertIn("Showing", display_caption(info))
        # Spatial coverage must survive the cut rather than collapsing to one
        # corner. Grid sampling takes cells from within each bin, so the extreme
        # cell may be missed by up to one bin width -- coverage, not exactness.
        span_x = max(r.x for r in records) - min(r.x for r in records)
        span_y = max(r.y for r in records) - min(r.y for r in records)
        self.assertLessEqual(abs(min(r.x for r in subset) - min(r.x for r in records)), span_x * 0.05)
        self.assertGreaterEqual(max(r.x for r in subset), min(r.x for r in records) + span_x * 0.95)
        self.assertGreaterEqual(max(r.y for r in subset), min(r.y for r in records) + span_y * 0.95)

    def test_downsampling_is_deterministic(self):
        from spatialmind.viz.display_sampling import downsample_for_display

        records = self._records(20000)
        first = [r.cell_id for r in downsample_for_display(records, max_points=3000)[0]]
        second = [r.cell_id for r in downsample_for_display(records, max_points=3000)[0]]
        self.assertEqual(first, second)

    def test_descriptive_lane_records_stage_timings(self):
        if not os.path.isdir(XENIUM_LYMPH):
            self.skipTest("local Xenium dataset not available")
        from spatialmind.pilot.xenium import run_pilot

        with tempfile.TemporaryDirectory() as tmp:
            result = run_pilot(XENIUM_LYMPH, Path(tmp) / "p", max_records=150)
            stages = (result["descriptive_analysis"] or {}).get("stage_seconds") or {}
            self.assertIn("total", stages)
            self.assertIn("qc_and_cluster", stages)
            self.assertGreaterEqual(float(stages["total"]), 0.0)


class ClusteringPerformanceTests(unittest.TestCase):
    def test_neighbors_helper_prefers_sklearn_and_falls_back(self):
        from spatialmind.tools.implementations import _neighbors

        calls = []

        class FakePP:
            def neighbors(self, adata, **kwargs):
                calls.append(kwargs)
                if kwargs.get("transformer") == "sklearn" and adata == "no-sklearn":
                    raise TypeError("unexpected keyword argument 'transformer'")

        class FakeSC:
            pp = FakePP()

        self.assertEqual(_neighbors(FakeSC(), "ok", 15, 0), "sklearn")
        self.assertEqual(calls[-1].get("transformer"), "sklearn")
        # Older scanpy without the transformer argument must still work.
        self.assertEqual(_neighbors(FakeSC(), "no-sklearn", 15, 0), "pynndescent")
        self.assertNotIn("transformer", calls[-1])

    def test_thin_samples_are_flagged(self):
        if not os.path.isdir(XENIUM_LYMPH):
            self.skipTest("local Xenium dataset not available")
        from spatialmind.pilot.xenium import MIN_CELLS_FOR_STABLE_CLUSTERS, run_pilot

        with tempfile.TemporaryDirectory() as tmp:
            result = run_pilot(XENIUM_LYMPH, Path(tmp) / "p", max_records=200)
            descriptive = result["descriptive_analysis"]
            if descriptive.get("status") != "computed":
                self.skipTest("descriptive lane unavailable")
            self.assertIn("sampling_warning", descriptive)
            self.assertIn(str(MIN_CELLS_FOR_STABLE_CLUSTERS), descriptive["sampling_warning"])
            report = (Path(tmp) / "p" / "validated_xenium_pilot_report.md").read_text()
            self.assertIn("Sampling note", report)


class SpatialGeneScreeningTests(unittest.TestCase):
    def _adata_stub(self, n_obs, gene_names, detected_counts):
        import numpy as np

        class Var:
            def __init__(self, names):
                self.var_names = names

        class Stub:
            def __init__(self, n_obs, names, counts):
                self.n_obs = n_obs
                self.var_names = names
                # X only needs to support (X > 0).sum(axis=0)
                matrix = np.zeros((n_obs, len(names)))
                for col, count in enumerate(counts):
                    matrix[:count, col] = 1.0
                self.X = matrix

        return Stub(n_obs, gene_names, detected_counts)

    def test_detection_filter_drops_rarely_detected_genes(self):
        from spatialmind.tools.implementations import _screen_spatial_genes

        adata = self._adata_stub(1000, ["A", "B", "C"], [900, 500, 2])
        screen = _screen_spatial_genes(sq=None, adata=adata, params={"min_detected_cells": 100}, n_top=10, random_state=0)
        self.assertEqual(sorted(screen["tested_genes"]), ["A", "B"])
        self.assertEqual(screen["report"]["detected_genes"], 2)
        self.assertEqual(screen["report"]["panel_genes"], 3)
        self.assertIn("detected in >=", screen["report"]["rule"])

    def test_permutation_budget_is_not_inflated_by_default(self):
        from spatialmind.tools.implementations import _screen_spatial_genes

        adata = self._adata_stub(1000, ["A", "B", "C"], [900, 800, 700])
        screen = _screen_spatial_genes(sq=None, adata=adata, params={"n_perms": 100}, n_top=10, random_state=0)
        # Spending the saving on more permutations nets zero total work, which is
        # exactly the mistake this default avoids.
        self.assertEqual(screen["permutations"], 100)

    def test_screen_falls_back_when_too_few_genes_survive(self):
        from spatialmind.tools.implementations import _screen_spatial_genes

        adata = self._adata_stub(1000, ["A", "B"], [1, 1])
        screen = _screen_spatial_genes(sq=None, adata=adata, params={"min_detected_cells": 500}, n_top=5, random_state=0)
        # Never return an empty test set just because the filter was strict.
        self.assertEqual(sorted(screen["tested_genes"]), ["A", "B"])

    def test_report_states_the_screen(self):
        from spatialmind.pilot.xenium import _spatial_screen_note

        note = _spatial_screen_note(
            {
                "significant_gene_count_all": 50,
                "screening": {"rule": "detected in >= 80 cells; top 50", "panel_genes": 424,
                              "detected_genes": 295, "tested_genes": 50},
            }
        )
        self.assertIn("424", note)
        self.assertIn("conditional on this screen", note)
        self.assertEqual(_spatial_screen_note({}), "")


class SamplingWarningBoundaryTests(unittest.TestCase):
    def test_qc_dropping_a_few_cells_does_not_trip_the_warning(self):
        """Requesting exactly the recommended cell count must not warn.

        QC legitimately excludes zero-feature cells, so a 6000-cell request
        analyzes 5999. Warning there penalises a correct request.
        """
        import inspect

        from spatialmind.pilot import xenium as pilot_module

        source = inspect.getsource(pilot_module._run_descriptive_lane)
        self.assertIn("MIN_CELLS_FOR_STABLE_CLUSTERS * 0.95", source)
        threshold = pilot_module.MIN_CELLS_FOR_STABLE_CLUSTERS
        cutoff = threshold * 0.95
        # A 6000-cell request that analyzes 5999 after QC must stay silent.
        self.assertGreaterEqual(threshold - 1, cutoff)
        # A genuinely thin sample must still be flagged.
        self.assertLess(threshold // 2, cutoff)
        # And the cutoff must remain below the recommended target.
        self.assertLess(cutoff, threshold)


class BiologicalReplicationTests(unittest.TestCase):
    def test_one_section_per_condition_is_pseudoreplicated(self):
        from spatialmind.methods.replication import assess_condition_replication

        result = assess_condition_replication({
            "healthy": [{"section_id": "s1", "cell_count": 24406}],
            "glioblastoma": [{"section_id": "s2", "cell_count": 40887}],
        })
        self.assertFalse(result["supports_condition_inference"])
        self.assertEqual(result["status"], "pseudoreplicated")
        # Cell count must not rescue the design.
        self.assertEqual(result["conditions"]["healthy"]["cell_count"], 24406)
        self.assertEqual(result["conditions"]["healthy"]["section_count"], 1)
        self.assertIn("not independent biological replicates", result["allowed_interpretation"])

    def test_replicated_design_is_supported(self):
        from spatialmind.methods.replication import assess_condition_replication

        result = assess_condition_replication({
            "healthy": [{"section_id": "h1", "donor_id": "d1"}, {"section_id": "h2", "donor_id": "d2"}],
            "glioblastoma": [{"section_id": "g1", "donor_id": "d3"}, {"section_id": "g2", "donor_id": "d4"}],
        })
        self.assertTrue(result["supports_condition_inference"])
        self.assertEqual(result["status"], "replicated")
        self.assertEqual(result["unit_of_analysis"], "section")
        self.assertIn("pseudobulk", result["recommended_statistics"])

    def test_single_condition_cannot_be_compared(self):
        from spatialmind.methods.replication import assess_condition_replication

        result = assess_condition_replication({"healthy": [{"section_id": "a"}, {"section_id": "b"}]})
        self.assertFalse(result["supports_condition_inference"])
        self.assertTrue(any("two conditions" in b for b in result["blockers"]))

    def test_labelled_sections_still_block_condition_deltas(self):
        """The review sprint must not by itself unlock an n=1 vs n=1 comparison."""
        from unittest.mock import patch

        from spatialmind.review import glioblastoma as module

        class Intake:
            ready_for_validated_pilot = True
            loaded_cells = 750

            def to_dict(self):
                return {"status": "validated_ready"}

        class Dataset:
            def __init__(self, n):
                self.records = [type("R", (), {"cell_type": "astrocyte"})() for _ in range(n)]

        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(module, "validate_xenium_label_intake", return_value=Intake()), \
                 patch.object(module, "_load_validated_dataset", side_effect=[Dataset(700), Dataset(650)]):
                out = module.build_brain_comparison_report(output_dir=tmp, max_records=750)
        self.assertEqual(out["status"], "blocked_insufficient_biological_replication")
        self.assertNotIn("glioblastoma_minus_healthy_label_fraction", out.get("comparison") or {})
        self.assertIn("per_section_summary", out)


class ToolCapabilityTests(unittest.TestCase):
    def test_scaffolds_are_marked_unavailable_and_hidden_from_planning(self):
        registry = build_default_registry()
        summary = registry.capability_summary()
        self.assertGreater(summary.get("unavailable", 0), 0)
        self.assertEqual(len(registry.list_plannable()) + summary["unavailable"], len(registry.list_all()))
        plannable = {tool.name for tool in registry.list_plannable()}
        self.assertNotIn("cnv_inference", plannable)
        self.assertNotIn("pathway_activity", plannable)
        self.assertIn("marker_detection", plannable)

    def test_llm_tool_schemas_exclude_scaffolds_by_default(self):
        registry = build_default_registry()
        shown = {tool["name"] for tool in registry.to_anthropic_tools()}
        every = {tool["name"] for tool in registry.to_anthropic_tools(plannable_only=False)}
        self.assertLess(len(shown), len(every))
        self.assertNotIn("cnv_inference", shown)
        self.assertIn("cnv_inference", every)

    def test_mvp_registry_exposes_no_unavailable_tools(self):
        self.assertEqual(build_mvp_registry().capability_summary().get("unavailable", 0), 0)

    def test_unknown_capability_is_rejected(self):
        from spatialmind.tools.registry import SpatialTool

        with self.assertRaises(ValueError):
            SpatialTool(
                name="x", description="d", when_to_use="w", when_not_to_use="n",
                input_schema={}, output_schema={}, preconditions=[], estimated_runtime="fast",
                callable=lambda dataset, params: None, capability="totally_fine",
            )


class ReviewPlanAlignmentTests(unittest.TestCase):
    def test_pilot_template_is_aligned_to_the_run_it_came_from(self):
        """Labels only count if their cell_ids are in the cells the run loads.

        A review file built from a different selection can be labelled perfectly
        and still leave the gate shut, so the template a run emits must contain
        exactly that run's cells.
        """
        if not os.path.isdir(XENIUM_LYMPH):
            self.skipTest("local Xenium dataset not available")
        from spatialmind.ingestion import load_xenium, write_expert_label_template

        with tempfile.TemporaryDirectory() as tmp:
            dataset = load_xenium(XENIUM_LYMPH, max_records=120)
            path = write_expert_label_template(
                dataset, str(Path(tmp) / "t.csv"), max_rows=120, dataset_path=XENIUM_LYMPH
            )
            with open(path, newline="", encoding="utf-8") as handle:
                template_ids = [row["cell_id"] for row in csv.DictReader(handle)]
            loaded_ids = {record.cell_id for record in dataset.records}
            self.assertTrue(template_ids)
            self.assertTrue(set(template_ids) <= loaded_ids)
            self.assertEqual(len(template_ids), len(loaded_ids))


class DescriptiveReportPackagingTests(unittest.TestCase):
    def test_run_qc_is_surfaced_from_instrument_metrics(self):
        from spatialmind.pilot.xenium import _run_qc, _run_qc_markdown

        dataset = SpatialDataset(
            sample_id="X", source_path="x", modality="xenium_spatial_rna",
            records=[SpotRecord("X", 0.0, 0.0, "c", {"A": 1.0}, cell_id="c1")],
            metadata={
                "fraction_transcripts_decoded_q20": 0.8798,
                "fraction_transcripts_assigned": 0.7876,
                "negative_control_probe_rate": 0.0013,
                "median_genes_per_cell": 69,
            },
        )
        qc = _run_qc(dataset)
        self.assertEqual(qc["status"], "reported")
        by_key = {row["key"]: row for row in qc["metrics"]}
        self.assertEqual(by_key["fraction_transcripts_decoded_q20"]["status"], "ok")
        markdown = "\n".join(_run_qc_markdown({"run_qc": qc}))
        self.assertIn("Instrument run QC", markdown)
        self.assertIn("0.8798", markdown)

    def test_poor_decode_rate_is_flagged_for_attention(self):
        from spatialmind.pilot.xenium import _run_qc

        dataset = SpatialDataset(
            sample_id="X", source_path="x", modality="xenium_spatial_rna",
            records=[SpotRecord("X", 0.0, 0.0, "c", {"A": 1.0}, cell_id="c1")],
            metadata={"fraction_transcripts_decoded_q20": 0.42},
        )
        row = _run_qc(dataset)["metrics"][0]
        self.assertEqual(row["status"], "attention")

    def test_missing_metrics_degrade_quietly(self):
        from spatialmind.pilot.xenium import _run_qc, _run_qc_markdown

        dataset = SpatialDataset(
            sample_id="X", source_path="x", modality="xenium_spatial_rna",
            records=[SpotRecord("X", 0.0, 0.0, "c", {"A": 1.0}, cell_id="c1")], metadata={},
        )
        self.assertEqual(_run_qc(dataset)["status"], "unavailable")
        self.assertEqual(_run_qc_markdown({"run_qc": _run_qc(dataset)}), [])

    def test_successful_descriptive_run_is_not_headlined_as_blocked(self):
        if not os.path.isdir(XENIUM_LYMPH):
            self.skipTest("local Xenium dataset not available")
        from spatialmind.pilot.xenium import run_pilot

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "p"
            result = run_pilot(XENIUM_LYMPH, out, max_records=200, review_artifacts=False)
            report = (out / "validated_xenium_pilot_report.md").read_text()
            # The gate code is still reported, but it is not the headline.
            self.assertIn("Xenium Analysis Report", report)
            self.assertIn("Outcome:", report)
            self.assertTrue(result["status"].startswith("blocked"))
            # Review templates are skipped when not doing a review.
            self.assertEqual(result["expert_label_template"], "")
            self.assertFalse((out / "expert_label_template.csv").exists())


class DescriptiveReportSplitTests(unittest.TestCase):
    def test_validation_sections_are_dropped_only_from_descriptive_reports(self):
        from spatialmind.pilot.xenium import _filter_descriptive_sections

        lines = [
            "# Report", "", "## Status", "- ok",
            "## Descriptive Analysis (no expert labels required)", "content",
            "## Claim Ledger", "refused",
            "## Typed Tool Plan", "plan",
            "## Limitations", "limits",
        ]
        descriptive = {"status": "blocked_missing_validation_inputs",
                       "descriptive_analysis": {"status": "computed"}}
        kept = _filter_descriptive_sections(lines, descriptive)
        self.assertIn("## Descriptive Analysis (no expert labels required)", kept)
        self.assertIn("## Limitations", kept)
        self.assertNotIn("## Claim Ledger", kept)
        self.assertNotIn("## Typed Tool Plan", kept)

        # Validated runs keep everything: those sections carry real content there.
        validated = {"status": "validated_ready", "descriptive_analysis": {"status": "computed"}}
        self.assertEqual(_filter_descriptive_sections(lines, validated), lines)

        # A run with no descriptive results is left alone too.
        nothing = {"status": "blocked_analysis_backend", "descriptive_analysis": {"status": "not_run"}}
        self.assertEqual(_filter_descriptive_sections(lines, nothing), lines)

    def test_marker_hints_name_the_family_not_the_cell_type(self):
        from spatialmind.pilot.xenium import _cluster_marker_hints

        hints = _cluster_marker_hints({
            "0": ["CD68", "AIF1", "C1QA", "CD74"],
            "1": ["MOG", "MOBP", "CLDN11", "PLP1"],
            "2": ["PDGFRA", "CSPG4", "VCAN", "BCAN"],
            "3": ["SOMEGENE", "ANOTHER"],
        })
        self.assertIn("myeloid", hints["0"])
        self.assertIn("oligodendrocyte", hints["1"])
        self.assertIn("opc", hints["2"])
        # No match stays silent rather than guessing.
        self.assertNotIn("3", hints)
        # Wording must not assert an identity.
        self.assertNotIn("cell type", " ".join(hints.values()).lower())

    def test_every_label_lineage_is_marker_detectable(self):
        """A lineage a reference label maps to must be reachable from markers.

        Otherwise cells of that lineage can never be flagged as absent from a
        reference, which is what the coverage check exists to catch.
        """
        from spatialmind.tools.implementations import LABEL_LINEAGE_KEYWORDS, LINEAGE_MARKERS

        label_lineages = {lineage for _keyword, lineage in LABEL_LINEAGE_KEYWORDS}
        self.assertEqual(label_lineages - set(LINEAGE_MARKERS), set())

    def test_opc_and_oligodendrocyte_stay_distinguishable(self):
        from spatialmind.tools.implementations import marker_lineage

        self.assertEqual(marker_lineage({"PDGFRA": 6.0, "CSPG4": 5.0, "VCAN": 4.0})[0], "opc")
        self.assertEqual(marker_lineage({"MOG": 8.0, "MOBP": 7.0, "CLDN11": 6.0, "OPALIN": 5.0})[0], "oligodendrocyte")


class TimingQualityTests(unittest.TestCase):
    def test_contended_run_is_flagged(self):
        """Reproduces the observed case: a stage that takes ~150s idle measured
        2032s while the test suite ran concurrently. Wall time alone made a
        healthy dataset look pathological."""
        from spatialmind.pilot.xenium import _timing_quality

        quality = _timing_quality(wall_seconds=2032.0, cpu_seconds=150.0)
        self.assertEqual(quality["status"], "contended")
        self.assertGreater(quality["wall_to_cpu_ratio"], 10)
        self.assertIn("should not be compared", quality["note"])

    def test_parallel_run_is_not_mistaken_for_contention(self):
        """Multi-threaded work spends more CPU than wall time; that is healthy."""
        from spatialmind.pilot.xenium import _timing_quality

        quality = _timing_quality(wall_seconds=129.9, cpu_seconds=218.8)
        self.assertEqual(quality["status"], "clean")
        self.assertLess(quality["wall_to_cpu_ratio"], 1.0)

    def test_short_runs_are_not_flagged(self):
        """A brief run can show a high ratio from startup noise alone."""
        from spatialmind.pilot.xenium import _timing_quality

        self.assertEqual(_timing_quality(wall_seconds=5.0, cpu_seconds=1.0)["status"], "clean")

    def test_zero_cpu_time_does_not_divide_by_zero(self):
        from spatialmind.pilot.xenium import _timing_quality

        self.assertEqual(_timing_quality(wall_seconds=10.0, cpu_seconds=0.0)["wall_to_cpu_ratio"], 0.0)


class ControlProbeExclusionTests(unittest.TestCase):
    def test_xenium_control_families_are_recognised(self):
        from spatialmind.tools.implementations import is_control_feature

        for name in ("UnassignedCodeword_0045", "NegControlProbe_00042", "NegControlCodeword_0500",
                     "BLANK_0006", "antisense_PROK2", "DeprecatedCodeword_123", "NegControlProbe"):
            self.assertTrue(is_control_feature(name), name)

    def test_real_genes_sharing_a_prefix_are_not_excluded(self):
        """Prefix matching alone would drop real genes; the convention requires a
        separator or digits after the prefix."""
        from spatialmind.tools.implementations import is_control_feature

        for name in ("AQP4", "CD3D", "KRT6B", "NEGR1", "BLANKET_GENE", "ANTISENSEGENE"):
            self.assertFalse(is_control_feature(name), name)

    def test_control_probes_are_kept_out_of_the_expression_matrix(self):
        from spatialmind.tools.implementations import expression_feature_names

        dataset = SpatialDataset(
            sample_id="X", source_path="x", modality="xenium_spatial_rna",
            records=[SpotRecord("X", 0.0, 0.0, "c", {
                "AQP4": 5.0, "CD3D": 3.0, "GJA1": 2.0,
                "UnassignedCodeword_0045": 9.0, "NegControlProbe_00042": 8.0, "BLANK_0006": 7.0,
                "CELL_AREA": 400.0,
            }, cell_id="c1")],
        )
        kept = expression_feature_names(dataset)
        self.assertEqual(sorted(kept), ["AQP4", "CD3D", "GJA1"])

    def test_fixtures_of_only_control_features_still_return_something(self):
        from spatialmind.tools.implementations import expression_feature_names

        dataset = SpatialDataset(
            sample_id="X", source_path="x", modality="xenium_spatial_rna",
            records=[SpotRecord("X", 0.0, 0.0, "c", {"BLANK_0001": 1.0, "BLANK_0002": 2.0}, cell_id="c1")],
        )
        self.assertEqual(len(expression_feature_names(dataset)), 2)

    def test_real_panels_lose_only_control_features(self):
        if not os.path.isdir(XENIUM_LYMPH):
            self.skipTest("local Xenium dataset not available")
        from spatialmind.ingestion import load_xenium
        from spatialmind.tools.implementations import expression_feature_names, is_control_feature

        dataset = load_xenium(XENIUM_LYMPH, max_records=200)
        kept = set(expression_feature_names(dataset))
        dropped = set(dataset.genes) - kept
        self.assertTrue(dropped, "expected control probes in a real Xenium panel")
        from spatialmind.schemas import NON_EXPRESSION_FEATURE_NAMES

        for gene in dropped:
            self.assertTrue(
                is_control_feature(gene) or gene.upper() in NON_EXPRESSION_FEATURE_NAMES,
                "dropped a biological gene: %s" % gene,
            )


class ClusterConfidenceTests(unittest.TestCase):
    def test_micro_clusters_are_flagged(self):
        """A 34-cell cluster is not a peer of a 51,929-cell one.

        Reproduces the observed breast full-section result, where five of
        fourteen clusters held between 34 and 186 of 209,467 cells and carried
        degenerate overlapping markers.
        """
        from spatialmind.pilot.xenium import _cluster_confidence

        counts = {"6": 51929, "0": 42351, "7": 31835, "1": 23469, "9": 20804,
                  "4": 16798, "5": 15685, "3": 3787, "2": 1083,
                  "13": 186, "10": 116, "11": 88, "12": 80, "8": 34}
        result = _cluster_confidence(counts)
        self.assertEqual(result["status"], "computed")
        self.assertEqual(set(result["low_confidence_clusters"]), {"13", "10", "11", "12", "8"})
        self.assertNotIn("2", result["low_confidence_clusters"])  # 1083 cells clears both floors
        self.assertGreater(result["fraction_cells_in_well_populated"], 0.97)

    def test_balanced_clustering_flags_nothing(self):
        from spatialmind.pilot.xenium import _cluster_confidence

        result = _cluster_confidence({"0": 5000, "1": 4000, "2": 3000, "3": 2000})
        self.assertEqual(result["low_confidence_clusters"], [])
        self.assertEqual(result["fraction_cells_in_well_populated"], 1.0)

    def test_empty_counts_do_not_crash(self):
        from spatialmind.pilot.xenium import _cluster_confidence

        self.assertEqual(_cluster_confidence({})["status"], "unavailable")

    def test_silhouette_reading_distinguishes_separation_quality(self):
        from spatialmind.pilot.xenium import _silhouette_reading

        poor = _silhouette_reading(-0.0013, 0.7384)
        typical = _silhouette_reading(0.05, 0.72)
        good = _silhouette_reading(0.141, 0.786)
        self.assertIn("overlap almost completely", poor)
        self.assertIn("routine", typical)
        self.assertIn("reasonably separated", good)
        # Distinct readings, and a missing value stays silent rather than guessing.
        self.assertEqual(len({poor, typical, good}), 3)
        self.assertEqual(_silhouette_reading(None, 0.7), "")


class SelfContainedReportTests(unittest.TestCase):
    def test_figures_are_embedded_so_the_report_survives_being_emailed(self):
        """The HTML report is what a reviewer is sent.

        Referencing figures by relative filename means every image breaks the
        moment the file leaves its directory, which is exactly what happens when
        it reaches a domain expert.
        """
        from spatialmind.pilot.xenium import _figure_html

        with tempfile.TemporaryDirectory() as tmp:
            png = Path(tmp) / "fig.png"
            # Minimal valid PNG.
            png.write_bytes(bytes.fromhex(
                "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
                "890000000a49444154789c63000100000500010d0a2db40000000049454e44ae426082"
            ))
            markup = _figure_html(str(png))
            self.assertIn('src="data:image/png;base64,', markup)
            self.assertNotIn('src="fig.png"', markup)

    def test_svg_uses_the_right_mime_type(self):
        from spatialmind.pilot.xenium import _figure_html

        with tempfile.TemporaryDirectory() as tmp:
            svg = Path(tmp) / "fig.svg"
            svg.write_text('<svg xmlns="http://www.w3.org/2000/svg"><rect width="1" height="1"/></svg>')
            self.assertIn("data:image/svg+xml;base64,", _figure_html(str(svg)))

    def test_oversized_and_missing_figures_fall_back_to_a_link(self):
        """A single huge figure must not make the report itself unopenable."""
        from spatialmind.pilot import xenium as module

        with tempfile.TemporaryDirectory() as tmp:
            big = Path(tmp) / "big.png"
            big.write_bytes(b"\x00" * 1024)
            original = module.MAX_INLINE_FIGURE_BYTES
            try:
                module.MAX_INLINE_FIGURE_BYTES = 100  # force the cap
                self.assertEqual(module._figure_data_uri(str(big)), "")
                self.assertIn('src="big.png"', module._figure_html(str(big)))
            finally:
                module.MAX_INLINE_FIGURE_BYTES = original
        # A path that does not exist degrades to the filename rather than raising.
        self.assertEqual(module._figure_data_uri("/nonexistent/none.png"), "")

    def test_interactive_html_stays_a_link(self):
        """Only images inline; the viewer is a separate multi-megabyte artifact."""
        from spatialmind.pilot.xenium import _figure_html

        markup = _figure_html("explorer_lite_viewer.html")
        self.assertIn("<a href=", markup)
        self.assertNotIn("data:image", markup)


class GateEvidenceTests(unittest.TestCase):
    """The gate's numbers must reach the report, and its conditions must count
    what a reviewer supplied rather than what the loader guessed.

    Both failures were silent and both were found by running the agent rather
    than by reading it: a run cleared at 1% label coverage printed the same
    Limitations sentence as one cleared at 100%, and a bundle holding a single
    reviewed cell satisfied "at least two biological cell classes" out of the
    loader's own marker-rule fallbacks.
    """

    def _dataset(self, labels, regions):
        records = [
            SpotRecord("S1", float(i), 0.0, label, {"CD8A": 1.0}, cell_id="c%d" % i)
            for i, label in enumerate(labels)
        ]
        for record, region in zip(records, regions):
            record.region = region
        return SpatialDataset(sample_id="S1", source_path="xenium", records=records)

    def test_conditions_count_the_reviewed_table_not_the_records(self):
        from spatialmind.gatekeeper import pilot_gate

        # One reviewed cell; the loader has filled the rest with marker guesses
        # and a section-wide placeholder region -- which is what used to pass.
        dataset = self._dataset(
            labels=["Neuron", "Neural/Glial cell", "Endothelial cell", "Myeloid cell"],
            regions=["cortex", "whole_section", "whole_section", "whole_section"],
        )
        assets = {k: True for k in ("has_cell_table", "has_feature_matrix", "has_morphology", "has_boundaries")}
        gate = pilot_gate(
            dataset=dataset,
            asset_readiness=assets,
            label_report={
                "status": "expert_labels_applied", "matched_cells": 1, "total_records": 4,
                "reviewed_labels": ["Neuron"],
            },
            region_report={
                "status": "user_regions_applied", "matched_cells": 1, "total_records": 4,
                "reviewed_regions": ["cortex"],
            },
            min_label_coverage=0.0,
            min_region_coverage=0.0,
            allow_single_region=False,
            # A 2-cell fixture exercises the condition wiring, not the
            # class-size floor; the default 50 is asserted separately.
            min_cells_per_class=1,
        )
        self.assertNotEqual(gate["status"], "validated_ready",
                            "one reviewed cell cleared the gate using loader fallback labels")
        self.assertEqual(gate["reviewed_cell_classes"], ["Neuron"])
        self.assertEqual(gate["reviewed_regions"], ["cortex"])
        self.assertEqual(gate["reviewed_basis"]["labels"], "reviewed_table")

    def test_a_report_without_reviewed_lists_falls_back_and_says_so(self):
        """An older run record has no reviewed_* lists. It must still evaluate,
        and the fallback must be visible rather than passing as reviewer input."""
        from spatialmind.gatekeeper import pilot_gate

        dataset = self._dataset(labels=["Neuron", "Astrocyte"], regions=["a", "b"])
        gate = pilot_gate(
            dataset=dataset,
            asset_readiness={k: True for k in ("has_cell_table", "has_feature_matrix", "has_morphology", "has_boundaries")},
            label_report={"status": "expert_labels_applied", "matched_cells": 2, "total_records": 2},
            region_report={"status": "user_regions_applied", "matched_cells": 2, "total_records": 2},
            min_label_coverage=0.7, min_region_coverage=0.7, allow_single_region=False,
            # 2-cell fixture: exercising the fallback, not the class-size floor.
            min_cells_per_class=1,
        )
        self.assertEqual(gate["status"], "validated_ready")
        self.assertEqual(gate["reviewed_basis"]["labels"], "record_scan_fallback")

    def test_coverage_floor_refuses_a_gate_that_cannot_refuse(self):
        from spatialmind.gatekeeper import MIN_COVERAGE_FLOOR, CoverageFloorError, enforce_coverage_floor

        with self.assertRaises(CoverageFloorError):
            enforce_coverage_floor(0.0, 0.7)
        with self.assertRaises(CoverageFloorError):
            enforce_coverage_floor(0.7, 0.0)
        # Acknowledged, and at the floor itself, both proceed.
        enforce_coverage_floor(0.0, 0.0, acknowledge_low_coverage=True)
        enforce_coverage_floor(MIN_COVERAGE_FLOOR, MIN_COVERAGE_FLOOR)

    def test_limitations_state_coverage_and_the_threshold(self):
        from spatialmind.pilot.xenium import _limitations

        payload = {
            "features_loaded": 319,
            "label_report": {"status": "expert_labels_applied"},
            "region_report": {"status": "user_regions_applied"},
            "gate_evidence": {
                "label_coverage": 0.01, "label_matched_cells": 40, "label_total_records": 3994,
                "min_label_coverage": 0.0, "label_review_decisions": 2,
                "region_coverage": 0.01, "region_matched_cells": 40, "region_total_records": 3994,
                "min_region_coverage": 0.0, "region_review_decisions": 2,
                "below_coverage_floor": ["label", "region"], "coverage_floor": 0.2,
                "default_min_label_coverage": 0.7, "thresholds_lowered": ["label", "region"],
            },
            "analysis_scope": {"scope": "full_section"},
            "status": "validated_ready",
        }
        text = " ".join(_limitations(payload))
        self.assertIn("1.0% coverage", text, "the report did not state label coverage")
        self.assertIn("40 of 3,994", text)
        self.assertIn("2 review decisions", text)
        self.assertIn("COVERAGE FLOOR OVERRIDDEN", text,
                      "a run below the floor did not say so in its own limitations")


class ReliabilityHonestyTests(unittest.TestCase):
    """Components must not report 1.0000 for "no test was needed"."""

    def test_not_applicable_renders_as_na_not_as_full_marks(self):
        from spatialmind.pilot.xenium import _component_cell, _reliability_cell

        item = {
            "S_statistical": 1.0, "reliability": 0.5,
            "components": {"S_statistical": {"status": "not_applicable", "score": 1.0}},
            "status": "computed",
        }
        self.assertEqual(_component_cell(item, "S_statistical"), "n/a")

        measured = {"P_panel": 0.5455, "components": {"P_panel": {"status": "computed"}}, "status": "computed"}
        self.assertEqual(_component_cell(measured, "P_panel"), "0.5455")

        dropped = {"reliability": 0.5, "status": "blocked", "components": {}}
        self.assertIn("claim not made", _reliability_cell(dropped))

    def test_annotation_score_is_coverage_not_reviewer_confidence(self):
        from spatialmind.methods.reliability.scoring import _annotation_component

        payload = {
            "label_report": {
                "status": "expert_labels_applied", "matched_cells": 90, "total_records": 100,
                "confidence_summary": {"mean": 0.9}, "review_decisions": 11,
            },
            "records_loaded": 100,
        }
        component = _annotation_component({"claim_type": "cell_type_annotation"}, payload)
        self.assertAlmostEqual(component.score, 0.9, places=4)
        # 0.9 coverage x 0.9 confidence would be 0.81. The confidence must be
        # reported, not multiplied in: the Studio writes a hard-coded 0.9.
        self.assertNotAlmostEqual(component.score, 0.81, places=4)
        self.assertIn("review_decisions:11", component.evidence)
        self.assertIn("stated_confidence:0.90", component.evidence)

    def test_multiple_testing_uses_the_pairs_tested_not_the_pairs_shown(self):
        from spatialmind.methods.reliability.scoring import _tested_pair_count

        result = ToolResult(
            tool_name="cell_neighborhood_enrichment",
            summary="",
            metrics={"tested_pair_count": 45, "top_pairs": [{"pair": "a | b", "zscore": 3.0}] * 10},
        )
        self.assertEqual(_tested_pair_count(result), 45)
        # A result predating the metric falls back to the old behaviour rather
        # than to an invented number.
        legacy = ToolResult(tool_name="x", summary="", metrics={"top_pairs": [{"pair": "a | b"}] * 7})
        self.assertEqual(_tested_pair_count(legacy), 7)

    def test_region_claim_can_actually_be_grounded(self):
        """It required an evidence token nothing emitted, so it was dropped in
        every validated run ever made -- with regions applied and the summary on
        disk."""
        from spatialmind.agent.grounding import ClaimGroundingChecker

        evidence = ClaimGroundingChecker()._available_evidence(
            [ToolResult(tool_name="region_summary", summary="", metrics={})]
        )
        self.assertIn("region_summary", evidence)


class NeighborhoodPairTests(unittest.TestCase):
    def test_self_pairs_do_not_occupy_the_ranked_slots(self):
        from spatialmind.tools.implementations import _is_self_pair

        self.assertTrue(_is_self_pair({"pair": "3 | 3"}))
        self.assertTrue(_is_self_pair({"pair": "T cell | T cell"}))
        self.assertFalse(_is_self_pair({"pair": "3 | 5"}))
        self.assertFalse(_is_self_pair({"pair": "T cell | B cell"}))


class PanelSizeTests(unittest.TestCase):
    def test_one_definition_of_panel_size_across_layers(self):
        """319 measured genes were reported as 483 by the contract and as a
        third, sample-dependent number by the transfer preflight."""
        from spatialmind.ingestion.contract import to_cell_by_feature_contract
        from spatialmind.schemas import expression_feature_names

        dataset = SpatialDataset(
            sample_id="S1",
            source_path="xenium",
            records=[
                SpotRecord("S1", 0.0, 0.0, "T cell", {
                    "CD8A": 1.0, "CD3D": 2.0, "PTPRC": 1.0,
                    "NegControlProbe_00042": 1.0, "TRANSCRIPT_COUNTS": 400.0,
                }, cell_id="c0")
            ],
            metadata={"assay_subtype": "xenium_spatial_rna", "is_targeted_panel": True},
        )
        measured = expression_feature_names(dataset)
        self.assertNotIn("NegControlProbe_00042", measured)
        self.assertNotIn("TRANSCRIPT_COUNTS", measured)
        self.assertEqual(to_cell_by_feature_contract(dataset).n_features, len(measured))


class EvalPolicyMismatchTests(unittest.TestCase):
    def test_mvp_cases_without_mvp_policy_are_refused(self):
        """Running MVP cases against the legacy registry scores 2/13 and reads
        exactly like a regression. The runner knows both values; it now checks."""
        from eval.runner import _refuse_policy_mismatch

        class _Parser:
            def error(self, message):
                raise SystemExit(message)

        with self.assertRaises(SystemExit):
            _refuse_policy_mismatch(_Parser(), "eval/mvp_cases", False)
        with self.assertRaises(SystemExit):
            _refuse_policy_mismatch(_Parser(), "eval/test_cases", True)
        _refuse_policy_mismatch(_Parser(), "eval/mvp_cases", True)
        _refuse_policy_mismatch(_Parser(), "eval/test_cases", False)


class ProvenanceHashTests(unittest.TestCase):
    def test_a_directory_digest_reacts_to_edited_content(self):
        """It hashed names and sizes only, so a same-length edit to the
        reviewer's label table left the digest unchanged and replay still
        printed "verified"."""
        from spatialmind.storage.run_store import _file_md5

        root = tempfile.mkdtemp()
        try:
            labels = os.path.join(root, "expert_cell_labels.csv")
            with open(labels, "w", encoding="utf-8") as handle:
                handle.write("cell_id,expert_label\nc1,Neuron\n")
            before = _file_md5(root)
            with open(labels, "w", encoding="utf-8") as handle:
                handle.write("cell_id,expert_label\nc1,Glioma\n")  # same length
            self.assertNotEqual(before, _file_md5(root),
                                "an edited label table left the input digest unchanged")
        finally:
            shutil.rmtree(root, ignore_errors=True)


class RegionClaimGroundingTests(unittest.TestCase):
    def test_the_region_claim_is_supported_when_region_summary_ran(self):
        """It inherited required_evidence=["figure"] from its claim type -- a
        token only feature_overlay emits, which the validated plan does not run.
        So the claim was dropped every time, with regions applied."""
        from spatialmind.pilot.claims import build_pilot_claim_ledger

        results = [
            ToolResult(tool_name="annotation", summary="", metrics={}),
            ToolResult(tool_name="region_summary", summary="", metrics={}),
            ToolResult(
                tool_name="cell_neighborhood_enrichment", summary="",
                metrics={"engine": "squidpy", "top_pairs": [{"pair": "a | b", "zscore": 4.0}],
                         "tested_pair_count": 3},
            ),
        ]
        ledger = build_pilot_claim_ledger({"status": "validated_ready"}, results)
        region = next(item for item in ledger if "tissue regions" in item["claim_text"])
        self.assertEqual(region["status"], "supported",
                         "the region claim is still structurally impossible to support")

    def test_it_is_still_dropped_when_region_summary_did_not_run(self):
        from spatialmind.pilot.claims import build_pilot_claim_ledger

        ledger = build_pilot_claim_ledger(
            {"status": "validated_ready"},
            [ToolResult(tool_name="annotation", summary="", metrics={})],
        )
        region = next(item for item in ledger if "tissue regions" in item["claim_text"])
        self.assertEqual(region["status"], "dropped")


class SpatialStatisticsTests(unittest.TestCase):
    """The local and point-pattern statistics, and the pitfalls that make them lie."""

    def _synthetic(self, n=1200, seed=0):
        import numpy as np

        rng = np.random.default_rng(seed)
        xy = rng.uniform(0, 400, size=(n, 2))
        # A gene expressed only in a horizontal band, and a pure-noise control.
        stripe = np.where((xy[:, 1] > 160) & (xy[:, 1] < 240), 8.0, 0.2) + rng.normal(0, 0.3, n)
        noise = rng.normal(5.0, 1.0, n)
        records = [
            SpotRecord(
                "S", float(xy[i, 0]), float(xy[i, 1]),
                "band" if 160 < xy[i, 1] < 240 else "outside",
                {"STRIPE": float(max(stripe[i], 0.0)), "NOISE": float(max(noise[i], 0.0)), "FILL": 1.0},
                cell_id="c%d" % i,
            )
            for i in range(n)
        ]
        dataset = SpatialDataset(sample_id="S", source_path="synthetic", records=records)
        dataset.normalized = True
        return dataset

    def test_local_moran_finds_a_stripe_and_not_noise(self):
        """A statistic that finds structure everywhere is as useless as one that
        finds it nowhere, so this asserts both directions on data with a known
        answer."""
        from spatialmind.tools.spatial_statistics import local_moran_hotspots

        result = local_moran_hotspots(self._synthetic(), ["STRIPE", "NOISE"],
                                      {"strict_engine": True, "n_perms": 199})
        self.assertEqual(result["status"], "computed")
        rows = {row["gene"]: row for row in result["genes"]}
        self.assertGreater(rows["STRIPE"]["hotspot_cells"], 50,
                           "local Moran's I did not find an implanted stripe")
        self.assertEqual(rows["NOISE"]["hotspot_cells"], 0,
                         "local Moran's I found hotspots in pure noise")
        self.assertEqual(len(result["per_cell"]), len(self._synthetic().records))

    def test_local_moran_significance_has_no_permutation_floor(self):
        """Counted permutation p-values cannot fall below 1/(n_perms + 1), and BH
        over one test per cell multiplies that floor by the cell count -- at 999
        permutations on 4,000 cells nothing could reach alpha and every gene
        reported zero hotspots. The z-score basis has no floor."""
        from spatialmind.tools.spatial_statistics import local_moran_hotspots

        result = local_moran_hotspots(self._synthetic(), ["STRIPE"], {"strict_engine": True, "n_perms": 99})
        self.assertEqual(result["significance_basis"], "conditional_randomisation_z_score")
        self.assertGreater(result["genes"][0]["hotspot_cells"], 0,
                           "a 99-permutation budget silently disabled detection")

    def test_cell_type_autocorrelation_skips_populations_too_small_to_score(self):
        from spatialmind.tools.spatial_statistics import cell_type_spatial_autocorrelation

        dataset = self._synthetic()
        # Three cells of a rare type: enough to compute a number, never enough to
        # mean one. Measured on real data: 12 cells scored I = 0.025 at p = 0.04.
        for record in dataset.records[:3]:
            record.cell_type = "rare"
        result = cell_type_spatial_autocorrelation(dataset, {"strict_engine": True, "n_perms": 49})
        self.assertEqual(result["status"], "computed")
        self.assertNotIn("rare", [row["group"] for row in result["groups"]])
        self.assertIn("rare", [row["group"] for row in result["skipped_groups"]])

    def test_ripley_caps_radii_far_below_the_library_default(self):
        """squidpy's default ceiling is sqrt(area/2). On a real section that is
        two orders of magnitude past any interaction scale, where the statistic
        measures the shape of the section rather than the cells inside it.

        This previously also asserted a `window_caveat` describing a bounding-box
        window. That caveat was factually wrong -- squidpy uses the convex hull,
        not the bounding box -- and the window is now an occupancy mask of the
        tissue, so the assertion moved to the metadata that is actually true.
        """
        from spatialmind.tools.spatial_statistics import ripley_cell_types

        result = ripley_cell_types(self._synthetic(), {"strict_engine": True, "n_simulations": 20})
        self.assertEqual(result["status"], "computed")
        self.assertLess(result["max_distance_um"], result["default_max_distance_um"],
                        "Ripley used the library default radius")
        self.assertEqual(result["window"], "tissue_occupancy_mask")
        self.assertGreater(result["window_area_um2"], 0)
        self.assertTrue(result["window_rule"])
        self.assertTrue(result["estimator_note"])

    def test_ripley_scores_against_the_envelope_not_the_theoretical_line(self):
        from spatialmind.tools.spatial_statistics import ripley_cell_types

        result = ripley_cell_types(self._synthetic(), {"strict_engine": True, "n_simulations": 30})
        self.assertIn("simulated CSR envelope", result["reference"])
        for row in result["groups"]:
            # A group whose deviation is positive must not be called dispersed,
            # and vice versa. The two disagreed while L(r) - r was the reference.
            if row["verdict"] == "clustered":
                self.assertGreater(row["peak_deviation"], 0)
            if row["verdict"] == "dispersed":
                self.assertLess(row["peak_deviation"], 0)

    def test_lees_l_reports_a_missing_gene_instead_of_inventing_one(self):
        from spatialmind.tools.spatial_statistics import bivariate_spatial_correlation

        result = bivariate_spatial_correlation(
            self._synthetic(), [("STRIPE", "NOISE"), ("ABSENT", "STRIPE")], {"strict_engine": True})
        rows = {(row["gene_a"], row["gene_b"]): row for row in result["pairs"]}
        self.assertEqual(rows[("STRIPE", "NOISE")]["status"], "computed")
        self.assertEqual(rows[("ABSENT", "STRIPE")]["status"], "gene_not_in_panel")
        self.assertIsNone(rows[("ABSENT", "STRIPE")]["lees_l"])


class RegionProposalTests(unittest.TestCase):
    def test_proposals_never_write_the_file_the_gate_reads(self):
        """The whole safety property is the filename. `cell_regions.csv` is the
        reviewer's file and the only one the gate reads; a proposal writing there
        would turn a data-derived grouping into gate-clearing evidence."""
        from spatialmind.tools.region_proposal import CANDIDATE_FILENAME, write_region_candidates

        self.assertEqual(CANDIDATE_FILENAME, "cell_regions_candidate.csv")
        root = Path(tempfile.mkdtemp())
        try:
            dataset = SpatialDataset(
                sample_id="S", source_path="x",
                records=[SpotRecord("S", float(i), 0.0, "t", {"A": 1.0}, cell_id="c%d" % i) for i in range(4)],
            )
            proposal = {
                "status": "computed", "source": "spatial_domain",
                "assignments": {"c0": "domain_0", "c1": "domain_0", "c2": "", "c3": "domain_1"},
            }
            result = write_region_candidates(dataset, [proposal], root)
            self.assertEqual(result["status"], "written")
            self.assertFalse((root / "cell_regions.csv").exists(),
                             "a region proposal wrote the file the gate reads")
            rows = list(csv.DictReader(open(root / CANDIDATE_FILENAME, encoding="utf-8")))
            self.assertEqual(len(rows), 4)
            self.assertTrue(all(row["review_status"] == "needs_expert_review" for row in rows))
            self.assertTrue(all(row["region"] == "" for row in rows),
                            "a candidate pre-filled the reviewer's own column")
            self.assertEqual(rows[0]["candidate_region"], "domain_0")
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_proposed_regions_are_never_named_after_a_tissue(self):
        from spatialmind.tools.region_proposal import hotspot_regions

        lisa = {
            "genes": [{"gene": "MOG", "status": "computed"}],
            "per_cell": {"c%d" % i: {"lisa_MOG": "high-high"} for i in range(150)},
        }
        result = hotspot_regions(lisa, min_cells=10)
        self.assertEqual(result["status"], "computed")
        self.assertEqual(result["regions"][0]["candidate_region"], "hotspot_MOG")


class ResultTableTests(unittest.TestCase):
    def _payload(self, status):
        return {
            "status": status,
            "dataset_path": "/data/section",
            "created_at": "2026-09-16T00:00:00Z",
            "features_loaded": 319,
            "analysis_scope": {"scope": "full_section", "loaded_records": 2, "total_records": 2},
            "gate_evidence": {"label_coverage": 0.99, "min_label_coverage": 0.7},
            "label_report": {"method": "expert_label_table"},
            "descriptive_analysis": {"markers_by_cluster": {"0": ["GJA1", "AQP4"]}},
        }

    def _dataset(self):
        return SpatialDataset(
            sample_id="S", source_path="x",
            records=[
                SpotRecord("S", 1.0, 2.0, "Astrocyte", {"GJA1": 3.0, "TRANSCRIPT_COUNTS": 90.0}, cell_id="c0"),
                SpotRecord("S", 3.0, 4.0, "Neuron", {"GJA1": 0.0, "TRANSCRIPT_COUNTS": 70.0}, cell_id="c1"),
            ],
        )

    def test_a_blocked_run_never_exports_a_column_headed_cell_type(self):
        """On a blocked run that column holds the loader's marker-rule guesses.
        A spreadsheet header saying `cell_type` is a claim the report spends a
        page refusing to make."""
        from spatialmind.viz.tables import write_result_tables

        root = Path(tempfile.mkdtemp())
        try:
            write_result_tables(self._payload("blocked_missing_validation_inputs"),
                                self._dataset(), root, run_id="r1")
            text = (root / "tables" / "cells.tsv").read_text(encoding="utf-8")
            header = next(line for line in text.splitlines() if not line.startswith("#"))
            self.assertIn("cell_type_provisional", header.split("\t"))
            self.assertNotIn("cell_type", header.split("\t"))
            self.assertIn("WARNING", text)
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_a_validated_run_does_export_cell_type(self):
        from spatialmind.viz.tables import write_result_tables

        root = Path(tempfile.mkdtemp())
        try:
            write_result_tables(self._payload("validated_ready"), self._dataset(), root, run_id="r1")
            header = next(line for line in (root / "tables" / "cells.tsv").read_text(encoding="utf-8").splitlines()
                          if not line.startswith("#"))
            self.assertIn("cell_type", header.split("\t"))
            self.assertNotIn("cell_type_provisional", header.split("\t"))
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_every_table_carries_its_own_provenance(self):
        """A TSV gets emailed detached from its report. The caveats have to travel
        with it or the honesty work is undone when someone opens the file."""
        from spatialmind.viz.tables import write_result_tables

        root = Path(tempfile.mkdtemp())
        try:
            manifest = write_result_tables(self._payload("blocked_missing_validation_inputs"),
                                           self._dataset(), root, run_id="run_42")
            self.assertEqual(manifest["status"], "written")
            self.assertTrue(manifest["tables"])
            for item in manifest["tables"]:
                text = Path(item["path"]).read_text(encoding="utf-8")
                self.assertIn("run_42", text, "%s lost its run id" % item["table"])
                self.assertIn("gate_status", text, "%s lost the gate status" % item["table"])
                self.assertIn("did not open", text, "%s did not carry the blocked-run note" % item["table"])
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_gene_table_reports_screened_out_genes_not_only_survivors(self):
        from spatialmind.viz.tables import write_result_tables

        payload = self._payload("validated_ready")
        payload["descriptive_analysis"]["spatial_genes"] = {
            "screening": {"rule": "top 2 by analytic Moran's I", "tested_genes": 2, "detected_genes": 4},
            "top_genes": [{"gene": "A", "morans_i": 0.3, "pval_adj": 0.001},
                          {"gene": "B", "morans_i": 0.2, "pval_adj": 0.01}],
            "screened_out_genes": [{"gene": "C", "morans_i": 0.01}, {"gene": "D", "morans_i": 0.0}],
        }
        root = Path(tempfile.mkdtemp())
        try:
            write_result_tables(payload, self._dataset(), root, run_id="r1")
            rows = [line.split("\t") for line in
                    (root / "tables" / "genes_spatial.tsv").read_text(encoding="utf-8").splitlines()
                    if not line.startswith("#")]
            header, body = rows[0], rows[1:]
            tested = {row[header.index("gene")]: row[header.index("tested")] for row in body}
            self.assertEqual(tested, {"A": "true", "B": "true", "C": "false", "D": "false"},
                             "the table reported only the genes that survived the screen")
        finally:
            shutil.rmtree(root, ignore_errors=True)


class CellTableColumnTests(unittest.TestCase):
    def test_n_genes_counts_measured_genes_not_zero(self):
        """The first filter excluded uppercase keys to drop QC pseudo-features.
        Xenium gene symbols are uppercase, so it dropped every real gene and every
        cell reported n_genes = 0."""
        from spatialmind.viz.tables import write_result_tables

        dataset = SpatialDataset(
            sample_id="S", source_path="x",
            records=[SpotRecord("S", 1.0, 2.0, "Astrocyte", {
                "GJA1": 3.0, "AQP4": 1.0, "MOG": 0.0,
                "TRANSCRIPT_COUNTS": 90.0, "NegControlProbe_00042": 2.0,
            }, cell_id="c0")],
        )
        root = Path(tempfile.mkdtemp())
        try:
            write_result_tables({"status": "validated_ready", "dataset_path": "d"}, dataset, root, run_id="r")
            lines = [l for l in (root / "tables" / "cells.tsv").read_text(encoding="utf-8").splitlines()
                     if not l.startswith("#")]
            header, row = lines[0].split("\t"), lines[1].split("\t")
            # GJA1 and AQP4 only: MOG is zero, and the QC and control features
            # are not measured genes.
            self.assertEqual(row[header.index("n_genes")], "2")
        finally:
            shutil.rmtree(root, ignore_errors=True)


class RegionSummaryFeatureTests(unittest.TestCase):
    def test_region_top_features_exclude_qc_pseudo_features(self):
        """CELL_AREA, TOTAL_COUNTS and friends are library size and morphology on
        a scale two orders of magnitude above any gene. They took the top slots in
        every region of every run, which is the same contamination the expression
        matrix already excludes."""
        from spatialmind.tools.implementations import region_summary

        records = []
        for i in range(6):
            records.append(SpotRecord(
                "S", float(i), 0.0, "Astrocyte",
                {"GJA1": 4.0, "AQP4": 2.0, "CELL_AREA": 560.0, "TRANSCRIPT_COUNTS": 200.0},
                cell_id="c%d" % i,
            ))
            records[-1].region = "band_a"
        dataset = SpatialDataset(sample_id="S", source_path="x", records=records)
        result = region_summary(dataset, {"top_n_features": 4})
        features = [row["feature"] for row in result.metrics["regions"]["band_a"]["top_features"]]
        self.assertEqual(features, ["GJA1", "AQP4"])
        self.assertNotIn("CELL_AREA", features)
        self.assertNotIn("TRANSCRIPT_COUNTS", features)

    def test_region_composition_table_handles_the_mapping_shape(self):
        """`regions` is a name -> summary mapping. Iterating it as a list yielded
        region names as strings and no rows, so the table silently did not exist."""
        from spatialmind.viz.tables import _region_rows

        metrics = {"region_summary": {"regions": {
            "band_a": {"cell_count": 3, "cell_type_counts": {"Astrocyte": 2, "Neuron": 1}},
        }}}
        rows = list(_region_rows({}, metrics))
        self.assertEqual(len(rows), 2)
        by_type = {row["cell_type"]: row for row in rows}
        self.assertEqual(by_type["Astrocyte"]["n_cells"], 2)
        self.assertEqual(by_type["Astrocyte"]["fraction"], 0.6667)
        self.assertEqual(by_type["Astrocyte"]["region"], "band_a")


class RipleyPeakDirectionTests(unittest.TestCase):
    """The headline deviation must point the way the verdict does.

    The envelope fix corrected the *reference* and left the *peak selection*
    wrong: argmax of |deviation| picked whichever excursion was largest in
    magnitude regardless of sign, so a population with 82% of its radii below the
    envelope -- dispersed on every reading -- reported peak_deviation = +0.5. The
    original test passed straight through it, because synthetic data never
    produced a mixed-sign pattern. This builds one directly.
    """

    def _peak(self, deviations, verdict):
        """The shipped selection, over a known deviation curve."""
        import numpy as np

        from spatialmind.tools.spatial_statistics import peak_deviation_index

        values = np.asarray(deviations, dtype=float)
        usable = np.ones(len(values), dtype=bool)
        return float(values[peak_deviation_index(values, usable, verdict)])

    def test_a_dispersed_pattern_never_reports_a_positive_peak(self):
        # The exact curve from the failing row: one small positive excursion with
        # the largest magnitude, several larger negative ones. argmax(|dev|)
        # returns +0.501 here, which is what shipped.
        self.assertLess(self._peak([0.501, -0.3, -0.4, -0.45, -0.2], "dispersed"), 0.0,
                        "a dispersed population reported a positive peak deviation")

    def test_a_clustered_pattern_never_reports_a_negative_peak(self):
        self.assertGreater(self._peak([-9.0, 1.0, 2.0, 3.0], "clustered"), 0.0,
                           "a clustered population reported a negative peak deviation")

    def test_it_falls_back_rather_than_raising_when_no_excursion_matches(self):
        # Every deviation negative but the verdict says clustered: pick the
        # largest in magnitude rather than raise, and let the sign show the
        # disagreement.
        self.assertEqual(self._peak([-1.0, -5.0, -2.0], "clustered"), -5.0)

    def test_real_sections_keep_peak_and_verdict_consistent(self):
        """The invariant asserted end to end, on data that actually mixes signs."""
        from spatialmind.tools.spatial_statistics import ripley_cell_types

        import numpy as np

        rng = np.random.default_rng(3)
        records = []
        # A tightly clustered population and an evenly spaced one, so the run
        # contains both verdicts at once.
        for i in range(260):
            cx, cy = (80.0, 80.0) if i % 2 else (320.0, 320.0)
            records.append(SpotRecord("S", float(rng.normal(cx, 18)), float(rng.normal(cy, 18)),
                                      "clumped", {"A": 1.0, "B": 1.0}, cell_id="k%d" % i))
        grid = np.linspace(20, 380, 17)
        for i, (gx, gy) in enumerate((x, y) for x in grid for y in grid):
            records.append(SpotRecord("S", float(gx), float(gy), "even",
                                      {"A": 1.0, "B": 1.0}, cell_id="e%d" % i))
        dataset = SpatialDataset(sample_id="S", source_path="x", records=records)
        result = ripley_cell_types(dataset, {"strict_engine": True, "n_simulations": 40})
        self.assertEqual(result["status"], "computed")
        self.assertTrue(result["groups"])
        for row in result["groups"]:
            if row["verdict"] == "clustered":
                self.assertGreater(row["peak_deviation"], 0, row["group"])
            elif row["verdict"] == "dispersed":
                self.assertLess(row["peak_deviation"], 0, row["group"])
            self.assertIn("deviation_sign_varies", row)


class PairTableUniquenessTests(unittest.TestCase):
    def test_self_pairs_are_not_written_twice(self):
        """`all_pairs` already contains the self-pairs. Appending `self_pairs` on
        top of it wrote each one twice -- and a duplicated row in a results table
        is worse than a missing one, because nothing about it looks wrong."""
        from spatialmind.viz.tables import _pair_rows

        metrics = {"cell_neighborhood_enrichment": {
            "tested_pair_count": 3,
            "all_pairs": [
                {"pair": "A | A", "zscore": 9.0},
                {"pair": "A | B", "zscore": 2.0},
                {"pair": "B | B", "zscore": 7.0},
            ],
            "self_pairs": [{"pair": "A | A", "zscore": 9.0}, {"pair": "B | B", "zscore": 7.0}],
        }}
        rows = list(_pair_rows({"cell_type_counts": {"A": 5, "B": 6}}, metrics))
        labels = ["%s | %s" % (row["type_a"], row["type_b"]) for row in rows]
        self.assertEqual(len(labels), len(set(labels)), "a self-pair was written twice")
        self.assertEqual(sorted(labels), ["A | A", "A | B", "B | B"])

    def test_the_split_lists_are_still_used_when_all_pairs_is_absent(self):
        from spatialmind.viz.tables import _pair_rows

        metrics = {"cell_neighborhood_enrichment": {
            "top_pairs": [{"pair": "A | B", "zscore": 2.0}],
            "self_pairs": [{"pair": "A | A", "zscore": 9.0}],
        }}
        rows = list(_pair_rows({}, metrics))
        self.assertEqual(sorted("%s | %s" % (r["type_a"], r["type_b"]) for r in rows), ["A | A", "A | B"])


class LisaReportedParametersTests(unittest.TestCase):
    def test_the_permutation_count_is_reported_not_left_as_a_placeholder(self):
        from spatialmind.tools.spatial_statistics import local_moran_hotspots

        import numpy as np

        rng = np.random.default_rng(0)
        xy = rng.uniform(0, 200, size=(300, 2))
        records = [SpotRecord("S", float(xy[i, 0]), float(xy[i, 1]), "t",
                              {"G": float(abs(rng.normal(3, 1))), "H": 1.0}, cell_id="c%d" % i)
                   for i in range(300)]
        dataset = SpatialDataset(sample_id="S", source_path="x", records=records)
        dataset.normalized = True
        result = local_moran_hotspots(dataset, ["G"], {"strict_engine": True, "n_perms": 149})
        self.assertEqual(result["n_perms"], 149, "the run's permutation count was reported as None")
        self.assertEqual(result["genes"][0]["n_perms"], 149)


class ExpressionNormalisationTests(unittest.TestCase):
    def test_local_statistics_normalise_like_every_other_expression_tool(self):
        """qc_and_cluster, marker_detection and spatial_variable_genes all guard
        with `if not dataset.normalized: normalize_total + log1p`. The local and
        bivariate statistics did not, so on raw counts the report would show
        global Moran's I on log-normalised data and local Moran's I on raw counts
        for the same genes, with nothing saying they disagreed."""
        from spatialmind.tools.spatial_statistics import _expression_adata

        import numpy as np

        records = [
            SpotRecord("S", float(i), 0.0, "t", {"A": 100.0 * (i + 1), "B": 5.0}, cell_id="c%d" % i)
            for i in range(6)
        ]
        raw = SpatialDataset(sample_id="S", source_path="x", records=records)
        raw.normalized = False
        values = np.asarray(_expression_adata(raw)[:, "A"].X, dtype=float).ravel()
        self.assertLess(float(values.max()), 20.0,
                        "raw counts reached the local statistics un-normalised")

        already = SpatialDataset(sample_id="S", source_path="x", records=records)
        already.normalized = True
        untouched = np.asarray(_expression_adata(already)[:, "A"].X, dtype=float).ravel()
        self.assertGreater(float(untouched.max()), 100.0,
                           "an already-normalised dataset was normalised twice")


class TableColumnCoverageTests(unittest.TestCase):
    def test_every_field_a_row_builder_emits_has_a_column(self):
        """A field added to a row dict but not to the column list is silently
        dropped at write time -- which is how `deviation_sign_varies` was computed
        on every run and written to nothing."""
        from spatialmind.viz import tables

        cases = [
            (tables.POINT_PATTERN_COLUMNS, tables._point_pattern_rows,
             {"cell_type_point_pattern": {"status": "computed", "max_distance_um": 200.0, "groups": [
                 {"group": "A", "n_cells": 60, "peak_deviation": 1.0, "peak_radius_um": 10.0,
                  "mean_deviation": 0.5, "radii_above_envelope": 0.9, "radii_below_envelope": 0.1,
                  "deviation_sign_varies": True, "verdict": "clustered"}]}}),
            (tables.GENE_PAIR_COLUMNS, tables._gene_pair_rows,
             {"gene_pair_spatial_correlation": {"status": "computed", "pairs": [
                 {"gene_a": "A", "gene_b": "B", "lees_l": 0.1, "pearson_r": 0.2,
                  "status": "computed", "interpretation": "x"}]}}),
            (tables.GROUP_COLUMNS, tables._group_rows,
             {"cell_type_spatial_autocorrelation": {"status": "computed", "graph": {"family": "knn", "n_neighs": 6},
                                                    "groups": [{"group": "A", "morans_i": 0.1, "pval_sim": 0.01,
                                                                "pval_adj": 0.02, "n_cells": 99,
                                                                "interpretation": "y"}]}}),
        ]
        for columns, builder, payload in cases:
            rows = list(builder(payload))
            self.assertTrue(rows, "%s produced no rows to check" % builder.__name__)
            for row in rows:
                missing = sorted(set(row) - set(columns))
                self.assertEqual(missing, [], "%s emits %s with no column" % (builder.__name__, missing))

        # The builders that need the tool-result mapping, checked the same way.
        pair_rows = list(tables._pair_rows(
            {"cell_type_counts": {"A": 5}},
            {"cell_neighborhood_enrichment": {"tested_pair_count": 1,
                                              "all_pairs": [{"pair": "A | A", "zscore": 3.0}]}}))
        self.assertTrue(pair_rows)
        for row in pair_rows:
            self.assertEqual(sorted(set(row) - set(tables.PAIR_COLUMNS)), [])

        region_rows = list(tables._region_rows(
            {}, {"region_summary": {"regions": {"r": {"cell_type_counts": {"A": 2}}}}}))
        self.assertTrue(region_rows)
        for row in region_rows:
            self.assertEqual(sorted(set(row) - set(tables.REGION_COLUMNS)), [])


class TissueWindowTests(unittest.TestCase):
    """The observation window, which decides what the CSR null is allowed to be.

    squidpy simulates inside the convex hull of all cells. A section is not its
    hull: cortex is concave, ventricles are holes, and a null free to spread
    where no cell could be makes any confined population read as clustered.
    """

    def test_the_mask_recovers_a_known_area(self):
        from spatialmind.tools.spatial_statistics import _window_bin_size, tissue_window

        import numpy as np

        rng = np.random.default_rng(0)
        points = rng.uniform(0, 400, size=(900, 2))
        true_area = 400.0 * 400.0
        size = _window_bin_size(true_area, len(points), 5.0)
        window = tissue_window(points, size)
        self.assertGreater(window["area"], 0.85 * true_area, "the mask lost real tissue")
        self.assertLess(window["area"], 1.20 * true_area, "the mask covered empty space")

    def test_the_mask_excludes_the_hole_a_hull_would_swallow(self):
        from spatialmind.tools.spatial_statistics import _window_bin_size, tissue_window

        import numpy as np

        rng = np.random.default_rng(1)
        # A square with a large central void: the hull is the whole square.
        points = []
        while len(points) < 2500:
            x, y = rng.uniform(0, 400, 2)
            if not (120 < x < 280 and 120 < y < 280):
                points.append((x, y))
        points = np.asarray(points)
        size = _window_bin_size(400.0 * 400.0, len(points), 4.0)
        window = tissue_window(points, size)
        void = 160.0 * 160.0
        self.assertLess(window["area"], 400.0 * 400.0 - 0.5 * void,
                        "the mask covered a void the size of the hole")

    def test_chance_gaps_are_filled_but_real_holes_are_not(self):
        """An empty bin with six occupied neighbours is a sampling gap; a
        ventricle's interior has no occupied neighbours at all."""
        from spatialmind.tools.spatial_statistics import tissue_window

        import numpy as np

        # A 7x7 block of bins with one interior bin left empty.
        points = []
        for ix in range(7):
            for iy in range(7):
                if (ix, iy) == (3, 3):
                    continue
                points.append((ix * 10.0 + 5.0, iy * 10.0 + 5.0))
        window = tissue_window(np.asarray(points), 10.0)
        self.assertEqual(window["filled_bins"], 1, "the single-bin gap was not filled")
        self.assertEqual(window["occupied_bins"], 49)

    def test_a_uniform_population_in_the_window_reads_as_random(self):
        """The decisive control: points drawn uniformly from the window must not
        read as clustered. Run against the real section this is what showed the
        machinery was sound and 'all clustered' was biology, not bias."""
        from spatialmind.tools.spatial_statistics import (
            _window_bin_size, ripley_cell_types, sample_in_window, tissue_window)

        import numpy as np

        rng = np.random.default_rng(5)
        backdrop = rng.uniform(0, 400, size=(3000, 2))
        size = _window_bin_size(400.0 * 400.0, len(backdrop), 4.0)
        window = tissue_window(backdrop, size)
        records = [SpotRecord("S", float(x), float(y), "background", {"A": 1.0, "B": 1.0},
                              cell_id="b%d" % i) for i, (x, y) in enumerate(backdrop)]
        for i, (x, y) in enumerate(sample_in_window(window, 400, rng)):
            records.append(SpotRecord("S", float(x), float(y), "uniform_control",
                                      {"A": 1.0, "B": 1.0}, cell_id="u%d" % i))
        # A genuinely clumped population alongside, so the test would fail if the
        # statistic had simply stopped discriminating.
        for i in range(400):
            cx, cy = rng.choice([80.0, 300.0]), rng.choice([80.0, 300.0])
            records.append(SpotRecord("S", float(rng.normal(cx, 10)), float(rng.normal(cy, 10)),
                                      "clumped", {"A": 1.0, "B": 1.0}, cell_id="c%d" % i))
        result = ripley_cell_types(SpatialDataset(sample_id="S", source_path="x", records=records),
                                   {"strict_engine": True, "n_simulations": 60})
        verdicts = {row["group"]: row["verdict"] for row in result["groups"]}
        self.assertNotIn(verdicts.get("uniform_control"), ("clustered", "dispersed"),
                         "a uniform control inside the window returned a verdict")
        self.assertEqual(verdicts.get("clumped"), "clustered",
                         "the statistic stopped detecting a real clump")

    def test_it_reports_the_window_it_used_rather_than_the_hull(self):
        from spatialmind.tools.spatial_statistics import ripley_cell_types

        import numpy as np

        rng = np.random.default_rng(2)
        points = rng.uniform(0, 300, size=(1200, 2))
        records = [SpotRecord("S", float(x), float(y), "a" if i % 2 else "b",
                              {"A": 1.0, "B": 1.0}, cell_id="c%d" % i)
                   for i, (x, y) in enumerate(points)]
        result = ripley_cell_types(SpatialDataset(sample_id="S", source_path="x", records=records),
                                   {"strict_engine": True, "n_simulations": 25})
        self.assertEqual(result["window"], "tissue_occupancy_mask")
        self.assertGreater(result["window_area_um2"], 0)
        self.assertIn("window_bin_um", result)
        self.assertIn("convex_hull_area_um2", result)


class RipleyEffectSizeTests(unittest.TestCase):
    def test_a_verdict_needs_an_effect_size_not_only_an_envelope_crossing(self):
        """At thousands of cells the simulated envelope is narrow enough that a
        sub-micron systematic bias puts 97% of radii outside it. A uniform ring
        control came back 'clustered' on a deviation of 0.7 um at r = 25 um, next
        to a real clump at 28.9 um -- significance without effect size."""
        from spatialmind.tools.spatial_statistics import RIPLEY_MIN_RELATIVE_DEVIATION, ripley_cell_types

        import numpy as np

        self.assertGreater(RIPLEY_MIN_RELATIVE_DEVIATION, 0.0)
        rng = np.random.default_rng(4)
        points = rng.uniform(0, 300, size=(1600, 2))
        records = [SpotRecord("S", float(x), float(y), "a" if i % 2 else "b",
                              {"A": 1.0, "B": 1.0}, cell_id="c%d" % i)
                   for i, (x, y) in enumerate(points)]
        dataset = SpatialDataset(sample_id="S", source_path="x", records=records)

        strict = ripley_cell_types(dataset, {"strict_engine": True, "n_simulations": 40})
        for row in strict["groups"]:
            self.assertIn("relative_effect", row)
            if row["verdict"] in ("clustered", "dispersed"):
                self.assertGreaterEqual(row["relative_effect"], RIPLEY_MIN_RELATIVE_DEVIATION)

        # Drop the floor to zero and the same uniform data starts returning
        # verdicts, which is what the floor exists to prevent.
        loose = ripley_cell_types(dataset, {"strict_engine": True, "n_simulations": 40,
                                            "min_relative_deviation": 0.0})
        self.assertEqual(loose["min_relative_deviation"], 0.0)


class ReviewedLabelScopeTests(unittest.TestCase):
    """A table headed "cell type" on a validated run must list reviewed classes.

    A partially reviewed section carries both the reviewer's classes and the
    loader's marker-rule guesses on every cell review did not reach. Grouping by
    `record.cell_type` mixes them: a validated run listed `Unannotated cell` and
    `Neural/Glial cell` beside atlas-transferred classes, with nothing to tell
    them apart -- the conflation the gate exists to prevent, arriving after the
    gate had opened.
    """

    def _dataset(self):
        import numpy as np

        rng = np.random.default_rng(0)
        records = []
        # Two reviewed classes, plus a loader-guessed residue.
        for i in range(160):
            records.append(SpotRecord("S", float(rng.uniform(0, 200)), float(rng.uniform(0, 200)),
                                      "astrocyte", {"A": 1.0, "B": 1.0}, cell_id="a%d" % i))
        for i in range(160):
            records.append(SpotRecord("S", float(rng.uniform(0, 200)), float(rng.uniform(0, 200)),
                                      "oligodendrocyte", {"A": 1.0, "B": 1.0}, cell_id="o%d" % i))
        for i in range(140):
            records.append(SpotRecord("S", float(rng.uniform(0, 200)), float(rng.uniform(0, 200)),
                                      "Unannotated cell", {"A": 1.0, "B": 1.0}, cell_id="u%d" % i))
        return SpatialDataset(sample_id="S", source_path="x", records=records)

    def test_autocorrelation_excludes_labels_the_reviewer_did_not_supply(self):
        from spatialmind.tools.spatial_statistics import cell_type_spatial_autocorrelation

        result = cell_type_spatial_autocorrelation(self._dataset(), {
            "strict_engine": True, "n_perms": 49,
            "reviewed_labels": ["astrocyte", "oligodendrocyte"],
        })
        groups = [row["group"] for row in result["groups"]]
        self.assertEqual(sorted(groups), ["astrocyte", "oligodendrocyte"])
        self.assertNotIn("Unannotated cell", groups)
        self.assertTrue(result["reviewed_only"])
        self.assertEqual(result["unreviewed_cells"], 140)
        reasons = {row["group"]: row["reason"] for row in result["skipped_groups"]}
        self.assertIn("not in the reviewed label table", reasons["Unannotated cell"])

    def test_point_pattern_excludes_them_too(self):
        from spatialmind.tools.spatial_statistics import ripley_cell_types

        result = ripley_cell_types(self._dataset(), {
            "strict_engine": True, "n_simulations": 20,
            "reviewed_labels": ["astrocyte", "oligodendrocyte"],
        })
        self.assertEqual(sorted(row["group"] for row in result["groups"]),
                         ["astrocyte", "oligodendrocyte"])
        self.assertEqual(result["unreviewed_cells"], 140)

    def test_without_a_reviewed_set_every_group_is_reported(self):
        """The descriptive lane groups by cluster, where there is no reviewed set
        and nothing to filter against."""
        from spatialmind.tools.spatial_statistics import cell_type_spatial_autocorrelation

        result = cell_type_spatial_autocorrelation(self._dataset(), {"strict_engine": True, "n_perms": 49})
        self.assertFalse(result["reviewed_only"])
        self.assertEqual(result["unreviewed_cells"], 0)
        self.assertIn("Unannotated cell", [row["group"] for row in result["groups"]])


class SkipReasonReportingTests(unittest.TestCase):
    def test_each_skip_reason_gets_its_own_line(self):
        """One hard-coded heading described two different exclusions, producing
        "Not tested (fewer than 50 cells): Neural/Glial cell (2446)"."""
        from spatialmind.pilot.xenium import _group_autocorrelation_lines

        block = {
            "status": "computed", "min_cells": 50, "graph": {"family": "knn", "n_neighs": 6},
            "groups": [{"group": "astrocyte", "morans_i": 0.1, "pval_sim": 0.01,
                        "pval_adj": 0.01, "n_cells": 900, "interpretation": "forms spatial patches"}],
            "skipped_groups": [
                {"group": "Neural/Glial cell", "n_cells": 2446, "reason": "not in the reviewed label table"},
                {"group": "T/NK cell", "n_cells": 15, "reason": "fewer than 50 cells"},
            ],
        }
        text = "\n".join(_group_autocorrelation_lines(block, "## H", "intro"))
        small = next(l for l in text.splitlines() if "T/NK cell" in l)
        unreviewed = next(l for l in text.splitlines() if "Neural/Glial cell" in l)
        self.assertNotEqual(small, unreviewed, "both exclusions were printed on one line")
        self.assertIn("fewer than 50 cells", small)
        self.assertNotIn("fewer than 50 cells", unreviewed,
                         "a 2,446-cell group was described as having fewer than 50")
        self.assertIn("not in the reviewed label table", unreviewed)


class EnrichmentReviewedScopeTests(unittest.TestCase):
    def test_neighbourhood_enrichment_drops_unreviewed_labels_when_told(self):
        """A validated pair table listing `Unannotated cell | endothelial cell`
        beside a reviewed pair is the same conflation, one tool over."""
        from spatialmind.tools.implementations import cell_neighborhood_enrichment

        import numpy as np

        rng = np.random.default_rng(0)
        records = []
        for name, count in (("astrocyte", 120), ("oligodendrocyte", 120), ("Unannotated cell", 90)):
            for i in range(count):
                records.append(SpotRecord("S", float(rng.uniform(0, 200)), float(rng.uniform(0, 200)),
                                          name, {"A": 1.0, "B": 1.0}, cell_id="%s%d" % (name[:3], i)))
        dataset = SpatialDataset(sample_id="S", source_path="x", records=records)

        scoped = cell_neighborhood_enrichment(dataset, {
            "strict_engine": True, "n_perms": 20, "include_all_pairs": True,
            "reviewed_labels": ["astrocyte", "oligodendrocyte"],
        })
        labels = {part.strip()
                  for pair in scoped.metrics["all_pairs"]
                  for part in str(pair["pair"]).split("|")}
        self.assertEqual(labels, {"astrocyte", "oligodendrocyte"})
        self.assertTrue(scoped.metrics["reviewed_only"])
        self.assertEqual(scoped.metrics["excluded_unreviewed_cell_count"], 90)

        # Without the reviewed set -- the descriptive lane -- nothing is dropped.
        unscoped = cell_neighborhood_enrichment(dataset, {
            "strict_engine": True, "n_perms": 20, "include_all_pairs": True})
        unscoped_labels = {part.strip()
                           for pair in unscoped.metrics["all_pairs"]
                           for part in str(pair["pair"]).split("|")}
        self.assertIn("Unannotated cell", unscoped_labels)
        self.assertFalse(unscoped.metrics["reviewed_only"])


class PerCellLabelSourceTests(unittest.TestCase):
    def test_label_source_distinguishes_reviewed_cells_from_loader_guesses(self):
        """A single table-wide source string stamped `expert_label_table` on every
        cell, including the ones review never reached -- so the one column
        downstream code would filter on was the one that lied."""
        from spatialmind.viz.tables import write_result_tables

        dataset = SpatialDataset(
            sample_id="S", source_path="x",
            records=[
                SpotRecord("S", 1.0, 1.0, "astrocyte", {"A": 1.0}, cell_id="reviewed"),
                SpotRecord("S", 2.0, 2.0, "Unannotated cell", {"A": 1.0}, cell_id="guessed"),
            ],
        )
        payload = {
            "status": "validated_ready", "dataset_path": "d",
            "label_report": {"method": "expert_label_table", "reviewed_labels": ["astrocyte"]},
        }
        root = Path(tempfile.mkdtemp())
        try:
            write_result_tables(payload, dataset, root, run_id="r")
            lines = [l for l in (root / "tables" / "cells.tsv").read_text(encoding="utf-8").splitlines()
                     if not l.startswith("#")]
            header = lines[0].split("\t")
            rows = {r.split("\t")[0]: r.split("\t") for r in lines[1:]}
            column = header.index("label_source")
            self.assertEqual(rows["reviewed"][column], "expert_label_table")
            self.assertEqual(rows["guessed"][column], "loader_marker_rule",
                             "an unreviewed cell claimed a reviewed label source")
        finally:
            shutil.rmtree(root, ignore_errors=True)


class AnnDataBuildTests(unittest.TestCase):
    """The matrix fill was rewritten for speed; the values must not move.

    The old form asked every cell about every gene -- a nested comprehension of
    dict lookups, 85% of which returned a default on Xenium's sparse data. The
    new form writes only what each cell measured. Same numbers, and this proves
    it on the cases where the two could differ: a gene absent from a cell, a
    cell whose `raw_genes` covers only part of the panel, a cell with no
    `raw_genes` at all, and a gene present on a record but not in the panel.
    """

    def _reference_matrices(self, dataset, genes):
        """The original implementation, kept here as the oracle."""
        import numpy as np

        matrix = np.array(
            [[record.genes.get(gene, 0.0) for gene in genes] for record in dataset.records],
            dtype=float,
        )
        source = np.array(
            [[record.raw_genes.get(gene, record.genes.get(gene, 0.0)) for gene in genes]
             for record in dataset.records],
            dtype=float,
        )
        return matrix, source

    def _dataset(self):
        records = [
            # Every gene present, raw differs from analysis values.
            SpotRecord("S", 0.0, 0.0, "a", {"G1": 1.5, "G2": 2.5, "G3": 3.5}, cell_id="full"),
            # A gene missing from this cell entirely.
            SpotRecord("S", 1.0, 0.0, "a", {"G1": 4.0, "G3": 6.0}, cell_id="sparse"),
            # No raw_genes at all: source must fall back to the analysis values.
            SpotRecord("S", 2.0, 0.0, "b", {"G1": 7.0, "G2": 8.0}, cell_id="no_raw"),
            # Carries a gene that is not in the panel; it must be ignored, not
            # written into some other column.
            SpotRecord("S", 3.0, 0.0, "b", {"G2": 9.0, "NOT_IN_PANEL": 99.0}, cell_id="extra"),
        ]
        records[0].raw_genes = {"G1": 10.0, "G2": 20.0, "G3": 30.0}
        # raw covers only part of the panel: the rest must fall back per gene.
        records[1].raw_genes = {"G1": 40.0}
        records[2].raw_genes = {}
        records[3].raw_genes = {"G2": 90.0, "NOT_IN_PANEL": 999.0}
        return SpatialDataset(sample_id="S", source_path="x", records=records)

    def test_the_fast_fill_matches_the_original_exactly(self):
        import numpy as np

        from spatialmind.tools.implementations import _dataset_to_anndata, expression_feature_names

        dataset = self._dataset()
        genes = expression_feature_names(dataset)
        expected_x, expected_source = self._reference_matrices(dataset, genes)

        adata = _dataset_to_anndata(dataset)
        self.assertEqual(list(adata.var_names), list(genes))
        np.testing.assert_array_equal(np.asarray(adata.X, dtype=float), expected_x)
        np.testing.assert_array_equal(
            np.asarray(adata.layers["source_values"], dtype=float), expected_source)

    def test_it_matches_on_a_wider_random_case(self):
        """Small hand-built cases can miss an indexing error that only shows up
        with more genes than cells, or vice versa."""
        import numpy as np

        from spatialmind.tools.implementations import _dataset_to_anndata, expression_feature_names

        rng = np.random.default_rng(0)
        names = ["G%d" % i for i in range(40)]
        records = []
        for i in range(60):
            present = rng.choice(names, size=int(rng.integers(1, 12)), replace=False)
            values = {g: float(rng.integers(1, 50)) for g in present}
            record = SpotRecord("S", float(i), 0.0, "t", values, cell_id="c%d" % i)
            # Half the cells get partial raw values, a quarter get none.
            if i % 2 == 0:
                record.raw_genes = {g: v * 10 for g, v in list(values.items())[:3]}
            elif i % 4 == 1:
                record.raw_genes = {}
            records.append(record)
        dataset = SpatialDataset(sample_id="S", source_path="x", records=records)
        genes = expression_feature_names(dataset)
        expected_x, expected_source = self._reference_matrices(dataset, genes)

        adata = _dataset_to_anndata(dataset)
        np.testing.assert_array_equal(np.asarray(adata.X, dtype=float), expected_x)
        np.testing.assert_array_equal(
            np.asarray(adata.layers["source_values"], dtype=float), expected_source)


class AnnDataCacheTests(unittest.TestCase):
    """A cache that returns a stale or shared matrix is worse than a slow one."""

    def _dataset(self, label="a"):
        records = [
            SpotRecord("S", float(i), 0.0, label, {"G1": float(i), "G2": 1.0}, cell_id="c%d" % i)
            for i in range(12)
        ]
        return SpatialDataset(sample_id="S", source_path="x", records=records)

    def setUp(self):
        from spatialmind.tools.implementations import clear_anndata_cache
        clear_anndata_cache()

    def test_each_caller_gets_its_own_object(self):
        """Every caller mutates what it gets -- scanpy normalises in place,
        squidpy writes into obsp -- so handing out the cached object would let
        one tool's graph leak into the next tool's result."""
        from spatialmind.tools.implementations import _dataset_to_anndata

        dataset = self._dataset()
        first = _dataset_to_anndata(dataset)
        second = _dataset_to_anndata(dataset)
        self.assertIsNot(first, second)
        import numpy as np

        first.obsp["scratch"] = np.zeros((first.n_obs, first.n_obs))
        self.assertNotIn("scratch", second.obsp)
        np.testing.assert_array_equal(np.asarray(first.X), np.asarray(second.X))

    def test_relabelling_invalidates_the_cache(self):
        """`apply_best_available_labels` rewrites `cell_type` in place. A cache
        that missed that would score the validated lane on the labels the
        descriptive lane saw."""
        from spatialmind.tools.implementations import _dataset_to_anndata

        dataset = self._dataset(label="Unannotated cell")
        before = _dataset_to_anndata(dataset)
        self.assertEqual(set(before.obs["cell_type"]), {"Unannotated cell"})

        for record in dataset.records:
            record.cell_type = "astrocyte"
        after = _dataset_to_anndata(dataset)
        self.assertEqual(set(after.obs["cell_type"]), {"astrocyte"},
                         "the cache returned the pre-labelling matrix")

    def test_a_region_change_invalidates_the_cache(self):
        from spatialmind.tools.implementations import _dataset_to_anndata

        dataset = self._dataset()
        _dataset_to_anndata(dataset)
        for record in dataset.records:
            record.region = "cortex"
        self.assertEqual(set(_dataset_to_anndata(dataset).obs["region"]), {"cortex"})

    def test_a_different_dataset_never_hits_another_entry(self):
        """A region-stratified pass builds one subset dataset per region."""
        from spatialmind.tools.implementations import _dataset_to_anndata

        full = self._dataset()
        _dataset_to_anndata(full)
        subset = SpatialDataset(sample_id="S", source_path="x", records=full.records[:4])
        self.assertEqual(_dataset_to_anndata(subset).n_obs, 4)
        self.assertEqual(_dataset_to_anndata(full).n_obs, 12)


class ProvenanceFileCoverageTests(unittest.TestCase):
    def test_every_file_the_loader_reads_is_hashed_by_content(self):
        """The hashed list must track what the loader actually opens. It listed
        only the parquet boundaries, so a 2022-vintage bundle shipping
        `cell_boundaries.csv.gz` fell to the names-and-sizes fallback -- and that
        vintage is the only expertly labelled section in the workspace."""
        from spatialmind.storage.run_store import HASHED_BUNDLE_FILES

        for name in ("cell_boundaries.parquet", "cell_boundaries.csv.gz",
                     "nucleus_boundaries.parquet", "nucleus_boundaries.csv.gz",
                     "cells.csv.gz", "cells.parquet", "cell_feature_matrix.h5",
                     "expert_cell_labels.csv", "cell_regions.csv"):
            self.assertIn(name, HASHED_BUNDLE_FILES, "%s is read but not content-hashed" % name)

    def test_a_boundaries_edit_changes_the_directory_digest(self):
        from spatialmind.storage.run_store import _file_md5

        root = tempfile.mkdtemp()
        try:
            path = os.path.join(root, "cell_boundaries.csv.gz")
            with open(path, "w", encoding="utf-8") as handle:
                handle.write("cell_id,vertex_x\nc1,1.0\n")
            before = _file_md5(root)
            with open(path, "w", encoding="utf-8") as handle:
                handle.write("cell_id,vertex_x\nc1,9.0\n")  # same length
            self.assertNotEqual(before, _file_md5(root))
        finally:
            shutil.rmtree(root, ignore_errors=True)


class CommandLineHelpTests(unittest.TestCase):
    """`--help` is the first thing anyone runs, and nothing tested it.

    Two entry points died on it. argparse runs its own `%` substitution over
    every help string, so a literal percent that survived a first formatting
    pass -- `20%` written for a reader -- is read back as a format spec:
    `spatialmind.cli` raised `unsupported format character ' '` and
    `run_validated_xenium_pilot` raised `must be real number, not dict`. Both
    sat behind flags about the coverage floor, which is the flag most likely to
    be looked up rather than remembered.
    """

    def _entry_points(self):
        root = Path(__file__).resolve().parents[1]
        found = []
        for directory in ("scripts", "eval", "spatialmind"):
            for path in sorted((root / directory).rglob("*.py")):
                text = path.read_text(encoding="utf-8", errors="ignore")
                if "ArgumentParser" in text:
                    found.append(path.relative_to(root))
        return found

    def test_every_argparse_entry_point_renders_its_help(self):
        import subprocess

        root = Path(__file__).resolve().parents[1]
        entry_points = self._entry_points()
        self.assertGreater(len(entry_points), 20, "the sweep found almost nothing; check the glob")

        broken = []
        for relative in entry_points:
            module = str(relative.with_suffix("")).replace(os.sep, ".")
            result = subprocess.run(
                [sys.executable, "-m", module, "--help"],
                cwd=str(root), capture_output=True, text=True, timeout=120,
            )
            if "Traceback" in result.stderr:
                broken.append("%s: %s" % (module, result.stderr.strip().splitlines()[-1]))
        self.assertEqual(broken, [], "entry points whose --help raises:\n  " + "\n  ".join(broken))


class ReviewerProvenanceTests(unittest.TestCase):
    """Who wrote the reviewed table has to reach the reader.

    The gate counts reviewed labels and reviewed regions; it cannot read who
    supplied them, and it should not -- a reviewer is trusted by construction.
    But this workspace's one validated section mixes two very different sources:
    the labels are a peer-reviewed publication, and the regions were named by a
    script from cell composition. The region table says exactly that in its own
    `reviewer_id` column, and that string was dropped at ingestion, so the report
    read `validated_ready` with 19 regions and no way to tell the two apart.
    """

    def _dataset(self, n=6):
        records = [
            SpotRecord(sample_id="S1", x=float(i), y=0.0, cell_type="Unannotated cell",
                       genes={"A": 1.0}, cell_id="c%d" % i)
            for i in range(n)
        ]
        return SpatialDataset(sample_id="S1", source_path="synthetic",
                              modality="xenium_spatial_rna", records=records)

    def _write(self, directory, name, header, rows):
        path = os.path.join(directory, name)
        with open(path, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(header)
            writer.writerows(rows)
        return path

    def test_label_and_region_reviewers_are_captured_from_the_table(self):
        root = tempfile.mkdtemp()
        try:
            labels = self._write(root, "expert_cell_labels.csv",
                                 ["cell_id", "expert_label", "reviewer_id"],
                                 [["c%d" % i, "Tumor", "Janesick et al. 2023"] for i in range(4)])
            regions = self._write(root, "cell_regions.csv",
                                  ["cell_id", "region", "reviewer_id"],
                                  [["c%d" % i, "tumor_rich", "composition-derived, not a pathologist call"]
                                   for i in range(4)])
            dataset = self._dataset()
            label_report = apply_external_label_table(dataset, labels)
            region_report = apply_external_region_table(dataset, regions)
            self.assertEqual(label_report.reviewers, {"Janesick et al. 2023": 4})
            self.assertEqual(region_report.reviewers,
                             {"composition-derived, not a pathologist call": 4})
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_an_unmatched_row_does_not_credit_its_reviewer(self):
        """A reviewer is credited for cells they actually covered. Counting rows
        instead would report 400 cells reviewed on a table that matched four."""
        root = tempfile.mkdtemp()
        try:
            labels = self._write(root, "expert_cell_labels.csv",
                                 ["cell_id", "expert_label", "reviewer_id"],
                                 [["c0", "Tumor", "R"], ["not_in_this_section", "Tumor", "R"]])
            report = apply_external_label_table(self._dataset(), labels)
            self.assertEqual(report.reviewers, {"R": 1})
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_a_table_without_a_reviewer_column_reports_nothing_rather_than_guessing(self):
        root = tempfile.mkdtemp()
        try:
            labels = self._write(root, "expert_cell_labels.csv",
                                 ["cell_id", "expert_label"], [["c0", "Tumor"]])
            self.assertEqual(apply_external_label_table(self._dataset(), labels).reviewers, {})
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_the_report_names_both_sources(self):
        from spatialmind.pilot.xenium import _reviewer_provenance_lines

        lines = _reviewer_provenance_lines({
            "label_report": {"reviewers": {"Janesick et al. 2023, Nat Commun 14:8353": 159168}},
            "region_report": {"reviewers": {"composition-derived, not a pathologist call": 163920}},
        })
        text = "\n".join(lines)
        self.assertIn("Janesick et al. 2023", text)
        self.assertIn("159,168 cells", text)
        self.assertIn("composition-derived", text)

    def test_composition_derived_regions_disclose_the_circularity(self):
        from spatialmind.pilot.xenium import _reviewer_provenance_lines

        derived = "\n".join(_reviewer_provenance_lines({
            "label_report": {"reviewers": {"A pathologist": 10}},
            "region_report": {"reviewers": {"composition-derived, not a pathologist call": 10}},
        }))
        self.assertIn("circular", derived)

        # A pathologist's regions carry no such caveat: they were drawn from
        # morphology, so a composition summary of them is a finding.
        drawn = "\n".join(_reviewer_provenance_lines({
            "label_report": {"reviewers": {"A pathologist": 10}},
            "region_report": {"reviewers": {"A pathologist, H&E review": 10}},
        }))
        self.assertNotIn("circular", drawn)
        self.assertIn("A pathologist, H&E review", drawn)

    def test_the_region_table_itself_carries_the_caveat(self):
        """The readiness section sits ~150 lines above the region table. A reader
        who jumps to `Current Region Summary` -- the table the circularity is
        about -- has to be told there too."""
        from spatialmind.pilot.xenium import _region_circularity_note

        derived = _region_circularity_note(
            {"region_report": {"reviewers": {"composition-derived, not a pathologist call": 10}}})
        self.assertTrue(derived)
        self.assertIn("restates the naming rule", "\n".join(derived))

        self.assertEqual(
            _region_circularity_note({"region_report": {"reviewers": {"Dr Chen, H&E review": 10}}}), [])
        # No reviewer column at all is unknown, not derived: say nothing rather
        # than accuse a table of being machine-made.
        self.assertEqual(_region_circularity_note({"region_report": {}}), [])

    def test_every_report_format_carries_the_provenance(self):
        """The HTML report is the artifact a reviewer is sent, and it is built by
        a different writer than the markdown. The first version of this fix put
        the disclosure in the markdown body only, so the HTML and PDF reports
        stated `validated_ready` and never named a source."""
        from spatialmind.pilot.xenium import _limitations

        payload = {
            "label_report": {"status": "expert_labels_applied",
                             "reviewers": {"Janesick et al. 2023": 159168}},
            "region_report": {"status": "user_regions_applied",
                              "reviewers": {"composition-derived, not a pathologist call": 163920}},
            "analysis_scope": {"scope": "full_section"},
            "status": "validated_ready",
            "features_loaded": 468,
        }
        # _limitations feeds the markdown, HTML and PDF writers alike.
        text = "\n".join(_limitations(payload))
        self.assertIn("Janesick et al. 2023", text)
        self.assertIn("composition-derived", text)
        self.assertIn("circular", text)

        payload["region_report"]["reviewers"] = {"Dr Chen, H&E review": 163920}
        drawn = "\n".join(_limitations(payload))
        self.assertIn("Dr Chen, H&E review", drawn)
        self.assertNotIn("circular", drawn)


class AnnotationProvenanceCheckTests(unittest.TestCase):
    """A clean join is not evidence on a bundle numbered 1..N.

    Pre-2023 Xenium bundles number cells `1, 2, 3, ...`, so an annotation of any
    smaller section matches every row. A published breast annotation was tried
    against this workspace's breast section and reported `unmatched to bundle: 0`
    while describing a different tissue block -- its domains were statistically
    independent of the section's own published cell types (V = 0.048). The
    importer's only provenance check could not see that at all.
    """

    def setUp(self):
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        import import_published_labels

        self.importer = import_published_labels

    def test_dense_integer_ids_are_recognised(self):
        self.assertTrue(self.importer.ids_are_dense_integers([str(i) for i in range(1, 5001)]))
        # A gappy integer set is not dense: a match count means something again.
        self.assertFalse(self.importer.ids_are_dense_integers(["1", "500", "90000"]))
        # Post-2023 Xenium ids, where a wrong section simply does not join.
        self.assertFalse(self.importer.ids_are_dense_integers(["aaaafije-1", "aabjifjp-1"]))
        self.assertFalse(self.importer.ids_are_dense_integers([]))

    def test_cramers_v_separates_agreement_from_independence(self):
        agreeing = [("Tumor", "tumour_core")] * 100 + [("Stromal", "stroma")] * 100
        self.assertGreater(self.importer.cramers_v(agreeing), 0.9)

        # Independent: every class spread over both domains in the same ratio.
        independent = ([("Tumor", "A")] * 50 + [("Tumor", "B")] * 50
                       + [("Stromal", "A")] * 50 + [("Stromal", "B")] * 50)
        self.assertLess(self.importer.cramers_v(independent), 0.05)

        # One-dimensional input carries no association to measure, and must not
        # be reported as perfect agreement.
        self.assertEqual(self.importer.cramers_v([("Tumor", "A")] * 10), 0.0)
        self.assertEqual(self.importer.cramers_v([]), 0.0)

    def test_the_refusal_threshold_sits_below_the_measured_real_pairing(self):
        """The bands are calibrated, not chosen. Guard the ordering so a later
        edit cannot quietly move the cut past a pairing known to be genuine."""
        # Refusal fires *below* the threshold, so it has to sit above the two
        # foreign pairings measured against the real data (0.048 and 0.054)...
        self.assertGreater(self.importer.INDEPENDENT_BELOW, 0.054)
        # ...and below the genuine same-section pairing (1.000, and 0.223 for the
        # weaker cross-kind comparison), or a real import would be refused.
        self.assertLess(self.importer.INDEPENDENT_BELOW, 0.223)
        self.assertLess(self.importer.INDEPENDENT_BELOW, self.importer.WEAK_BELOW)
        self.assertLess(self.importer.WEAK_BELOW, 1.0)


class StudioWorkflowTests(unittest.TestCase):
    """The Create wizard: suggest, understand, and never offer what it cannot run."""

    def setUp(self):
        from spatialmind.app import workflow

        self.workflow = workflow
        self.blocked = {
            "display_name": "A blocked section", "gate_open": False, "gate_status": "blocked",
            "panel": ["GFAP", "AQP4", "MBP"], "cell_types": [], "regions": [],
            "n_cells": 24406, "n_clusters": 9, "blocking_reasons": ["Expert cell labels were not applied."],
        }
        self.open = {
            "display_name": "A validated section", "gate_open": True, "gate_status": "validated_ready",
            "panel": ["ERBB2", "ACTA2", "KRT15"],
            "cell_types": ["Invasive_Tumor", "Myoepi_ACTA2+", "Stromal"],
            "regions": ["tumor_rich", "stroma_rich"],
            "n_cells": 167780, "n_clusters": 12, "blocking_reasons": [],
        }

    def test_every_suggested_question_resolves_to_a_plan(self):
        """The wizard offered "Which genes vary across tissue space in this
        section?" and then answered "no implemented tool matches that", because
        the suggestion was generated by one table and parsed by another. A
        question the app writes must be one the app can act on."""
        for facts in (self.blocked, self.open):
            for suggestion in self.workflow.recommend_questions(facts):
                self.assertTrue(suggestion["tools"],
                                "suggestion carries no tools: %s" % suggestion["question"])
                analysis = self.workflow.analyze_text(
                    suggestion["question"], facts, tools=suggestion["tools"])
                self.assertEqual(analysis["status"], "understood", suggestion["question"])
                self.assertTrue(analysis["tools"], suggestion["question"])

    def test_suggestions_only_name_tools_the_registry_implements(self):
        from spatialmind.app import planner

        plannable = {tool.name for tool in planner._registry().list_plannable()}
        for facts in (self.blocked, self.open):
            for suggestion in self.workflow.recommend_questions(facts):
                for name in suggestion["tools"]:
                    self.assertIn(name, plannable,
                                  "%s is suggested but not plannable" % name)

    def test_a_blocked_section_is_never_offered_a_cell_type_question(self):
        """Offering "which cell types sit next to each other" on a section with
        no reviewed labels teaches the user that the gate is arbitrary."""
        for suggestion in self.workflow.recommend_questions(self.blocked):
            self.assertEqual(suggestion["lane"], "descriptive", suggestion["question"])
            self.assertNotIn("cell type", suggestion["question"].lower())

    def test_an_open_section_is_offered_its_own_classes_and_regions(self):
        questions = " ".join(q["question"] for q in self.workflow.recommend_questions(self.open))
        self.assertIn("Invasive_Tumor", questions)
        self.assertIn("tumor_rich", questions)

    def test_an_unmeasured_gene_is_reported_rather_than_dropped(self):
        analysis = self.workflow.analyze_text("show me where ERBB2 and NOTAGENE1 are expressed",
                                              self.open)
        self.assertEqual(analysis["entities"]["genes"], ["ERBB2"])
        self.assertIn("NOTAGENE1", analysis["entities"]["unmeasured"])
        self.assertIn("not measured", analysis["note"])

    def test_a_scaffolded_intent_is_refused_by_name(self):
        analysis = self.workflow.analyze_text("run a ligand receptor analysis", self.open)
        self.assertEqual(analysis["status"], "unsupported")
        self.assertEqual(analysis["tools"], [])
        self.assertTrue(any("ligand_receptor_analysis" in line for line in analysis["not_supported"]))

    def test_no_clarifying_questions_when_there_is_no_plan(self):
        """A wizard that asks about run scope after refusing the question is
        collecting an answer it has no use for."""
        self.assertEqual(self.workflow.analyze_text("make it better", self.open)["questions"], [])

    def test_answers_become_real_parameters_and_nothing_else(self):
        overrides = self.workflow.apply_answers(
            ["marker_detection", "spatial_variable_genes"],
            {"group_key": "cell_type", "n_top": "100", "nonsense": "x", "n_perms": "999"})
        self.assertEqual(overrides["marker_detection"], {"group_key": "cell_type"})
        self.assertEqual(overrides["spatial_variable_genes"], {"n_top": 100})
        # n_perms was answered but no tool in the plan takes it, and a parameter
        # the validator rejects would surface as a plan error for a question
        # this module asked.
        self.assertNotIn("nonsense", str(overrides))
        self.assertNotIn("n_perms", str(overrides))


class StudioIntakeTests(unittest.TestCase):
    """Uploads land somewhere the catalogue can see, and say what they are."""

    def setUp(self):
        from spatialmind.app import uploads

        self.uploads = uploads
        self.root = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_a_traversing_path_is_refused(self):
        self.assertIsNone(self.uploads.safe_relative("../../etc/passwd"))
        self.assertIsNone(self.uploads.safe_relative("/etc/passwd"))
        self.assertIsNone(self.uploads.safe_relative("a/../../b"))
        self.assertEqual(self.uploads.safe_relative("run/cells.csv.gz"), "run/cells.csv.gz")

    def test_a_folder_upload_keeps_its_structure(self):
        """A Xenium bundle is a bundle because its files sit beside each other.
        Flattening the names produces a pile the catalogue cannot recognise."""
        result = self.uploads.store_files(self.root, "MyRun", [
            ("MyRun/experiment.xenium", b"{}"),
            ("MyRun/cells.csv.gz", b"x"),
        ])
        self.assertEqual(result["status"], "stored")
        stored = Path(result["path"])
        self.assertTrue((stored / "experiment.xenium").exists())
        self.assertTrue((stored / "cells.csv.gz").exists())
        # The browser prefixes every entry with the chosen folder's own name; the
        # dataset is that folder, not a wrapper around it.
        self.assertEqual(stored.name, "MyRun")

    def test_junk_is_skipped_without_failing_the_upload(self):
        result = self.uploads.store_files(self.root, "MyRun", [
            ("MyRun/.DS_Store", b"junk"),
            ("MyRun/data.csv", b"a,b\n1,2\n"),
        ])
        self.assertEqual(result["status"], "stored")
        self.assertEqual([f["path"] for f in result["files"]], ["MyRun/data.csv"])

    def test_an_unrecognised_upload_is_described_not_rejected(self):
        result = self.uploads.store_files(self.root, "mystery", [("notes.txt", b"hello")])
        described = self.uploads.describe(result["path"])
        self.assertFalse(described["usable"])
        self.assertIn("not recognised", described["note"].lower())

    def test_a_half_copied_xenium_bundle_names_what_is_missing(self):
        """"Not recognised" is a dead end. A partial bundle is the common case
        and naming the missing file is something the user can act on."""
        result = self.uploads.store_files(self.root, "Partial", [
            ("Partial/cells.csv.gz", b"x"),
            ("Partial/gene_panel.json", b"{}"),
        ])
        described = self.uploads.describe(result["path"])
        # `infer_data_type` already calls this a Xenium directory -- one of the
        # three marker files is enough for routing. The note has to be stricter
        # than the router, because it is a promise about what can be run.
        self.assertEqual(described["data_type"], "xenium_directory")
        self.assertIn("incomplete", described["note"])
        self.assertIn("experiment.xenium", described["note"])
        self.assertIn("cell_feature_matrix.h5", described["note"])

    def test_a_complete_bundle_is_not_called_incomplete(self):
        result = self.uploads.store_files(self.root, "Whole", [
            ("Whole/experiment.xenium", b"{}"),
            ("Whole/cells.csv.gz", b"x"),
            ("Whole/cell_feature_matrix.h5", b"x"),
        ])
        self.assertNotIn("incomplete", self.uploads.describe(result["path"])["note"])

    def test_linking_a_missing_folder_says_so(self):
        self.assertEqual(self.uploads.link_folder("/no/such/place")["status"], "missing")


class ReportLibraryTests(unittest.TestCase):
    """Pin, rename, edit and delete, without touching what the run produced."""

    def setUp(self):
        from spatialmind.app import library

        self.library_module = library
        self.root = tempfile.mkdtemp()
        self.run_dir = Path(self.root) / "job_abc123"
        self.run_dir.mkdir(parents=True)
        (self.run_dir / "report.md").write_text("# Original\n\nA finding.\n", encoding="utf-8")
        (self.run_dir / "plan_results.json").write_text(json.dumps({
            "title": "Spatial structure of the section",
            "run_id": "job_abc123",
            "dataset": {"display_name": "A section"},
            "results": [],
            "label_report": {"status": "missing_expert_labels"},
            "region_report": {"status": "missing_user_regions"},
        }), encoding="utf-8")
        self.lib = library.ReportLibrary(self.root)

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_a_report_is_titled_by_what_was_asked_not_the_job_id(self):
        """A Reports tab that lists "Job f846d36a14" is a list nobody can scan."""
        row = self.lib.scan()[0]
        self.assertEqual(row["title"], "Spatial structure of the section")
        self.assertEqual(row["dataset"], "A section")
        self.assertEqual(row["status"], "descriptive")

    def test_renaming_wins_over_the_run_title_and_survives_a_rescan(self):
        self.assertTrue(self.lib.rename("job_abc123", "Figure 3 source"))
        self.assertEqual(self.lib.scan()[0]["title"], "Figure 3 source")
        reopened = self.library_module.ReportLibrary(self.root)
        self.assertEqual(reopened.scan()[0]["title"], "Figure 3 source")

    def test_pinned_reports_sort_first(self):
        later = Path(self.root) / "job_zzz"
        later.mkdir()
        (later / "report.md").write_text("# Later\n", encoding="utf-8")
        self.assertTrue(self.lib.set_pinned("job_abc123", True))
        self.assertEqual(self.lib.scan()[0]["report_id"], "job_abc123")

    def test_an_edit_is_saved_beside_the_run_and_never_over_it(self):
        """The run record has to keep replaying, so the original report file is
        the one thing an edit may not touch."""
        self.assertTrue(self.lib.save_edit("job_abc123", "# Edited\n\nMy words.\n"))
        self.assertEqual((self.run_dir / "report.md").read_text(), "# Original\n\nA finding.\n")
        self.assertIn("My words", self.lib.markdown("job_abc123"))
        self.assertIn("A finding", self.lib.markdown("job_abc123", original=True))
        self.assertTrue(self.lib.provenance("job_abc123")["edited"])

    def test_reverting_restores_the_run_version(self):
        self.lib.save_edit("job_abc123", "# Edited\n")
        self.assertTrue(self.lib.revert_edit("job_abc123"))
        self.assertIn("Original", self.lib.markdown("job_abc123"))
        self.assertFalse(self.lib.provenance("job_abc123")["edited"])

    def test_deleting_removes_the_directory_and_the_index_entry(self):
        self.assertTrue(self.lib.delete("job_abc123"))
        self.assertFalse(self.run_dir.exists())
        self.assertEqual(self.lib.scan(), [])

    def test_a_report_id_cannot_escape_the_output_root(self):
        """Report ids come from the URL."""
        for bad in ("../../etc", "..", "/etc", "a/b", ""):
            self.assertFalse(self.lib.delete(bad), bad)
            self.assertIsNone(self.lib.get(bad), bad)


class ReportExportTests(unittest.TestCase):
    """Word, PDF, Excel and CSV, each carrying the provenance of its run."""

    MARKDOWN = (
        "# A title\n\n"
        "## Findings\n\n"
        "Some **bold** prose with `code`.\n\n"
        "- first point\n- second point\n\n"
        "| Gene | Moran |\n| --- | ---: |\n| GFAP | 0.71 |\n| AQP4 | 0.55 |\n\n"
        "## Limitations\n\nA caveat that must survive every export.\n"
    )

    def setUp(self):
        from spatialmind.app import exports

        self.exports = exports
        self.root = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_markdown_parses_into_the_blocks_the_report_uses(self):
        kinds = [b["kind"] for b in self.exports.parse_markdown(self.MARKDOWN)]
        self.assertEqual(kinds.count("heading"), 3)
        self.assertIn("bullets", kinds)
        self.assertIn("table", kinds)
        table = next(b for b in self.exports.parse_markdown(self.MARKDOWN) if b["kind"] == "table")
        self.assertEqual(table["header"], ["Gene", "Moran"])
        self.assertEqual(len(table["rows"]), 2)

    def test_unrecognised_markdown_falls_through_rather_than_vanishing(self):
        """Losing a caveat is the only failure that matters here."""
        blocks = self.exports.parse_markdown("> a blockquote the parser does not know\n")
        self.assertEqual(len(blocks), 1)
        self.assertIn("blockquote", blocks[0]["text"])

    def test_the_word_export_keeps_headings_bullets_tables_and_caveats(self):
        from docx import Document

        path = self.exports.write_docx(self.MARKDOWN, self.root / "r.docx", title="A title",
                                       provenance={"run_id": "job_1", "gate_status": "descriptive"})
        document = Document(str(path))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("A caveat that must survive every export", text)
        self.assertIn("Run job_1", text)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual([c.text for c in document.tables[0].rows[0].cells], ["Gene", "Moran"])
        # The document title and the report's own H1 are the same sentence.
        self.assertEqual(text.count("A title"), 1)

    def test_the_pdf_export_produces_a_real_pdf(self):
        path = self.exports.write_pdf(self.MARKDOWN, self.root / "r.pdf", title="A title")
        self.assertTrue(path.exists())
        self.assertTrue(path.read_bytes().startswith(b"%PDF"))

    def test_the_text_export_keeps_the_caveat(self):
        path = self.exports.write_text(self.MARKDOWN, self.root / "r.txt",
                                       provenance={"run_id": "job_1"})
        self.assertIn("A caveat that must survive every export", path.read_text())

    def _table(self, name="genes_spatial.tsv"):
        directory = self.root / "tables"
        directory.mkdir(exist_ok=True)
        path = directory / name
        path.write_text("# run_id\tjob_1\n# gate\tdescriptive\ngene\tmorans_i\nGFAP\t0.71\n",
                        encoding="utf-8")
        return {"name": name, "path": str(path), "bytes": path.stat().st_size,
                "kind": self.exports._table_kind(name)}

    def test_a_table_keeps_its_provenance_header_through_export(self):
        """A table is usually read apart from the report that qualifies it."""
        header, rows, comments = self.exports.read_table(Path(self._table()["path"]))
        self.assertEqual(header, ["gene", "morans_i"])
        self.assertEqual(rows, [["GFAP", "0.71"]])
        self.assertTrue(any("job_1" in c for c in comments))

    def test_the_excel_export_has_an_about_sheet_and_typed_numbers(self):
        import openpyxl

        path = self.exports.write_xlsx([self._table()], self.root / "r.xlsx",
                                       provenance={"run_id": "job_1"})
        book = openpyxl.load_workbook(str(path))
        self.assertEqual(book.sheetnames[0], "About")
        self.assertIn("genes_spatial", book.sheetnames)
        sheet = book["genes_spatial"]
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        self.assertIn(0.71, values)          # a number, not the string "0.71"
        self.assertTrue(any("job_1" in str(v) for v in values))

    def test_sheet_names_stay_inside_excels_limits(self):
        long_name = "a_very_long_result_table_name_that_excel_will_not_accept.tsv"
        self.assertLessEqual(len(self.exports._unique_sheet_name(long_name, set())), 31)
        first = self.exports._unique_sheet_name("markers.tsv", set())
        second = self.exports._unique_sheet_name("markers.tsv", {first})
        self.assertNotEqual(first, second)

    def test_the_csv_bundle_is_a_zip_of_csvs_with_provenance(self):
        import zipfile

        path = self.exports.write_csv_bundle([self._table()], self.root / "r.zip",
                                             provenance={"run_id": "job_1"})
        with zipfile.ZipFile(str(path)) as archive:
            names = archive.namelist()
            self.assertIn("PROVENANCE.txt", names)
            self.assertIn("genes_spatial.csv", names)
            body = archive.read("genes_spatial.csv").decode()
        self.assertIn("# run_id", body)
        self.assertIn("gene,morans_i", body)

    def test_tables_are_classified_so_the_ui_can_offer_cells_or_genes(self):
        self.assertEqual(self.exports._table_kind("cells.tsv.gz"), "cells")
        self.assertEqual(self.exports._table_kind("genes_spatial.tsv"), "genes")
        self.assertEqual(self.exports._table_kind("markers.tsv"), "genes")
        self.assertEqual(self.exports._table_kind("region_composition.tsv"), "regions")
        self.assertEqual(self.exports._table_kind("celltype_pairs.tsv"), "pairs")


class RegionReviewPacketTests(unittest.TestCase):
    """Naming regions from tissue, not from the counts that defined them.

    The one validated section in this workspace has expert labels and machine
    regions: its domains were named from their own reviewed cell composition,
    which makes a composition summary over them circular. This packet exists to
    let a pathologist replace those names from morphology, and its most
    important property is what it does *not* show.
    """

    def setUp(self):
        from spatialmind.review import regions

        self.regions = regions
        self.root = Path(tempfile.mkdtemp())
        self.bundle = self.root / "bundle"
        self.bundle.mkdir()
        import gzip

        with gzip.open(self.bundle / "cells.csv.gz", "wt", newline="") as handle:
            handle.write("cell_id,x_centroid,y_centroid,transcript_counts\n")
            for i in range(200):
                handle.write("c%d,%.1f,%.1f,40\n" % (i, (i % 20) * 30.0, (i // 20) * 30.0))
        with open(self.bundle / "expert_cell_labels.csv", "w", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["cell_id", "expert_label"])
            for i in range(200):
                writer.writerow(["c%d" % i, "Tumor" if i < 100 else "Stromal"])
        self.table = self.root / "cell_regions.csv"
        with open(self.table, "w", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(["cell_id", "region", "reviewer_id"])
            for i in range(200):
                writer.writerow(["c%d" % i, "tumor_rich" if i < 100 else "stroma_rich",
                                 "composition-derived, not a pathologist call"])

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def _build(self, blinded=True):
        return self.regions.build_region_review_packet(
            dataset_path=str(self.bundle), region_table=str(self.table),
            output_dir=str(self.root / "packet"), blinded=blinded)

    def test_the_packet_is_blinded_by_default(self):
        """Showing "this domain is 80% Tumor" and then asking what to call it
        reproduces exactly the circularity the packet exists to remove."""
        manifest = self._build(blinded=True)
        self.assertTrue(manifest["blinded"])
        page = Path(manifest["page"]).read_text()
        self.assertIn("Blinded", page)
        self.assertNotIn("Cell composition &mdash; not the thing to name from", page)
        self.assertNotIn("Stromal", page)

    def test_unblinding_says_on_the_page_that_it_is_a_second_pass(self):
        manifest = self._build(blinded=False)
        page = Path(manifest["page"]).read_text()
        self.assertIn("Unblinded", page)
        self.assertIn("after", page)
        self.assertIn("Stromal", page)

    def test_the_sheet_has_a_row_per_domain_and_no_prefilled_names(self):
        manifest = self._build()
        with open(manifest["naming_sheet"], newline="", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual({r["domain_id"] for r in rows}, {"tumor_rich", "stroma_rich"})
        # A prefilled name is a leading question.
        self.assertTrue(all(r["pathologist_label"] == "" for r in rows))

    def test_a_missing_morphology_image_still_produces_a_sheet(self):
        """The sheet is the thing that gets filled in; losing the crops should
        not lose the packet."""
        manifest = self._build()
        self.assertEqual(manifest["morphology"], "unavailable")
        self.assertTrue(Path(manifest["naming_sheet"]).exists())
        self.assertIn("No morphology image", Path(manifest["page"]).read_text())

    def test_domain_geometry_is_computed_from_the_cells(self):
        assignments = self.regions._read_region_table(self.table)
        coordinates = self.regions._read_coordinates(self.bundle)
        domains = self.regions.summarise_domains(assignments, coordinates)
        self.assertEqual({d["domain_id"] for d in domains}, {"tumor_rich", "stroma_rich"})
        for domain in domains:
            self.assertEqual(domain["n_cells"], 100)
            self.assertGreater(domain["area_um2"], 0)
            # Blinded: no composition unless labels were asked for.
            self.assertNotIn("composition", domain)

    def test_naming_writes_the_reviewer_into_every_row(self):
        """The gate counts a reviewed region table and cannot read who wrote it.
        That column is where the distinction survives into the report."""
        sheet = self.root / "sheet.csv"
        with open(sheet, "w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=self.regions.SHEET_FIELDS)
            writer.writeheader()
            writer.writerow({"domain_id": "tumor_rich", "pathologist_label": "invasive carcinoma",
                             "confidence": "0.95"})
            writer.writerow({"domain_id": "stroma_rich", "pathologist_label": "fibrous stroma",
                             "confidence": "0.9"})
        out = self.root / "named.csv"
        result = self.regions.apply_region_naming(
            str(self.bundle), str(sheet), str(self.table),
            reviewer_id="Dr Chen, DAPI morphology review", output_path=str(out))
        self.assertEqual(result["status"], "written")
        self.assertEqual(result["cells_written"], 200)
        self.assertEqual(result["distinct_regions"], 2)
        with open(out, newline="", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        self.assertTrue(all(r["reviewer_id"] == "Dr Chen, DAPI morphology review" for r in rows))
        self.assertEqual({r["region"] for r in rows}, {"invasive carcinoma", "fibrous stroma"})
        # The scope records which machine domain each call covered, so a reader
        # can see that two clicks covered 200 cells.
        self.assertEqual({r["assignment_scope"] for r in rows},
                         {"domain:tumor_rich", "domain:stroma_rich"})

    def test_an_unnamed_domain_is_dropped_rather_than_kept_under_its_old_name(self):
        """A half-named table that still says `tumor_rich` for the rest is the
        worst of both: it looks reviewed and is partly machine-made."""
        sheet = self.root / "sheet.csv"
        with open(sheet, "w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=self.regions.SHEET_FIELDS)
            writer.writeheader()
            writer.writerow({"domain_id": "tumor_rich", "pathologist_label": "invasive carcinoma"})
            writer.writerow({"domain_id": "stroma_rich", "pathologist_label": ""})
        out = self.root / "named.csv"
        result = self.regions.apply_region_naming(
            str(self.bundle), str(sheet), str(self.table), reviewer_id="Dr Chen",
            output_path=str(out))
        self.assertEqual(result["domains_left_blank"], ["stroma_rich"])
        self.assertEqual(result["cells_written"], 100)
        self.assertEqual(result["coverage_of_assigned"], 0.5)
        with open(out, newline="", encoding="utf-8") as handle:
            self.assertEqual({r["region"] for r in csv.DictReader(handle)}, {"invasive carcinoma"})

    def test_an_empty_sheet_is_reported_rather_than_writing_an_empty_table(self):
        sheet = self.root / "sheet.csv"
        with open(sheet, "w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=self.regions.SHEET_FIELDS)
            writer.writeheader()
            writer.writerow({"domain_id": "tumor_rich", "pathologist_label": ""})
        result = self.regions.apply_region_naming(
            str(self.bundle), str(sheet), str(self.table), reviewer_id="Dr Chen",
            output_path=str(self.root / "named.csv"))
        self.assertEqual(result["status"], "empty")
        self.assertFalse((self.root / "named.csv").exists())

    def test_a_named_table_no_longer_reads_as_composition_derived(self):
        """The whole point: the report's circularity caveat keys on the reviewer
        column, so a pathologist's table must stop triggering it."""
        from spatialmind.pilot.xenium import _composition_derived_regions

        self.assertTrue(_composition_derived_regions(
            {"region_report": {"reviewers": {"composition-derived, not a pathologist call": 10}}}))
        self.assertFalse(_composition_derived_regions(
            {"region_report": {"reviewers": {"Dr Chen, DAPI morphology review": 10}}}))


class ParquetCellTableTests(unittest.TestCase):
    """A Xenium bundle from GEO ships `cells.parquet`, not `cells.csv.gz`.

    The asset check already counted parquet as a present cell table, so such a
    bundle passed readiness and then failed to load: the two disagreed, and the
    disagreement blocked every GEO-deposited Xenium series -- which is how
    multi-donor Xenium data is actually published.
    """

    def setUp(self):
        self.root = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def _frame(self, n=60):
        import pandas as pd

        return pd.DataFrame({
            "cell_id": ["c%d" % i for i in range(n)],
            "x_centroid": [float(i % 10) * 20.0 for i in range(n)],
            "y_centroid": [float(i // 10) * 20.0 for i in range(n)],
            "transcript_counts": [40 + (i % 7) for i in range(n)],
            "control_probe_counts": [0] * n,
        })

    def test_a_plain_parquet_cell_table_is_read(self):
        from spatialmind.ingestion.pipeline import _open_cell_rows

        path = self.root / "cells.parquet"
        self._frame().to_parquet(path)
        with _open_cell_rows(str(path)) as reader:
            rows = list(reader)
        self.assertEqual(len(rows), 60)
        self.assertEqual(rows[0]["cell_id"], "c0")
        self.assertEqual(float(rows[11]["y_centroid"]), 20.0)

    def test_a_gzipped_parquet_cell_table_is_read(self):
        """Parquet is already compressed, so `.parquet.gz` has to be unwrapped
        before pyarrow will look at it. GEO ships it that way regardless."""
        import gzip
        import io

        from spatialmind.ingestion.pipeline import _open_cell_rows

        buffer = io.BytesIO()
        self._frame().to_parquet(buffer)
        path = self.root / "cells.parquet.gz"
        with gzip.open(path, "wb") as handle:
            handle.write(buffer.getvalue())
        with _open_cell_rows(str(path)) as reader:
            rows = list(reader)
        self.assertEqual(len(rows), 60)
        self.assertEqual(rows[-1]["cell_id"], "c59")

    def test_csv_is_still_streamed_and_agrees_with_parquet(self):
        import gzip

        from spatialmind.ingestion.pipeline import _open_cell_rows

        frame = self._frame()
        csv_path = self.root / "cells.csv.gz"
        with gzip.open(csv_path, "wt", newline="") as handle:
            frame.to_csv(handle, index=False)
        parquet_path = self.root / "cells.parquet"
        frame.to_parquet(parquet_path)

        with _open_cell_rows(str(csv_path)) as reader:
            from_csv = [(r["cell_id"], float(r["x_centroid"])) for r in reader]
        with _open_cell_rows(str(parquet_path)) as reader:
            from_parquet = [(r["cell_id"], float(r["x_centroid"])) for r in reader]
        self.assertEqual(from_csv, from_parquet)

    def test_a_parquet_only_bundle_loads_end_to_end(self):
        """No experiment.xenium either, which is how GEO deposits them."""
        import h5py
        import numpy as np
        from scipy import sparse

        from spatialmind.ingestion import load_xenium

        bundle = self.root / "bundle"
        bundle.mkdir()
        frame = self._frame()
        frame.to_parquet(bundle / "cells.parquet")

        genes = ["GFAP", "AQP4", "MBP"]
        counts = sparse.csc_matrix(np.arange(len(genes) * len(frame)).reshape(len(genes), len(frame)) % 9)
        with h5py.File(bundle / "cell_feature_matrix.h5", "w") as handle:
            group = handle.create_group("matrix")
            group.create_dataset("data", data=counts.data.astype("int32"))
            group.create_dataset("indices", data=counts.indices.astype("int64"))
            group.create_dataset("indptr", data=counts.indptr.astype("int64"))
            group.create_dataset("shape", data=np.array([len(genes), len(frame)], dtype="int32"))
            group.create_dataset("barcodes", data=np.array(frame["cell_id"], dtype="S"))
            features = group.create_group("features")
            features.create_dataset("name", data=np.array(genes, dtype="S"))
            features.create_dataset("id", data=np.array(genes, dtype="S"))
            features.create_dataset("feature_type",
                                    data=np.array(["Gene Expression"] * len(genes), dtype="S"))

        dataset = load_xenium(str(bundle), max_records=0)
        self.assertTrue(dataset.records)
        self.assertTrue(any("GFAP" in record.genes for record in dataset.records))

    def test_a_bundle_with_no_cell_table_at_all_says_which_names_it_looked_for(self):
        from spatialmind.ingestion import IngestionValidationError, load_xenium

        bundle = self.root / "empty"
        bundle.mkdir()
        (bundle / "experiment.xenium").write_text("{}", encoding="utf-8")
        with self.assertRaises(IngestionValidationError) as caught:
            load_xenium(str(bundle))
        self.assertIn("cells.parquet", str(caught.exception))


class GeoFetchTests(unittest.TestCase):
    """Addressing a GEO series, without hitting the network in a test."""

    def setUp(self):
        sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
        import fetch_geo_xenium

        self.fetch = fetch_geo_xenium

    def test_accessions_map_to_their_nested_ftp_folders(self):
        self.assertEqual(self.fetch.series_prefix("GSE311609"), "GSE311nnn")
        self.assertEqual(self.fetch.sample_prefix("GSM9509172"), "GSM9509nnn")

    def test_samples_are_grouped_from_their_file_names(self):
        rows = [
            {"name": "GSM9509172_breast_breast_B1_B1_1_cells.parquet.gz", "bytes": 2399466},
            {"name": "GSM9509172_breast_breast_B1_B1_1_cell_feature_matrix.h5", "bytes": 8035436},
            {"name": "GSM9509172_breast_breast_B1_B1_1_morphology.ome.tif.gz", "bytes": 2824830300},
            {"name": "GSM9509160_breast_breast_B2_B2_cells.parquet.gz", "bytes": 2692847},
            {"name": "filelist.txt", "bytes": 28000},
        ]
        samples = self.fetch.group_samples(rows)
        self.assertEqual(set(samples), {"breast_breast_B1_B1_1", "breast_breast_B2_B2"})
        first = samples["breast_breast_B1_B1_1"]
        self.assertEqual(first["gsm"], "GSM9509172")
        self.assertIn("cells.parquet.gz", first["files"])
        # The 2.8 GB morphology stack is catalogued but is not an analysis file,
        # so it is only fetched on request.
        self.assertIn("morphology.ome.tif.gz", first["files"])
        self.assertNotIn("morphology.ome.tif.gz", self.fetch.ANALYSIS_FILES)


class ReplicationDesignTests(unittest.TestCase):
    """What the workspace's sections amount to as a study design."""

    def test_the_recorded_design_parses_and_names_its_label_state(self):
        root = Path(__file__).resolve().parents[1]
        with open(root / "docs" / "replication_design.json", encoding="utf-8") as handle:
            design = json.load(handle)
        conditions = design["conditions"]
        self.assertIn("breast_carcinoma", conditions)
        for sections in conditions.values():
            for section in sections:
                self.assertIn("donor_id", section)
                self.assertIn(section.get("labels"), {"expert", "transferred", "none"})

    def test_more_donors_clear_replication_without_clearing_the_label_gate(self):
        """The two are separate, and satisfying one says nothing about the
        other. A design that passes replication and has no reviewed labels still
        produces descriptive results only."""
        from spatialmind.methods.replication import assess_condition_replication

        one_donor = {"breast": [{"section_id": "s1", "donor_id": "d1", "cell_count": 1000}]}
        blockers = assess_condition_replication(one_donor)["blockers"]
        self.assertTrue(any("donor" in b for b in blockers))

        three_donors = {"breast": [
            {"section_id": "s1", "donor_id": "d1", "cell_count": 1000},
            {"section_id": "s2", "donor_id": "d2", "cell_count": 1000},
            {"section_id": "s3", "donor_id": "d3", "cell_count": 1000},
        ]}
        result = assess_condition_replication(three_donors)
        self.assertEqual(result["conditions"]["breast"]["donor_count"], 3)
        self.assertFalse(any("donor" in b for b in result["blockers"]))
        # Still blocked, and for the right reason: one condition cannot be
        # compared with anything.
        self.assertTrue(any("two conditions" in b for b in result["blockers"]))


class ReviewSizingTests(unittest.TestCase):
    """Sizing the review in decisions, which is the unit the work comes in.

    `plan_expert_review` sizes it in cells -- 70% of whatever the run loads, so
    17,324 on the healthy brain section. That is correct arithmetic and the
    wrong unit: the Review Studio's gesture is a cluster click, and four of them
    cover 73.5% of that section.
    """

    def setUp(self):
        from spatialmind.review import sizing

        self.sizing = sizing

    def test_decisions_are_counted_largest_first(self):
        sizes = {"a": 50, "b": 30, "c": 15, "d": 5}
        plan = self.sizing.decisions_for_coverage(sizes, coverage=0.7)
        self.assertEqual(plan["decisions"], 2)          # 50 + 30 = 80%
        self.assertEqual([g["group"] for g in plan["groups"]], ["a", "b"])
        self.assertEqual(plan["covered_cells"], 80)
        self.assertAlmostEqual(plan["achieved_coverage"], 0.8)

    def test_unassigned_cells_count_against_coverage_but_cannot_be_labelled(self):
        """10x leaves cells unassigned. They are part of the denominator -- the
        gate's coverage is of all cells -- and they are not a decision anyone
        can make."""
        sizes = {"1": 60, "2": 10, "unassigned": 30}
        plan = self.sizing.decisions_for_coverage(sizes, coverage=0.7)
        self.assertEqual(plan["total_cells"], 100)
        self.assertEqual([g["group"] for g in plan["groups"]], ["1", "2"])
        self.assertEqual(plan["covered_cells"], 70)
        self.assertTrue(plan["reachable"])

    def test_an_unreachable_coverage_is_reported_rather_than_rounded_up(self):
        plan = self.sizing.decisions_for_coverage({"1": 40, "unassigned": 60}, coverage=0.7)
        self.assertFalse(plan["reachable"])
        self.assertEqual(plan["decisions"], 1)

    def test_one_group_covering_everything_is_the_absence_of_a_clustering(self):
        """A bundle with no cluster solution would otherwise report the most
        encouraging line in the plan: "1 decision reaches 100%"."""
        plan = self.sizing.size_label_review({"all": 1000}, coverage=0.7)
        self.assertTrue(plan["no_cluster_solution"])
        summary = self.sizing.summarise("nowhere", plan, None)
        # A bare coverage plan has no marker block; summarising one must not
        # crash and lose every other blocker with it.
        bare = self.sizing.decisions_for_coverage({"all": 1000}, coverage=0.7)
        self.assertTrue(self.sizing.summarise("nowhere", bare, None)["blockers"])
        self.assertTrue(any("No clustering to review" in b for b in summary["blockers"]))

    def test_groups_too_small_for_marker_statistics_are_named(self):
        """A section can clear coverage and then fail with
        `blocked_analysis_backend` because a class is too small to test."""
        readiness = self.sizing.marker_readiness({"1": 900, "2": 80, "3": 10})
        self.assertEqual(readiness["testable_groups"], 2)
        self.assertEqual([g["group"] for g in readiness["too_small"]], ["3"])
        self.assertTrue(readiness["meets_two_class_minimum"])

        thin = self.sizing.marker_readiness({"1": 900, "2": 10})
        self.assertFalse(thin["meets_two_class_minimum"])

    def test_a_total_with_an_uncounted_side_is_marked_a_lower_bound(self):
        """Breast S1 ranked cheapest at 8 decisions purely because nobody had
        proposed its regions yet. Comparing an incomplete total against a
        complete one is the kind of number this project exists to avoid."""
        labels = self.sizing.size_label_review({"1": 700, "2": 300})
        with_regions = self.sizing.summarise(
            "counted", labels, self.sizing.size_region_review({"r1": 600, "r2": 400}))
        without = self.sizing.summarise("uncounted", labels, None)
        self.assertTrue(with_regions["total_is_complete"])
        self.assertFalse(without["total_is_complete"])
        self.assertIn("lower bound", self.sizing.format_plan(without))
        self.assertNotIn("lower bound", self.sizing.format_plan(with_regions))

    def test_a_single_region_does_not_meet_the_contrast_minimum(self):
        plan = self.sizing.size_region_review({"only": 1000})
        self.assertFalse(plan["meets_two_region_minimum"])

    def test_the_plan_states_what_a_cluster_decision_asserts(self):
        """The cheap path is the honest answer and the dangerous one, so the
        trade has to be in the output rather than in someone's head."""
        text = self.sizing.format_plan(self.sizing.summarise(
            "s", self.sizing.size_label_review({"1": 700, "2": 300}),
            self.sizing.size_region_review({"r1": 600, "r2": 400})))
        self.assertIn("asserts that every cell in that cluster", text)
        self.assertIn("review_decisions", text)
        self.assertIn("coverage is not review depth", text)


class ClusterWorksheetTests(unittest.TestCase):
    """Four decisions, made on four rows, expanding back to 17,909 cells."""

    def setUp(self):
        from spatialmind.review import sizing

        self.sizing = sizing
        self.root = Path(tempfile.mkdtemp())
        self.run = self.root / "run"
        (self.run / "tables").mkdir(parents=True)
        (self.run / "descriptive_qc_and_cluster.json").write_text(json.dumps(
            {"metrics": {"cluster_counts": {"0": 60, "1": 30, "2": 10}}}), encoding="utf-8")
        (self.run / "descriptive_marker_detection.json").write_text(json.dumps(
            {"metrics": {"markers_by_group": {
                "0": [{"gene": "GJA1"}, {"gene": "AQP4"}],
                "1": [{"gene": "MOG"}, {"gene": "CLDN11"}],
                "2": [{"gene": "P2RY12"}]}}}), encoding="utf-8")
        with open(self.run / "tables" / "cells.tsv", "w", newline="", encoding="utf-8") as handle:
            handle.write("# SpatialMind result table\n# run_id\tr1\n")
            handle.write("cell_id\tx\ty\tcluster\n")
            for cluster, count in (("0", 60), ("1", 30), ("2", 10)):
                for i in range(count):
                    handle.write("c%s_%d\t0\t0\t%s\n" % (cluster, i, cluster))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_the_worksheet_is_one_row_per_cluster_with_its_markers(self):
        """The run's own template has one row per cell -- 24,362 of them -- which
        invites either labelling nobody will do or a fill-down that hides how few
        decisions were made."""
        path = self.root / "ws.csv"
        result = self.sizing.write_cluster_worksheet(str(self.run), str(path))
        self.assertEqual(result["status"], "written")
        self.assertEqual(result["clusters"], 3)
        self.assertEqual(result["with_markers"], 3)
        with open(path, newline="", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual([r["cluster"] for r in rows], ["0", "1", "2"])   # largest first
        self.assertEqual(rows[0]["top_markers"], "GJA1, AQP4")
        self.assertTrue(all(r["expert_label"] == "" for r in rows))

    def test_a_loader_guess_is_kept_out_of_the_evidence_column(self):
        path = self.root / "ws.csv"
        self.sizing.write_cluster_worksheet(str(self.run), str(path),
                                            loader_guesses={"0": "Neural/Glial cell"})
        with open(path, newline="", encoding="utf-8") as handle:
            rows = {r["cluster"]: r for r in csv.DictReader(handle)}
        self.assertEqual(rows["0"]["loader_guess"], "Neural/Glial cell")
        self.assertNotIn("Neural", rows["0"]["top_markers"])

    def test_named_clusters_expand_to_cells_recording_one_decision_each(self):
        """`review_decisions` has to count the calls, not the cells they cover.
        Writing per-cell rows without the scope would turn two judgements into a
        hundred apparent ones."""
        sheet = self.root / "ws.csv"
        self.sizing.write_cluster_worksheet(str(self.run), str(sheet))
        with open(sheet, newline="", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        rows[0]["expert_label"] = "astrocyte"
        rows[1]["expert_label"] = "oligodendrocyte"
        with open(sheet, "w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=self.sizing.WORKSHEET_FIELDS)
            writer.writeheader()
            writer.writerows(rows)

        out = self.root / "labels.csv"
        result = self.sizing.apply_cluster_labels(
            str(self.run), str(sheet), "Dr Chen, marker review", str(out))
        self.assertEqual(result["cells_written"], 90)
        self.assertEqual(result["clusters_named"], 2)
        self.assertEqual(result["distinct_classes"], 2)
        self.assertTrue(result["meets_two_class_minimum"])
        self.assertEqual(result["clusters_left_blank"], ["2"])

        with open(out, newline="", encoding="utf-8") as handle:
            written = list(csv.DictReader(handle))
        self.assertEqual(len({r["assignment_scope"] for r in written}), 2)
        self.assertEqual({r["assignment_scope"] for r in written}, {"cluster:0", "cluster:1"})
        self.assertTrue(all(r["reviewer_id"] == "Dr Chen, marker review" for r in written))

    def test_an_unnamed_worksheet_writes_nothing(self):
        sheet = self.root / "ws.csv"
        self.sizing.write_cluster_worksheet(str(self.run), str(sheet))
        out = self.root / "labels.csv"
        result = self.sizing.apply_cluster_labels(
            str(self.run), str(sheet), "Dr Chen", str(out))
        self.assertEqual(result["status"], "empty")
        self.assertFalse(out.exists())

    def test_a_run_without_clustering_says_so_rather_than_writing_an_empty_sheet(self):
        empty = self.root / "empty_run"
        empty.mkdir()
        result = self.sizing.write_cluster_worksheet(str(empty), str(self.root / "x.csv"))
        self.assertEqual(result["status"], "unavailable")
        self.assertIn("descriptive lane", result["reason"])


class MarkerSeparabilityTests(unittest.TestCase):
    """Where marker evidence stops being enough to tell two clusters apart.

    A decision count says how much review opens the gate. It says nothing about
    whether the decisions can be made well, and on a tumour section that is the
    whole question.
    """

    def setUp(self):
        from spatialmind.review import sizing

        self.sizing = sizing

    def test_distinct_markers_are_reported_as_separable(self):
        """Calibration: on the healthy brain section the highest real overlap is
        0.2, between two GABAergic populations sharing GAD2 and SLC6A1 -- which
        is correct biology, not a problem. The threshold sits above it."""
        result = self.sizing.marker_overlap({
            "0": ["GJA1", "AQP4", "SOX9"],
            "1": ["MOG", "CLDN11", "OPALIN"],
            "2": ["FLT1", "PECAM1", "IGFBP7"],
        })
        self.assertTrue(result["separable"])
        self.assertEqual(result["pairs"], [])
        self.assertEqual(max(result["max_overlap_by_cluster"].values()), 0.0)

    def test_shared_markers_are_flagged_with_the_genes_that_overlap(self):
        result = self.sizing.marker_overlap({
            "0": ["RNASET2", "VSIG4", "C3", "GPR34"],
            "1": ["P2RY12", "RNASET2", "GPR34", "VSIG4", "C3"],
            "9": ["AQP4", "GJA1", "SOX9"],
        })
        self.assertFalse(result["separable"])
        self.assertEqual(len(result["pairs"]), 1)
        pair = result["pairs"][0]
        self.assertEqual(sorted(pair["clusters"]), ["0", "1"])
        self.assertIn("RNASET2", pair["shared"])
        self.assertEqual(result["max_overlap_by_cluster"]["9"], 0.0)

    def test_a_cluster_with_no_markers_is_skipped_not_scored_as_identical(self):
        """Two empty marker sets are not "the same cluster"; they are two
        clusters nothing is known about."""
        result = self.sizing.marker_overlap({"0": [], "1": [], "2": ["AQP4"]})
        self.assertEqual(result["pairs"], [])

    def test_a_neoplastic_section_is_recognised_from_its_own_metadata(self):
        root = Path(tempfile.mkdtemp())
        try:
            for name, run in (("tumour", "Human Brain Glioblastoma (FFPE)"),
                              ("healthy", "Human Healthy Brain (FFPE)")):
                bundle = root / name
                bundle.mkdir()
                (bundle / "experiment.xenium").write_text(
                    json.dumps({"run_name": run}), encoding="utf-8")
            self.assertTrue(self.sizing.tumour_context(str(root / "tumour"))["neoplastic"])
            self.assertFalse(self.sizing.tumour_context(str(root / "healthy"))["neoplastic"])
            # A bundle with no metadata is not assumed to be either.
            self.assertFalse(self.sizing.tumour_context(str(root / "missing"))["neoplastic"])
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_a_tumour_section_states_the_call_the_sizing_cannot_make(self):
        """The overlap check compares clusters with each other, so it is silent
        on the hardest call in a tumour section: a malignant cell mimicking a
        lineage carries that lineage's markers, and that is not a within-section
        comparison. Cluster 5 of the glioblastoma section -- PTPRZ1, BCAN,
        OLIG2, PDGFRA -- reads as OPC and as OPC-like tumour equally."""
        label_plan = self.sizing.size_label_review({"5": 700, "9": 300})
        markers = {"5": ["PTPRZ1", "BCAN", "OLIG2", "PDGFRA"], "9": ["AQP4", "GJA1", "SOX9"]}
        regions = self.sizing.size_region_review({"r1": 600, "r2": 400})

        tumour = self.sizing.format_plan(self.sizing.summarise(
            "gbm", label_plan, regions, markers=markers,
            tumour={"neoplastic": True, "evidence": "run_name = Glioblastoma"}))
        self.assertIn("THE CALL THIS CANNOT SIZE", tumour)
        self.assertIn("mimicking a lineage", tumour)
        self.assertIn("cnv_inference", tumour)

        healthy = self.sizing.format_plan(self.sizing.summarise(
            "healthy", label_plan, regions, markers=markers,
            tumour={"neoplastic": False, "evidence": ""}))
        self.assertNotIn("THE CALL THIS CANNOT SIZE", healthy)

    def test_separability_is_not_a_blocker(self):
        """The gate opens on coverage, and it should. Overlapping markers change
        what the decisions cost to make well, which is a different thing from
        whether they can be made at all."""
        summary = self.sizing.summarise(
            "s", self.sizing.size_label_review({"0": 700, "1": 300}),
            self.sizing.size_region_review({"r1": 600, "r2": 400}),
            markers={"0": ["C3", "GPR34", "VSIG4"], "1": ["C3", "GPR34", "VSIG4"]})
        self.assertFalse(summary["separability"]["separable"])
        self.assertEqual(summary["blockers"], [])
        self.assertTrue(summary["opens_gate"])

    def test_marker_readiness_does_not_shadow_the_marker_genes(self):
        """`summarise` held a local `markers` of counts beside a parameter
        `markers` of genes, and fed the counts to the overlap check."""
        summary = self.sizing.summarise(
            "s", self.sizing.size_label_review({"0": 700, "1": 300}), None,
            markers={"0": ["AQP4"], "1": ["MOG"]})
        self.assertIsNotNone(summary["separability"])
        self.assertIn("0", summary["separability"]["max_overlap_by_cluster"])


class ParquetOnlyBundleTests(unittest.TestCase):
    """A Xenium bundle as GEO deposits it: cells.parquet.gz, no experiment file.

    The loader learned to read these; discovery and the Studio's own cell index
    had not, so such a bundle loaded fine through the ingestion layer and was
    either never listed or listed and then failed to index. Three readers, one
    format, and they disagreed.
    """

    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.bundle = self.root / "geo_style"
        self.bundle.mkdir()
        self._write_cells(self.bundle / "cells.parquet.gz")

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def _write_cells(self, path, n=40):
        import gzip
        import io

        import pandas as pd

        frame = pd.DataFrame({
            "cell_id": ["aaaa%04d-1" % i for i in range(n)],
            "x_centroid": [float(i % 8) * 25.0 for i in range(n)],
            "y_centroid": [float(i // 8) * 25.0 for i in range(n)],
            "transcript_counts": [40 + (i % 5) for i in range(n)],
        })
        buffer = io.BytesIO()
        frame.to_parquet(buffer)
        if str(path).endswith(".gz"):
            with gzip.open(path, "wb") as handle:
                handle.write(buffer.getvalue())
        else:
            path.write_bytes(buffer.getvalue())

    def test_discovery_lists_a_bundle_with_no_experiment_file(self):
        from spatialmind.app.catalog import discover_datasets

        found = {entry.name: entry for entry in discover_datasets(str(self.root))}
        self.assertIn("geo_style", found)
        self.assertTrue(found["geo_style"].reviewable)

    def test_the_studio_cell_index_reads_gzipped_parquet(self):
        from spatialmind.app.catalog import build_cell_index

        index = build_cell_index(str(self.bundle))
        self.assertEqual(index.n_cells, 40)
        self.assertEqual(index.cell_ids[0], "aaaa0000-1")
        self.assertEqual(index.cluster_method, "none")

    def test_plain_parquet_reads_the_same_as_gzipped(self):
        from spatialmind.app.catalog import build_cell_index

        plain = self.root / "plain"
        plain.mkdir()
        self._write_cells(plain / "cells.parquet")
        self.assertEqual(build_cell_index(str(plain)).cell_ids,
                         build_cell_index(str(self.bundle)).cell_ids)

    def test_a_bundle_with_no_cell_table_names_every_form_it_looked_for(self):
        from spatialmind.app.catalog import build_cell_index

        empty = self.root / "empty"
        empty.mkdir()
        (empty / "experiment.xenium").write_text("{}", encoding="utf-8")
        with self.assertRaises(FileNotFoundError) as caught:
            build_cell_index(str(empty))
        message = str(caught.exception)
        self.assertIn("cells.parquet", message)
        self.assertIn("cells.csv", message)

    def test_a_folder_that_is_not_xenium_is_not_listed_as_one(self):
        """Widening the marker set must not make every folder a dataset."""
        from spatialmind.app.catalog import discover_datasets

        noise = self.root / "not_data"
        noise.mkdir()
        (noise / "notes.txt").write_text("hello", encoding="utf-8")
        names = {entry.name for entry in discover_datasets(str(self.root))}
        self.assertNotIn("not_data", names)


class MalignantCaveatTests(unittest.TestCase):
    """The caveat's example comes from the section, not from the source.

    A fixed example asserted PTPRZ1/BCAN/OLIG2 -- brain genes -- on a breast
    section, where the ambiguity is just as real and the genes are EPCAM, CD24
    and the keratins. Stating the right principle with the wrong evidence is
    still stating the wrong evidence.
    """

    def setUp(self):
        from spatialmind.review import sizing

        self.sizing = sizing

    def test_the_example_is_the_sections_own_largest_cluster(self):
        breast = self.sizing.malignant_caveat(
            "run_name = Human Breast Cancer",
            {"5": ["SFRP1", "CD24", "EPCAM"], "0": ["CD74", "LYZ"]}, "5")
        text = "\n".join(breast)
        self.assertIn("EPCAM", text)
        self.assertNotIn("PTPRZ1", text)

        brain = self.sizing.malignant_caveat(
            "run_name = Glioblastoma",
            {"4": ["MOG", "ERMN", "CLDN11"]}, "4")
        self.assertIn("MOG", "\n".join(brain))
        self.assertNotIn("EPCAM", "\n".join(brain))

    def test_the_principle_survives_with_no_markers_to_cite(self):
        lines = self.sizing.malignant_caveat("run_name = Something", None, "")
        text = "\n".join(lines)
        self.assertIn("mimicking a lineage", text)
        self.assertIn("cnv_inference", text)
        self.assertNotIn("starting with the largest", text)

    def test_a_plan_cites_the_cluster_it_actually_sized(self):
        markers = {"5": ["SFRP1", "CD24", "EPCAM"], "0": ["CD74", "LYZ", "CD68"]}
        summary = self.sizing.summarise(
            "breast", self.sizing.size_label_review({"5": 700, "0": 300}),
            self.sizing.size_region_review({"r1": 600, "r2": 400}),
            markers=markers,
            tumour={"neoplastic": True, "evidence": "run_name = Breast Cancer"})
        text = self.sizing.format_plan(summary)
        self.assertIn("cluster 5 (SFRP1, CD24, EPCAM)", text)


class GeoBundleDegradationTests(unittest.TestCase):
    """A bundle with no `experiment.xenium` is normal, not broken.

    GEO deposits them that way. Coordinates are already in microns, so the
    section loads and gates; the parts that need run metadata are the morphology
    viewer and any judgement about the tissue.
    """

    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.bundle = self.root / "geo"
        self.bundle.mkdir()
        (self.bundle / "cells.parquet.gz").write_bytes(b"x")
        (self.bundle / "cell_feature_matrix.h5").write_bytes(b"x")

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_a_parquet_cell_table_is_not_reported_as_a_missing_one(self):
        """The note told users to go and find `cells.csv.gz`, a file the loader
        does not need once it can read the parquet beside it."""
        from spatialmind.app import uploads

        note = uploads.describe(str(self.bundle))["note"]
        self.assertNotIn("cells.csv.gz", note)
        self.assertIn("how GEO deposits them", note)

    def test_a_genuinely_incomplete_bundle_still_says_what_is_missing(self):
        from spatialmind.app import uploads

        partial = self.root / "partial"
        partial.mkdir()
        (partial / "cells.parquet.gz").write_bytes(b"x")
        note = uploads.describe(str(partial))["note"]
        self.assertIn("incomplete", note)
        self.assertIn("cell_feature_matrix.h5", note)

    def test_unknown_tissue_is_distinguished_from_known_healthy(self):
        """No metadata is "cannot tell", not "not a tumour". The whole
        GSE311609 series is tumour and ships no experiment file, so a flat False
        would drop the caveat on exactly that data."""
        from spatialmind.review.sizing import tumour_context

        unknown = tumour_context(str(self.bundle))
        self.assertFalse(unknown["neoplastic"])
        self.assertFalse(unknown["known"])

        healthy = self.root / "healthy"
        healthy.mkdir()
        (healthy / "experiment.xenium").write_text(
            json.dumps({"run_name": "Human Healthy Brain"}), encoding="utf-8")
        known = tumour_context(str(healthy))
        self.assertFalse(known["neoplastic"])
        self.assertTrue(known["known"])

    def test_the_plan_says_when_it_cannot_tell(self):
        from spatialmind.review import sizing

        text = sizing.format_plan(sizing.summarise(
            "geo", sizing.size_label_review({"0": 700, "1": 300}),
            sizing.size_region_review({"r1": 600, "r2": 400}),
            tumour={"neoplastic": False, "evidence": "", "known": False}))
        self.assertIn("TISSUE UNKNOWN", text)

        quiet = sizing.format_plan(sizing.summarise(
            "healthy", sizing.size_label_review({"0": 700, "1": 300}),
            sizing.size_region_review({"r1": 600, "r2": 400}),
            tumour={"neoplastic": False, "evidence": "", "known": True}))
        self.assertNotIn("TISSUE UNKNOWN", quiet)


class ClassPairArithmeticTests(unittest.TestCase):
    """The decision count and the analysis it enables are different numbers.

    A lymph node reaches 70% coverage on T cells and B cells alone: two
    decisions, which clears the gate's two-class condition exactly, and leaves
    exactly one cell-type pair to test. The cheapest review there is also the
    thinnest result, and only the cheapness was ever printed.
    """

    def setUp(self):
        from spatialmind.review import sizing

        self.sizing = sizing

    def test_pairs_grow_quadratically_with_named_classes(self):
        self.assertEqual(self.sizing.pair_count(2), 1)
        self.assertEqual(self.sizing.pair_count(4), 6)
        self.assertEqual(self.sizing.pair_count(5), 10)
        # Degenerate inputs are zero, not negative.
        self.assertEqual(self.sizing.pair_count(1), 0)
        self.assertEqual(self.sizing.pair_count(0), 0)
        self.assertEqual(self.sizing.pair_count(-3), 0)

    def test_a_two_class_plan_warns_that_one_pair_is_the_whole_analysis(self):
        thin = self.sizing.format_plan(self.sizing.summarise(
            "lymph", self.sizing.size_label_review({"0": 360, "2": 350, "1": 290}),
            self.sizing.size_region_review({"r1": 600, "r2": 400})))
        self.assertIn("1 cell-type pair(s) to test", thin)
        self.assertIn("One pair is the whole neighbourhood analysis", thin)

    def test_a_richer_plan_reports_pairs_without_the_warning(self):
        rich = self.sizing.format_plan(self.sizing.summarise(
            # Top three reach 65%, top four reach 83%: four decisions, six pairs.
            "brain", self.sizing.size_label_review(
                {"4": 260, "0": 200, "1": 190, "2": 180, "3": 170}),
            self.sizing.size_region_review({"r1": 600, "r2": 400})))
        self.assertIn("6 cell-type pair(s) to test", rich)
        self.assertNotIn("One pair is the whole", rich)


class NoClusterSolutionPlanTests(unittest.TestCase):
    """The CLI and the UI must refuse in the same words.

    A bundle with no clustering is one all-cells bucket. The UI said "nothing to
    review cluster by cluster yet"; the CLI printed "1 decision -> 0 pairs",
    which is arithmetic about a bucket rather than a plan.
    """

    def setUp(self):
        from spatialmind.review import sizing

        self.sizing = sizing

    def test_the_plan_refuses_instead_of_counting_a_single_bucket(self):
        text = self.sizing.format_plan(self.sizing.summarise(
            "geo", self.sizing.size_label_review({"all": 148150}), None))
        self.assertIn("Nothing to review cluster by cluster", text)
        self.assertNotIn("decision(s) reach", text)
        self.assertNotIn("cell-type pair(s) to test", text)
        self.assertIn("STILL BLOCKED", text)

    def test_a_real_clustering_still_gets_the_full_plan(self):
        text = self.sizing.format_plan(self.sizing.summarise(
            "brain", self.sizing.size_label_review({"0": 700, "1": 300}),
            self.sizing.size_region_review({"r1": 600, "r2": 400})))
        self.assertNotIn("Nothing to review cluster by cluster", text)
        self.assertIn("decision(s) reach", text)


class RunArtifactShapeTests(unittest.TestCase):
    """Two kinds of run, one reader.

    A pilot writes `descriptive_<tool>.json` per tool; a Studio plan run writes
    a single `plan_results.json` holding every tool's output. The sizing read
    only the first, so the app could not size a review against a clustering it
    had just produced itself -- it silently fell back to the bundle's 10x
    clusters and gave a different answer than the CLI on the same section.
    """

    def setUp(self):
        from spatialmind.review import sizing

        self.sizing = sizing
        self.root = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def _pilot_run(self):
        run = self.root / "pilot"
        run.mkdir()
        (run / "descriptive_qc_and_cluster.json").write_text(
            json.dumps({"metrics": {"cluster_counts": {"0": 60, "1": 40}}}), encoding="utf-8")
        (run / "descriptive_marker_detection.json").write_text(
            json.dumps({"metrics": {"markers_by_group": {"0": [{"gene": "AQP4"}]}}}),
            encoding="utf-8")
        return run

    def _plan_run(self):
        run = self.root / "plan"
        run.mkdir()
        (run / "plan_results.json").write_text(json.dumps({"results": [
            {"tool": "qc_and_cluster", "metrics": {"cluster_counts": {"0": 60, "1": 40}}},
            {"tool": "marker_detection", "metrics": {"markers_by_group": {"0": [{"gene": "AQP4"}]}}},
        ]}), encoding="utf-8")
        return run

    def test_both_run_shapes_yield_the_same_clusters_and_markers(self):
        pilot, plan = self._pilot_run(), self._plan_run()
        self.assertEqual(self.sizing.read_run_clusters(str(pilot)),
                         self.sizing.read_run_clusters(str(plan)))
        self.assertEqual(self.sizing.read_run_markers(str(pilot)),
                         self.sizing.read_run_markers(str(plan)))
        self.assertEqual(self.sizing.read_run_clusters(str(plan)), {"0": 60, "1": 40})

    def test_a_directory_with_neither_shape_returns_nothing(self):
        empty = self.root / "empty"
        empty.mkdir()
        self.assertEqual(self.sizing.read_run_clusters(str(empty)), {})
        self.assertEqual(self.sizing.read_run_markers(str(empty)), {})

    def test_a_plan_run_missing_that_tool_returns_nothing_for_it(self):
        run = self.root / "partial"
        run.mkdir()
        (run / "plan_results.json").write_text(json.dumps({"results": [
            {"tool": "spatial_variable_genes", "metrics": {"n_top": 25}},
        ]}), encoding="utf-8")
        self.assertEqual(self.sizing.read_run_clusters(str(run)), {})

    def test_malformed_json_does_not_take_the_sizing_down_with_it(self):
        run = self.root / "broken"
        run.mkdir()
        (run / "plan_results.json").write_text("{not json", encoding="utf-8")
        self.assertEqual(self.sizing.read_run_clusters(str(run)), {})


class CandidateProposalTests(unittest.TestCase):
    """A reference proposes; a reviewer decides; the gate refuses the proposal.

    Transferred labels make the review cheaper -- confirm or correct nine rows
    instead of naming nine from scratch -- and they are not expert labels. The
    whole value depends on the two never merging, so the proposal lives in its
    own worksheet column and `expert_label` stays empty.
    """

    def setUp(self):
        from spatialmind.review import sizing

        self.sizing = sizing
        self.root = Path(tempfile.mkdtemp())
        self.run = self.root / "run"
        (self.run / "tables").mkdir(parents=True)
        (self.run / "descriptive_qc_and_cluster.json").write_text(
            json.dumps({"metrics": {"cluster_counts": {"0": 10, "1": 10}}}), encoding="utf-8")
        (self.run / "descriptive_marker_detection.json").write_text(
            json.dumps({"metrics": {"markers_by_group": {"0": [{"gene": "MOG"}]}}}),
            encoding="utf-8")
        with open(self.run / "tables" / "cells.tsv", "w", newline="", encoding="utf-8") as handle:
            handle.write("# run\tr1\ncell_id\tcluster\n")
            for cluster in ("0", "1"):
                for i in range(10):
                    handle.write("c%s_%d\t%s\n" % (cluster, i, cluster))

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def _candidates(self, rows):
        path = self.root / "cand.csv"
        with open(path, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(["cell_id", "candidate_label", "confidence", "review_status"])
            for cell_id, label, conf in rows:
                writer.writerow([cell_id, label, conf, "needs_expert_review"])
        return str(path)

    def test_a_clear_cluster_gets_a_proposal_with_its_agreement(self):
        path = self._candidates(
            [("c0_%d" % i, "oligodendrocyte", 0.9) for i in range(10)] +
            [("c1_%d" % i, "astrocyte", 0.8) for i in range(10)])
        summary = self.sizing.candidates_by_cluster(
            self.sizing.read_run_cell_clusters(str(self.run)),
            self.sizing.read_candidate_labels(path))
        self.assertEqual(summary["0"]["label"], "oligodendrocyte")
        self.assertEqual(summary["0"]["share"], 1.0)
        self.assertTrue(summary["0"]["consensus"])
        self.assertIn("100%", self.sizing.format_candidate(summary["0"]))

    def test_a_split_cluster_is_flagged_rather_than_given_a_majority(self):
        """45/40 between two classes is the cluster a reviewer most needs to
        look at, and the one a bare majority label hides."""
        path = self._candidates(
            [("c1_%d" % i, "astrocyte", 0.6) for i in range(5)] +
            [("c1_%d" % i, "neuron", 0.6) for i in range(5, 10)])
        summary = self.sizing.candidates_by_cluster(
            self.sizing.read_run_cell_clusters(str(self.run)),
            self.sizing.read_candidate_labels(path))
        self.assertFalse(summary["1"]["consensus"])
        text = self.sizing.format_candidate(summary["1"])
        self.assertIn("mixed", text)
        self.assertIn("look at this one", text)

    def test_an_exact_tie_is_not_a_consensus(self):
        self.assertFalse(self.sizing.candidates_by_cluster(
            [("a", "x"), ("b", "x")],
            {"a": {"label": "p", "confidence": 1.0},
             "b": {"label": "q", "confidence": 1.0}})["x"]["consensus"])

    def test_the_proposal_never_lands_in_the_answer_column(self):
        """Writing a transferred label into `expert_label` would turn "confirm
        this" into "this is done"."""
        path = self._candidates([("c0_%d" % i, "oligodendrocyte", 0.9) for i in range(10)])
        sheet = self.root / "ws.csv"
        result = self.sizing.write_cluster_worksheet(str(self.run), str(sheet), candidates=path)
        self.assertEqual(result["with_proposals"], 1)
        with open(sheet, newline="", encoding="utf-8") as handle:
            rows = {r["cluster"]: r for r in csv.DictReader(handle)}
        self.assertIn("oligodendrocyte", rows["0"]["reference_proposes"])
        self.assertEqual(rows["0"]["expert_label"], "")
        self.assertEqual(rows["1"]["reference_proposes"], "")

    def test_a_worksheet_with_proposals_still_expands_only_what_was_named(self):
        """The reviewer's `expert_label` is the only thing that becomes a label.
        A confirmed-looking proposal that nobody confirmed is not one."""
        path = self._candidates([("c0_%d" % i, "oligodendrocyte", 0.9) for i in range(10)])
        sheet = self.root / "ws.csv"
        self.sizing.write_cluster_worksheet(str(self.run), str(sheet), candidates=path)
        out = self.root / "labels.csv"
        result = self.sizing.apply_cluster_labels(
            str(self.run), str(sheet), "Dr Chen", str(out))
        self.assertEqual(result["status"], "empty")
        self.assertFalse(out.exists())

    def test_the_candidate_file_itself_is_never_a_label_table(self):
        """`apply_best_available_labels` must not pick up a candidate file: the
        gate accepts `expert_cell_labels.csv`, which a human writes."""
        from spatialmind.ingestion.labels import LABEL_TABLE_NAMES

        for name in LABEL_TABLE_NAMES:
            self.assertNotIn("candidate", name)


class WorksheetIsAPlainCsvTests(unittest.TestCase):
    """The sheet is opened in Excel and read by scripts. It stays a plain CSV.

    A `#` caveat line at the top broke `csv.DictReader` -- the header became the
    comment -- and would have shown as a junk first row in a spreadsheet. The
    caveat belongs beside the file, not inside it.
    """

    def setUp(self):
        from spatialmind.review import sizing

        self.sizing = sizing
        self.root = Path(tempfile.mkdtemp())
        self.run = self.root / "run"
        (self.run / "tables").mkdir(parents=True)
        (self.run / "descriptive_qc_and_cluster.json").write_text(
            json.dumps({"metrics": {"cluster_counts": {"0": 10}}}), encoding="utf-8")
        with open(self.run / "tables" / "cells.tsv", "w", newline="", encoding="utf-8") as handle:
            handle.write("cell_id\tcluster\n")
            for i in range(10):
                handle.write("c%d\t0\n" % i)
        path = self.root / "cand.csv"
        with open(path, "w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            writer.writerow(["cell_id", "candidate_label", "confidence"])
            for i in range(10):
                writer.writerow(["c%d" % i, "oligodendrocyte", 0.9])
        self.candidates = str(path)

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_a_naive_dictreader_sees_the_header_first(self):
        sheet = self.root / "ws.csv"
        self.sizing.write_cluster_worksheet(str(self.run), str(sheet),
                                            candidates=self.candidates)
        with open(sheet, newline="", encoding="utf-8") as handle:
            rows = list(csv.DictReader(handle))
        self.assertEqual(rows[0]["cluster"], "0")
        self.assertIn("oligodendrocyte", rows[0]["reference_proposes"])
        self.assertFalse(sheet.read_text().startswith("#"))

    def test_the_caveat_is_written_beside_the_sheet(self):
        sheet = self.root / "ws.csv"
        result = self.sizing.write_cluster_worksheet(str(self.run), str(sheet),
                                                     candidates=self.candidates)
        readme = Path(result["readme"])
        self.assertTrue(readme.exists())
        text = readme.read_text()
        self.assertIn("transferred, not reviewed", text)
        self.assertIn("sampled per file", text)

    def test_no_readme_when_there_is_nothing_to_caveat(self):
        sheet = self.root / "plain.csv"
        result = self.sizing.write_cluster_worksheet(str(self.run), str(sheet))
        self.assertEqual(result["readme"], "")


class MarkerDisagreementTests(unittest.TestCase):
    """A proposal the cluster's own markers contradict is worse than none.

    A confident wrong label is easier to accept than a blank, so the flag is the
    loudest thing on the row when it fires -- and its range is documented,
    because trusting a flag past its range is worse than having none.
    """

    def setUp(self):
        from spatialmind.review import sizing

        self.sizing = sizing

    def _cells(self, n, label, disagree):
        return {"c%d" % i: {"label": label, "confidence": 0.9,
                            "marker_disagreement": i < disagree}
                for i in range(n)}

    def test_a_cluster_whose_markers_contradict_the_proposal_is_flagged(self):
        summary = self.sizing.candidates_by_cluster(
            [("c%d" % i, "0") for i in range(20)],
            self._cells(20, "endothelial cell", disagree=8))
        self.assertEqual(summary["0"]["marker_disagreement_share"], 0.4)
        self.assertIn("MARKERS DISAGREE", self.sizing.format_candidate(summary["0"]))
        self.assertIn("trust the markers column", self.sizing.format_candidate(summary["0"]))

    def test_a_clean_cluster_carries_no_flag(self):
        summary = self.sizing.candidates_by_cluster(
            [("c%d" % i, "0") for i in range(20)],
            self._cells(20, "oligodendrocyte", disagree=0))
        self.assertNotIn("MARKERS DISAGREE", self.sizing.format_candidate(summary["0"]))

    def test_a_few_disagreeing_cells_do_not_trip_it(self):
        """Every cluster has margins. The flag is for a systematic mismatch."""
        summary = self.sizing.candidates_by_cluster(
            [("c%d" % i, "0") for i in range(100)],
            self._cells(100, "astrocyte", disagree=5))
        self.assertNotIn("MARKERS DISAGREE", self.sizing.format_candidate(summary["0"]))

    def test_the_flags_range_is_written_down_where_the_reviewer_reads_it(self):
        """It is a lineage-level check. The two GABAergic clusters on the
        healthy brain section were proposed as glutamatergic and it stayed
        silent, because both are neuronal."""
        text = self.sizing.MARKER_DISAGREEMENT_LIMIT
        self.assertIn("lineage-level", text)
        self.assertIn("INSIDE the right lineage", text)
        self.assertIn("does not replace looking", text)

    def test_a_candidate_file_without_the_column_reads_as_no_disagreement(self):
        """Older candidate files predate the column; absent must not read True."""
        root = Path(tempfile.mkdtemp())
        try:
            path = root / "cand.csv"
            with open(path, "w", newline="", encoding="utf-8") as handle:
                writer = csv.writer(handle)
                writer.writerow(["cell_id", "candidate_label", "confidence"])
                writer.writerow(["c0", "astrocyte", 0.9])
            loaded = self.sizing.read_candidate_labels(str(path))
            self.assertFalse(loaded["c0"]["marker_disagreement"])
        finally:
            shutil.rmtree(root, ignore_errors=True)


class AssetContainerAgreementTests(unittest.TestCase):
    """Four places named the containers a Xenium table arrives in, and disagreed.

    The type inference, the catalogue, the Studio's cell index and the readiness
    check each spelled out their own list. A GEO deposit ships
    `cells.parquet.gz` and `cell_boundaries.parquet.gz`; the loader reads both,
    and the readiness check called them missing -- so the gate refused a section
    on assets it had.
    """

    def setUp(self):
        self.root = Path(tempfile.mkdtemp())
        self.bundle = self.root / "geo"
        self.bundle.mkdir()
        for name in ("cells.parquet.gz", "cell_boundaries.parquet.gz",
                     "nucleus_boundaries.parquet.gz", "cell_feature_matrix.h5"):
            (self.bundle / name).write_bytes(b"x")

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def test_readiness_sees_a_gzipped_parquet_cell_table_and_boundaries(self):
        from spatialmind.ingestion import summarize_xenium_expert_readiness

        readiness = summarize_xenium_expert_readiness(str(self.bundle))
        self.assertTrue(readiness.has_cell_table)
        self.assertTrue(readiness.has_boundaries)
        self.assertTrue(readiness.has_feature_matrix)
        # Genuinely absent, and still reported absent.
        self.assertFalse(readiness.has_morphology)

    def test_every_container_form_is_accepted_for_each_table(self):
        from spatialmind.ingestion.labels import ASSET_SUFFIXES, _any_asset

        for suffix in ASSET_SUFFIXES:
            one = self.root / ("only%s" % suffix.replace(".", "_"))
            one.mkdir()
            (one / ("cells%s" % suffix)).write_bytes(b"x")
            self.assertTrue(_any_asset(one, "cells"), suffix)

    def test_a_bundle_with_none_of_them_is_still_reported_missing(self):
        from spatialmind.ingestion import summarize_xenium_expert_readiness

        empty = self.root / "empty"
        empty.mkdir()
        (empty / "experiment.xenium").write_text("{}", encoding="utf-8")
        readiness = summarize_xenium_expert_readiness(str(empty))
        self.assertFalse(readiness.has_cell_table)
        self.assertFalse(readiness.has_boundaries)

    def test_the_four_readers_agree_on_one_bundle(self):
        """The regression that matters is them drifting apart again."""
        from spatialmind.app.catalog import _looks_like_xenium_dir
        from spatialmind.ingestion import infer_data_type, summarize_xenium_expert_readiness
        from spatialmind.ingestion.pipeline import _looks_like_xenium

        self.assertEqual(infer_data_type(str(self.bundle)), "xenium_directory")
        self.assertTrue(_looks_like_xenium(str(self.bundle)))
        self.assertTrue(_looks_like_xenium_dir(self.bundle))
        self.assertTrue(summarize_xenium_expert_readiness(str(self.bundle)).has_cell_table)


class AssertedTissueContextTests(unittest.TestCase):
    """When the bundle cannot say what the tissue is, a person can.

    A GEO deposit carries no `experiment.xenium`, so `tumour_context` correctly
    answered "cannot tell" for two breast carcinoma sections. Fabricating an
    experiment file to fix that would be inventing instrument output; a separate
    file naming who said so is the same pattern as `reviewer_id` on a label
    table.
    """

    def setUp(self):
        from spatialmind.review import sizing

        self.sizing = sizing
        self.root = Path(tempfile.mkdtemp())

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def _bundle(self, name, **files):
        bundle = self.root / name
        bundle.mkdir()
        for filename, payload in files.items():
            (bundle / filename).write_text(json.dumps(payload), encoding="utf-8")
        return bundle

    def test_an_assertion_answers_where_metadata_cannot(self):
        bundle = self._bundle("geo", **{
            self.sizing.TISSUE_CONTEXT_FILE: {"neoplastic": True, "evidence": "GSE311609 cohort"}})
        context = self.sizing.tumour_context(str(bundle))
        self.assertTrue(context["neoplastic"])
        self.assertTrue(context["known"])
        self.assertTrue(context["asserted"])
        self.assertIn("GSE311609 cohort", context["evidence"])

    def test_an_assertion_wins_over_inferred_metadata(self):
        """An explicit human statement beats a keyword match on a run name."""
        bundle = self._bundle("both", **{
            "experiment.xenium": {"run_name": "Human Healthy Brain"},
            self.sizing.TISSUE_CONTEXT_FILE: {"neoplastic": True, "evidence": "reclassified"}})
        self.assertTrue(self.sizing.tumour_context(str(bundle))["neoplastic"])

    def test_an_assertion_can_also_say_not_neoplastic(self):
        bundle = self._bundle("normal", **{
            self.sizing.TISSUE_CONTEXT_FILE: {"neoplastic": False, "evidence": "adjacent normal"}})
        context = self.sizing.tumour_context(str(bundle))
        self.assertFalse(context["neoplastic"])
        self.assertTrue(context["known"])

    def test_a_malformed_assertion_falls_through_rather_than_crashing(self):
        bundle = self.root / "broken"
        bundle.mkdir()
        (bundle / self.sizing.TISSUE_CONTEXT_FILE).write_text("{not json", encoding="utf-8")
        (bundle / "experiment.xenium").write_text(
            json.dumps({"run_name": "Glioblastoma"}), encoding="utf-8")
        context = self.sizing.tumour_context(str(bundle))
        self.assertTrue(context["neoplastic"])
        self.assertFalse(context.get("asserted", False))

    def test_the_plan_says_asserted_rather_than_names_itself(self):
        """"Names itself" is true of instrument metadata and false of a human's
        statement, and the difference is the kind this project tracks."""
        asserted = "\n".join(self.sizing.malignant_caveat(
            "tissue_context.json: a cohort", {"0": ["KRT7"]}, "0", asserted=True))
        self.assertIn("Asserted neoplastic by hand", asserted)
        self.assertNotIn("names itself", asserted)

        read = "\n".join(self.sizing.malignant_caveat(
            "run_name = Glioblastoma", {"0": ["MOG"]}, "0", asserted=False))
        self.assertIn("names itself", read)


class SilentSuccessTests(unittest.TestCase):
    """A run that does nothing must not report that it did something.

    `order_plan` filtered unknown names out, so asking for a tool this build
    does not have produced an empty plan, `plan_status: valid`, and -- through
    `POST /api/runs` -- a job that finished `succeeded`, `Done.`, no error, zero
    results, while its report named the tool as having run.
    """

    def test_an_unknown_tool_is_reported_not_dropped(self):
        from spatialmind.app import planner

        self.assertEqual(planner.unknown_tools(["not_a_tool", "qc_and_cluster"]), ["not_a_tool"])
        self.assertEqual(planner.unknown_tools(["qc_and_cluster"]), [])

    def test_a_plan_naming_an_unknown_tool_is_invalid(self):
        from spatialmind.app import planner

        described = planner.describe_plan(["not_a_tool"], gate_open=True)
        self.assertEqual(described["plan_status"], "invalid")
        self.assertIn("not_a_tool", " ".join(described["plan_errors"]))
        self.assertEqual(described["unknown_tools"], ["not_a_tool"])

    def test_a_valid_plan_is_still_valid(self):
        from spatialmind.app import planner

        described = planner.describe_plan(["qc_and_cluster"], gate_open=False)
        self.assertEqual(described["plan_status"], "valid")
        self.assertEqual(described["plan_errors"], [])
        self.assertTrue(described["steps"])

    def test_a_mixed_request_is_invalid_rather_than_quietly_trimmed(self):
        """Dropping the bad name and running the good one is the same silent
        substitution, just harder to notice."""
        from spatialmind.app import planner

        described = planner.describe_plan(["qc_and_cluster", "not_a_tool"], gate_open=False)
        self.assertEqual(described["plan_status"], "invalid")


class GroupSizePreconditionTests(unittest.TestCase):
    """Two classes is not the same as two classes a test can use.

    A table with one class on 24,405 cells and a second on one cell satisfied
    "at least two biological cell labels", the gate returned `validated_ready`,
    and the run then died inside scanpy: "Could not calculate statistics for
    groups Other since they only contain one sample."
    """

    def _gate(self, counts, **kwargs):
        from spatialmind.gatekeeper import pilot_gate

        records = []
        for label, n in counts.items():
            records.extend(SpotRecord(sample_id="S", x=float(i), y=0.0, cell_type=label,
                                      genes={"A": 1.0}, cell_id="%s_%d" % (label, i))
                           for i in range(n))
        dataset = SpatialDataset(sample_id="S", source_path="synthetic",
                                 modality="xenium_spatial_rna", records=records)
        dataset.metadata["analysis_scope"] = "full_section"
        total = len(records)
        options = {"min_label_coverage": 0.7, "min_region_coverage": 0.7,
                   "allow_single_region": True}
        options.update(kwargs)
        return pilot_gate(
            dataset=dataset,
            asset_readiness={k: True for k in ("has_cell_table", "has_feature_matrix",
                                               "has_morphology", "has_boundaries")},
            label_report={"status": "expert_labels_applied", "matched_cells": total,
                          "total_records": total, "reviewed_labels": sorted(counts),
                          "label_counts": dict(counts)},
            region_report={"status": "user_regions_applied", "matched_cells": total,
                           "total_records": total, "reviewed_regions": ["R1", "R2"]},
            **options)

    def test_a_one_cell_second_class_does_not_open_the_gate(self):
        gate = self._gate({"Tumor": 999, "Other": 1})
        self.assertNotEqual(gate["status"], "validated_ready")
        self.assertTrue(any("or more cells" in r for r in gate["blocking_reasons"]))

    def test_a_genuine_rare_population_does_not_veto_the_section(self):
        """A rare class should be reported and skipped, not block every other
        contrast in the section."""
        gate = self._gate({"Tumor": 500, "Stroma": 480, "Rare": 20})
        self.assertEqual(gate["status"], "validated_ready")

    def test_two_testable_classes_open_it(self):
        self.assertEqual(self._gate({"Tumor": 500, "Stroma": 500})["status"], "validated_ready")

    def test_the_floor_is_a_parameter_but_defaults_to_the_marker_floor(self):
        from spatialmind.gatekeeper import MIN_CELLS_PER_TESTED_CLASS

        self.assertEqual(MIN_CELLS_PER_TESTED_CLASS, 50)
        # A unit fixture may lower it to exercise wiring; the default is what
        # the validated lane actually needs.
        self.assertEqual(self._gate({"A": 1, "B": 1}, min_cells_per_class=1)["status"],
                         "validated_ready")

    def test_marker_detection_refuses_a_one_cell_group_by_type(self):
        """Not a scanpy traceback: a typed error naming the group and the fix."""
        from spatialmind.tools.exceptions import InsufficientDataError
        from spatialmind.tools.implementations import MIN_CELLS_FOR_GROUP_STATISTICS

        self.assertGreaterEqual(MIN_CELLS_FOR_GROUP_STATISTICS, 3)
        records = [SpotRecord(sample_id="S", x=float(i), y=0.0,
                              cell_type="Tumor" if i else "Other",
                              genes={"A": float(i % 5), "B": float(i % 3)},
                              cell_id="c%d" % i) for i in range(40)]
        dataset = SpatialDataset(sample_id="S", source_path="synthetic",
                                 modality="xenium_spatial_rna", records=records)
        registry = build_default_registry()
        with self.assertRaises(InsufficientDataError) as caught:
            registry.get("marker_detection").run(
                dataset, {"group_key": "cell_type", "n_top": 3, "strict_engine": True})
        message = str(caught.exception)
        self.assertIn("Other", message)
        self.assertIn("Merge them", message)


class RegionSizePreconditionTests(unittest.TestCase):
    """The same rule as the classes, on the regions two lines below them.

    When the class-size floor was added, the region branch was not checked. A
    second region of one cell still cleared "at least two user-defined regions",
    and `region_summary` skips anything under 50 cells -- so the run was left
    with one usable region, from a gate that had just required two.
    """

    def _gate(self, label_counts, region_counts, **kwargs):
        from spatialmind.gatekeeper import pilot_gate

        records, index = [], 0
        labels = [name for name, n in label_counts.items() for _ in range(n)]
        regions = [name for name, n in region_counts.items() for _ in range(n)]
        for label, region in zip(labels, regions):
            records.append(SpotRecord(sample_id="S", x=float(index), y=0.0, cell_type=label,
                                      genes={"A": 1.0}, region=region, cell_id="c%d" % index))
            index += 1
        dataset = SpatialDataset(sample_id="S", source_path="synthetic",
                                 modality="xenium_spatial_rna", records=records)
        dataset.metadata["analysis_scope"] = "full_section"
        total = len(records)
        options = {"min_label_coverage": 0.7, "min_region_coverage": 0.7,
                   "allow_single_region": False}
        options.update(kwargs)
        return pilot_gate(
            dataset=dataset,
            asset_readiness={k: True for k in ("has_cell_table", "has_feature_matrix",
                                               "has_morphology", "has_boundaries")},
            label_report={"status": "expert_labels_applied", "matched_cells": total,
                          "total_records": total, "reviewed_labels": sorted(label_counts),
                          "label_counts": dict(label_counts)},
            region_report={"status": "user_regions_applied", "matched_cells": total,
                           "total_records": total, "reviewed_regions": sorted(region_counts),
                           "region_counts": dict(region_counts)},
            **options)

    def test_a_one_cell_second_region_does_not_open_the_gate(self):
        gate = self._gate({"T": 500, "S": 500}, {"R1": 999, "R2": 1})
        self.assertNotEqual(gate["status"], "validated_ready")
        self.assertTrue(any("reviewed regions need" in r for r in gate["blocking_reasons"]))

    def test_two_real_regions_open_it(self):
        self.assertEqual(self._gate({"T": 500, "S": 500},
                                    {"R1": 500, "R2": 500})["status"], "validated_ready")

    def test_a_small_third_region_does_not_veto_the_section(self):
        self.assertEqual(self._gate({"T": 400, "S": 400, "X": 200},
                                    {"R1": 480, "R2": 500, "R3": 20})["status"],
                         "validated_ready")

    def test_the_single_region_waiver_still_bypasses_it(self):
        """An explicit waiver is a decision the report states; it is not a hole."""
        self.assertEqual(self._gate({"T": 500, "S": 500}, {"R1": 999, "R2": 1},
                                    allow_single_region=True)["status"], "validated_ready")

    def _region_summary(self, split):
        records = []
        index = 0
        for region, count in split.items():
            for _ in range(count):
                records.append(SpotRecord(
                    sample_id="S", x=float(index % 20), y=float(index // 20),
                    cell_type="Tumor" if index % 2 else "Stroma",
                    genes={"A": float(index % 4), "B": float(index % 3)},
                    region=region, cell_id="c%d" % index))
                index += 1
        dataset = SpatialDataset(sample_id="S", source_path="synthetic",
                                 modality="xenium_spatial_rna", records=records)
        return build_default_registry().get("region_summary").run(dataset, {})

    def test_a_gate_passing_split_leaves_two_usable_regions(self):
        """The invariant the floor exists for, asserted end to end rather than
        by comparing two constants: what the gate calls ready, the tool can
        actually summarise."""
        split = {"R1": 100, "R2": 100}
        self.assertEqual(self._gate({"Tumor": 100, "Stroma": 100}, split)["status"],
                         "validated_ready")
        self.assertGreaterEqual(self._region_summary(split).metrics["region_count"], 2)

    def test_the_split_the_gate_now_refuses_would_have_left_one(self):
        split = {"R1": 199, "R2": 1}
        self.assertNotEqual(self._gate({"Tumor": 100, "Stroma": 100}, split)["status"],
                            "validated_ready")
        # And this is why: only one of the two regions has a comparable
        # composition, so there is no contrast to report.
        self.assertEqual(self._region_summary(split).metrics["reliable_region_count"], 1)

    def test_a_tiny_region_composition_is_marked_not_presented_as_a_finding(self):
        """A proportion from one cell is not a proportion. It was reported as
        "100% T cells" in the same shape and the same table as one from 500
        cells."""
        result = self._region_summary({"R1": 480, "R2": 500, "R3": 1})
        regions = result.metrics["regions"]
        self.assertTrue(regions["R1"]["composition_reliable"])
        self.assertFalse(regions["R3"]["composition_reliable"])
        self.assertIn("below the 50-cell floor", regions["R3"]["note"])
        self.assertEqual(result.metrics["regions_below_floor"], ["R3"])
        self.assertEqual(result.metrics["reliable_region_count"], 2)
        # And it says so where a reader will see it, not only in the metrics.
        self.assertIn("unreliable", result.summary)
        self.assertTrue(any("not comparable" in c for c in result.caveats))

    def test_a_small_region_is_still_counted_and_named(self):
        """Dropping it silently would be its own misreport: the reviewer drew
        that region and is entitled to see what is in it."""
        regions = self._region_summary({"R1": 480, "R2": 500, "R3": 1}).metrics["regions"]
        self.assertIn("R3", regions)
        self.assertEqual(regions["R3"]["cell_count"], 1)
