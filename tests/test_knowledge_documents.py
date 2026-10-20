"""Document imports use synthetic fixtures, never research or personal files."""

import hashlib
import io
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient

from spatialmind.app.server import create_studio_app
from spatialmind.knowledge import documents
from spatialmind.knowledge.store import KnowledgeStore


def fixture(extension):
    stream = io.BytesIO()
    if extension == "docx":
        from docx import Document
        doc = Document()
        doc.add_paragraph("Fixturegene is an experimental marker.")
        doc.add_table(rows=1, cols=2).cell(0, 0).text = "Fixturetable"
        doc.save(stream)
    elif extension == "xlsx":
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = "Markers"
        book.active.append(["Fixturegene", "=1+2"])
        book.create_sheet("Hidden").sheet_state = "hidden"
        book["Hidden"].append(["Fixturehidden"])
        book.save(stream)
        book.close()
    elif extension == "pptx":
        from pptx import Presentation
        from pptx.util import Inches
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, Inches(4), Inches(1)).text = "Fixturegene spatial evidence"
        slide.shapes.add_table(1, 1, 0, Inches(2), Inches(4), Inches(1)).table.cell(0, 0).text = "Fixturetable"
        slide.notes_slide.notes_text_frame.text = "Fixturenotes"
        deck.save(stream)
    elif extension == "pdf":
        from reportlab.pdfgen import canvas
        pdf = canvas.Canvas(stream)
        pdf.drawString(40, 700, "Fixturegene spatial evidence")
        pdf.showPage()
        pdf.drawString(40, 700, "Fixturesecondpage")
        pdf.save()
    else:
        return {"csv": b'gene,description\nFixturegene,"quoted, value"\n',
                "tsv": b"gene\tdescription\nFixturegene\tmarker\n",
                "txt": b"Fixturegene spatial evidence", "md": b"# Fixturegene\nSpatial evidence"}[extension]
    return stream.getvalue()


class ExtractionTests(unittest.TestCase):
    def test_all_supported_formats_have_provenance_and_locators(self):
        locators = {"docx": "Paragraph 1", "xlsx": "Sheet Markers row 1", "pptx": "Slide 1",
                    "pdf": "Page 1", "csv": "Row 2", "tsv": "Row 2", "txt": "Line 1", "md": "Line 1"}
        for extension, locator in locators.items():
            with self.subTest(extension=extension):
                payload = fixture(extension)
                result = documents.extract_document("source." + extension.upper(), payload)
                self.assertIn("Fixturegene", result["body"])
                self.assertIn("[" + locator + "]", result["body"])
                self.assertEqual(result["metadata"]["file_sha256"], hashlib.sha256(payload).hexdigest())
                self.assertFalse(result["metadata"]["original_retained"])

    def test_office_tables_formulas_notes_and_hidden_sheet(self):
        doc = documents.extract_document("a.docx", fixture("docx"))
        self.assertIn("Fixturetable", doc["body"])
        book = documents.extract_document("a.xlsx", fixture("xlsx"))
        self.assertIn("=1+2", book["body"])
        self.assertIn("Fixturehidden", book["body"])
        self.assertTrue(any("hidden" in w for w in book["metadata"]["warnings"]))
        slides = documents.extract_document("a.pptx", fixture("pptx"))
        self.assertIn("Fixturetable", slides["body"])
        self.assertIn("[Slide 1 notes]\nFixturenotes", slides["body"])

    def test_unicode_and_quoted_csv(self):
        self.assertIn("quoted, value", documents.extract_document("a.csv", fixture("csv"))["body"])
        self.assertIn("marker", documents.extract_document("a.txt", "marker".encode("utf-16"))["body"])
        with self.assertRaisesRegex(ValueError, "UTF-8"):
            documents.extract_document("a.txt", b"\x80\x81")

    def test_legacy_and_mislabeled_binary_rejected(self):
        for name in ("a.doc", "a.xls", "a.ppt", "a.docm", "a.html", "a.exe"):
            with self.subTest(name=name), self.assertRaisesRegex(ValueError, "Unsupported format"):
                documents.extract_document(name, b"not a document")
        with self.assertRaisesRegex(ValueError, "signature"):
            documents.extract_document("a.pdf", b"not a PDF")
        with self.assertRaisesRegex(ValueError, "Binary"):
            documents.extract_document("a.txt", b"binary\x00")
        with self.assertRaisesRegex(ValueError, "contents"):
            documents.extract_document("a.docx", fixture("xlsx"))

    def test_file_archive_table_and_text_limits(self):
        with patch.object(documents, "MAX_FILE_BYTES", 2), self.assertRaisesRegex(ValueError, "25 MiB"):
            documents.extract_document("a.txt", b"123")
        payload = fixture("docx")
        with patch.object(documents, "MAX_ARCHIVE_BYTES", 10), self.assertRaisesRegex(ValueError, "expanded"):
            documents.extract_document("a.docx", payload)
        with patch.object(documents, "MAX_UNITS", 2), self.assertRaisesRegex(ValueError, "fields"):
            documents.extract_document("a.csv", fixture("csv"))
        with patch.object(documents, "MAX_TEXT_BYTES", 10), self.assertRaisesRegex(ValueError, "Extracted text"):
            documents.extract_document("a.txt", b"Fixturegene evidence")

    def test_macros_rejected_even_with_docx_extension(self):
        payload = io.BytesIO(fixture("docx"))
        with zipfile.ZipFile(payload, "a") as archive:
            archive.writestr("word/vbaProject.bin", b"dummy")
        with self.assertRaisesRegex(ValueError, "Macro"):
            documents.extract_document("a.docx", payload.getvalue())

    def test_blank_encrypted_and_partial_pdf(self):
        from pypdf import PdfReader, PdfWriter
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        stream = io.BytesIO()
        writer.write(stream)
        with self.assertRaisesRegex(ValueError, "OCR"):
            documents.extract_document("a.pdf", stream.getvalue())
        writer.add_page(PdfReader(io.BytesIO(fixture("pdf"))).pages[0])
        stream = io.BytesIO()
        writer.write(stream)
        result = documents.extract_document("a.pdf", stream.getvalue())
        self.assertIn("[Page 2]", result["body"])
        self.assertTrue(any("pages: 1" in w for w in result["metadata"]["warnings"]))
        writer.encrypt("fixture-password")
        stream = io.BytesIO()
        writer.write(stream)
        with self.assertRaisesRegex(ValueError, "Encrypted"):
            documents.extract_document("a.pdf", stream.getvalue())

    def test_long_page_locators_survive_retrieval_chunking(self):
        text = documents.ExtractedText()
        text.add("Page 9", "Fixturegene " * 600)
        with tempfile.TemporaryDirectory() as root:
            store = KnowledgeStore(Path(root) / "kb.sqlite")
            identity = store.add("Long page", "\n\n".join(text.parts))
            store.set_status(identity, True, True, "fixture-reviewer")
            passages = store.search("Fixturegene", 8)
            self.assertGreater(len(passages), 1)
            self.assertTrue(all(p["passage"].startswith("[Page 9]") for p in passages))

    def test_old_database_migrates_without_losing_sources(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "kb.sqlite"
            with sqlite3.connect(str(path)) as db:
                db.execute("CREATE TABLE sources(id TEXT PRIMARY KEY,title TEXT,url TEXT,version TEXT,sha256 TEXT,body TEXT,status TEXT,enabled INTEGER,reviewer TEXT,reviewed_at TEXT)")
                db.execute("INSERT INTO sources VALUES('old','Saved source','','1','digest','Saved body','pending',1,'','')")
            store = KnowledgeStore(path)
            self.assertEqual(store.source("old")["body"], "Saved body")
            self.assertEqual(store.source("old")["metadata"], {})


class UploadTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.app = create_studio_app(self.temp.name + "/data", self.temp.name + "/out", secure=True)
        self.client = TestClient(self.app, base_url="http://127.0.0.1:9876")
        self.page = self.client.get("/").text

    def tearDown(self):
        self.client.close()
        self.temp.cleanup()

    def test_upload_review_search_delete_all_formats_with_real_worker(self):
        for extension in ("txt", "md", "csv", "tsv", "docx", "xlsx", "pptx", "pdf"):
            with self.subTest(extension=extension):
                response = self.client.post("/api/knowledge/uploads", files={"file": ("fixture." + extension, fixture(extension))})
                self.assertEqual(response.status_code, 200, response.text)
                source = response.json()
                self.assertEqual(source["status"], "pending")
                identity = source["source_id"]
                self.assertFalse(self.client.get("/api/knowledge/search?q=Fixturegene").json()["passages"])
                endpoint = "/api/knowledge/source?source_id=" + identity
                self.assertEqual(self.client.get(endpoint).json()["metadata"]["format"], extension)
                self.assertEqual(self.client.patch(endpoint, json={"approve": True, "reviewer": "test-reviewer"}).status_code, 200)
                self.assertTrue(self.client.get("/api/knowledge/search?q=Fixturegene").json()["passages"])
                self.client.delete(endpoint)
                self.assertFalse(self.client.get("/api/knowledge/search?q=Fixturegene").json()["passages"])

    def test_failed_import_leaves_no_source_and_protects_endpoint(self):
        before = self.client.get("/api/knowledge/sources").json()["sources"]
        for name, payload in (("broken.docx", b"bad zip"), ("empty.txt", b""), ("old.doc", b"legacy")):
            response = self.client.post("/api/knowledge/uploads", files={"file": (name, payload)})
            self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(self.client.get("/api/knowledge/sources").json()["sources"], before)
        self.assertEqual(self.client.post("/api/knowledge/uploads", content=b"x", headers={"Content-Length": str(27 * 1024 ** 2)}).status_code, 413)
        self.assertEqual(self.client.post("/api/knowledge/uploads", files={"file": ("a.txt", b"text")}, headers={"Origin": "https://evil.example"}).status_code, 403)

    def test_biobot_and_multiple_upload_ui(self):
        self.assertIn('data-p="ask">BioBot</button>', self.page)
        self.assertIn('<h2>BioBot</h2>', self.page)
        self.assertNotIn('>Ask</', self.page)
        self.assertIn('id="knowledgeFile" multiple', self.page)
        self.assertIn('.docx,.xlsx,.csv,.tsv,.pdf,.pptx,.md,.txt', self.page)

    def test_worker_timeout_is_actionable(self):
        with patch.object(documents, "PARSE_TIMEOUT", 0), self.assertRaisesRegex(ValueError, "timed out"):
            documents.extract_isolated("a.txt", b"text")


if __name__ == "__main__":
    unittest.main()
